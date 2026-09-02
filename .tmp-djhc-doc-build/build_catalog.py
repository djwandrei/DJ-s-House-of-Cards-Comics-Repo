from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\djwan\Downloads\djshouseofcards-next-fixes-applied\DJHC_Fan_Analytics_Recommendation_Catalog.docx")

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5A6872"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "7A5A00"
RISK = "9B1C1C"
INK = "1A1A1A"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
LABEL_DXA = 2700
DETAIL_DXA = 6660


def set_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph(paragraph, before=0, after=6, line=1.25, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths=(LABEL_DXA, DETAIL_DXA), header_fill=LIGHT_BLUE, header_row=True):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if header_row and row_index == 0:
                set_cell_shading(cell, header_fill)
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "C7D0DA")
    if header_row:
        first_row_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        first_row_pr.append(repeat)


def prevent_row_split(row):
    """Keep a recommendation record together instead of orphaning its guardrail."""
    row_pr = row._tr.get_or_add_trPr()
    cant_split = row_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=8.5, color=MUTED)


def add_rule(paragraph, color="C7D0DA", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_doc_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1

    if "Catalog Note" not in styles:
        note = styles.add_style("Catalog Note", WD_STYLE_TYPE.PARAGRAPH)
        note.base_style = styles["Normal"]
        note.font.name = "Calibri"
        note._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        note._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        note.font.size = Pt(9)
        note.font.color.rgb = RGBColor.from_string(MUTED)
        note.paragraph_format.space_after = Pt(4)
        note.paragraph_format.line_spacing = 1.15


def add_header_footer(section):
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph(p, before=0, after=0, line=1.0)
    r = p.add_run("DJ'S HOUSE OF CARDS & COMICS  |  FAN ANALYTICS ROADMAP")
    set_font(r, size=8.5, color=MUTED, bold=True)
    add_rule(p, color="D7DEE7", size="4")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph(p, before=0, after=0, line=1.0)
    r = p.add_run("Recommendation Catalog  |  Page ")
    set_font(r, size=8.5, color=MUTED)
    add_page_field(p)


def add_body(doc, text, before=0, after=6, italic=False, color=INK, size=11):
    p = doc.add_paragraph()
    set_paragraph(p, before=before, after=after, line=1.25)
    r = p.add_run(text)
    set_font(r, size=size, color=color, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_font(r, size=11, color=INK)
    return p


def add_callout(doc, label, text, color=CALLOUT):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, widths=(TABLE_WIDTH_DXA,), header_fill=color)
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    set_paragraph(p, before=1, after=1, line=1.2)
    r = p.add_run(label + " ")
    set_font(r, size=10.5, color=NAVY, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)
    after = doc.add_paragraph()
    set_paragraph(after, before=0, after=4, line=1.0)


def add_table_header(cell, text):
    p = cell.paragraphs[0]
    set_paragraph(p, before=0, after=0, line=1.1)
    r = p.add_run(text)
    set_font(r, size=9, color=NAVY, bold=True)


def add_detail_paragraph(cell, label, text, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    set_paragraph(p, before=0, after=2, line=1.12)
    label_run = p.add_run(label + " ")
    set_font(label_run, size=8.8, color=DARK_BLUE, bold=True)
    text_run = p.add_run(text)
    set_font(text_run, size=8.8, color=INK)


def add_catalog_table(doc, heading, intro, entries, new_page=True):
    if new_page:
        doc.add_page_break()
    heading_p = doc.add_paragraph(heading, style="Heading 2")
    heading_p.paragraph_format.keep_with_next = True
    intro_p = add_body(doc, intro, after=5, size=10.5, color=MUTED)
    intro_p.paragraph_format.keep_with_next = True
    for index, entry in enumerate(entries):
        if index and index % 3 == 0:
            # Emit a real page-break paragraph.  Word occasionally positions a
            # paragraph with page_break_before at the previous page's boundary,
            # which can leave a continuation card clipped at the top of the
            # following rendered page.
            doc.add_page_break()
            continued = doc.add_paragraph(heading + " - continued", style="Heading 3")
            continued.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=2, cols=1)
        set_table_geometry(table, widths=(TABLE_WIDTH_DXA,))
        header_cell = table.cell(0, 0)
        detail_cell = table.cell(1, 0)
        for cell in (header_cell, detail_cell):
            set_cell_width(cell, TABLE_WIDTH_DXA)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = header_cell.paragraphs[0]
        set_paragraph(p, before=0, after=1, line=1.05)
        r = p.add_run(entry["name"])
        set_font(r, size=10.2, color=NAVY, bold=True)
        p2 = header_cell.add_paragraph()
        set_paragraph(p2, before=0, after=0, line=1.05)
        r = p2.add_run("Readiness: " + entry["readiness"])
        set_font(r, size=8.5, color=entry.get("readiness_color", MUTED), bold=True)
        add_detail_paragraph(detail_cell, "What it is:", entry["what"], first=True)
        add_detail_paragraph(detail_cell, "Data needed:", entry["data"])
        add_detail_paragraph(detail_cell, "Why implement:", entry["why"])
        if entry.get("guardrail"):
            add_detail_paragraph(detail_cell, "Guardrail:", entry["guardrail"])
        prevent_row_split(table.rows[0])
        prevent_row_split(table.rows[1])
        spacer = doc.add_paragraph()
        set_paragraph(spacer, before=0, after=2, line=1.0)
    spacing = doc.add_paragraph()
    set_paragraph(spacing, before=0, after=3, line=1.0)


def entry(name, readiness, what, data, why, guardrail=None, color=MUTED):
    return {
        "name": name,
        "readiness": readiness,
        "what": what,
        "data": data,
        "why": why,
        "guardrail": guardrail,
        "readiness_color": color,
    }


def build_document():
    doc = Document()
    section = doc.sections[0]
    add_header_footer(section)
    set_doc_styles(doc)
    doc.core_properties.title = "DJHC Fan Analytics, Games, Tools, and Services"
    doc.core_properties.subject = "Consolidated recommendation catalog"
    doc.core_properties.author = "DJ's House of Cards & Comics"
    doc.core_properties.comments = "Consolidated planning reference"

    # Editorial-cover opening, selected for a polished long-form reference guide.
    spacer = doc.add_paragraph()
    set_paragraph(spacer, before=80, after=0, line=1.0)
    kicker = doc.add_paragraph()
    set_paragraph(kicker, before=0, after=16, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = kicker.add_run("PRODUCT ROADMAP")
    set_font(r, size=10.5, color=GOLD, bold=True)
    title = doc.add_paragraph()
    set_paragraph(title, before=0, after=8, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = title.add_run("Fan Analytics, Games,\nTools, and Services")
    set_font(r, size=28, color=NAVY, bold=True)
    subtitle = doc.add_paragraph()
    set_paragraph(subtitle, before=0, after=24, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = subtitle.add_run("A consolidated recommendation catalog for DJ's House of Cards & Comics")
    set_font(r, size=14, color=DARK_BLUE)
    cover_rule = doc.add_paragraph()
    set_paragraph(cover_rule, before=0, after=15, line=1.0)
    add_rule(cover_rule, color="A8BDCE", size="8")
    cover_note = doc.add_paragraph()
    set_paragraph(cover_note, before=0, after=8, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = cover_note.add_run("Scope: existing roadmap, new games and quizzes, collector features, data services, readiness gates, and implementation guardrails")
    set_font(r, size=10.5, color=MUTED, italic=True)
    date_p = doc.add_paragraph()
    set_paragraph(date_p, before=14, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = date_p.add_run("September 2, 2026")
    set_font(r, size=11, color=NAVY, bold=True)
    doc.add_page_break()

    doc.add_paragraph("Executive summary", style="Heading 1")
    add_body(
        doc,
        "The opportunity is not a generic statistics dashboard. The strongest DJHC fan experience is a repeatable choice-and-reveal loop: a visitor builds, predicts, or solves something; the site explains the result with source-labeled evidence; and the visitor can save or share a lightweight result. This catalog consolidates every recommendation discussed so far into one operating reference.",
    )
    add_callout(
        doc,
        "Primary recommendation:",
        "Use Lineup DNA to strengthen the existing Lineup Lab, then launch Rotation Rescue as the first independent game. Pair it with Season Signature for collector value and Scout's Call or Statline Sleuth for replayable discovery.",
    )
    doc.add_paragraph("Recommended launch sequence", style="Heading 2")
    for item in (
        "Strengthen the live foundation: add Lineup DNA and shareable run cards to Lineup Lab.",
        "Launch the first independent game: Rotation Rescue, using the exact optimizer and deterministic challenge definitions.",
        "Add collector value: Season Signature for exact mapped NBA cards, then Player & Card Matchups and Team-Year Reverse Lookup as mapping coverage permits.",
        "Add repeatable discovery: Scout's Call, Statline Sleuth, Evidence Court, and Two Truths, One Box Score.",
        "Expand into data-prep and mapping-dependent concepts before touching raw or partial play-by-play.",
        "Unlock PBP-native experiences only from a compact, validated, versioned derived output with explicit reliability and coverage labels.",
    ):
        add_bullet(doc, item)

    doc.add_paragraph("Readiness and data gates", style="Heading 2")
    add_callout(
        doc,
        "Build now:",
        "Uses current Lineup Lab inputs, transparent role/rate helpers, approved browser-safe season summaries, or a curated static question bank.",
        color="EAF3EC",
    )
    add_callout(
        doc,
        "Mapping-dependent:",
        "Requires an active, verified athlete-to-product relationship and buyer-safe stat projection. Ambiguous and multi-subject identity cases remain unresolved unless every subject is verified.",
        color="FFF8E8",
    )
    add_callout(
        doc,
        "PBP validation-gated:",
        "Requires a current successful validation of compact derived play-by-play outputs. Raw provider payloads, lineups, stints, possessions, and model tables remain private.",
        color="FBEDED",
    )
    doc.add_paragraph("Non-negotiable product boundaries", style="Heading 2")
    for item in (
        "Each experience belongs in its own fan-tool route and uses browser-safe, read-only data.",
        "Scenario, challenge, and collection state should be local-first until a separate persistence decision is approved.",
        "Games never change catalog quantity, sale state, pricing, checkout, or Shopify state.",
        "A five-player PBP row is an exact lineup; two- through four-player rows are shared-floor co-presence, not an exact unit.",
        "RAPM, on/off, and WOWY are descriptive or model-based evidence, not causal proof or a final player ranking.",
    ):
        add_bullet(doc, item)

    core = [
        entry(
            "NBA Lineup Lab",
            "Current core / registry labeled live",
            "Build a five-player lineup or full rotation from a historical team-season, game plan, and user constraints.",
            "Verified NBA player and team-season data; browser-safe read-only views; exact 240-minute role-constrained optimizer.",
            "It is the anchor experience and creates the scenarios, explanations, and share links that other tools can extend.",
            "Keep exact solver behavior and source disclosure separate from storefront checkout code.",
            BLUE,
        ),
        entry(
            "Player & Card Matchups",
            "Existing planned collector tool",
            "Connect a player or Lineup Lab scenario to verified DJHC cards and useful player-season context.",
            "Active athlete identities, verified product mappings, and a single buyer-safe catalog projection.",
            "It makes sports analysis directly useful to collectors without reducing cards to title-search results.",
            "Never infer a card match from title similarity; unresolved identities stay unmatched.",
            GOLD,
        ),
        entry(
            "Build from Your Collection",
            "Existing planned collector game",
            "Turn owned or saved player cards into an eligible roster, then build the best lineup from that pool.",
            "Verified player-card relationships and local-first collection state with explicit opt-in persistence.",
            "It is the clearest bridge between collecting and the existing optimizer.",
            "No inventory, sales, quantity, or checkout mutation.",
            GOLD,
        ),
        entry(
            "Era & Roster Challenges",
            "Existing planned history game",
            "Recreate an iconic roster, satisfy a historic team brief, or solve a card-set challenge with fixed rules.",
            "Stable historical team-season facts and versioned challenge definitions with deterministic scoring.",
            "It creates shareable debate and replay without relying on speculative simulations.",
            "Version every challenge so an old shared result stays explainable after data refreshes.",
            GOLD,
        ),
        entry(
            "Trade & Package Builder",
            "Existing planned what-if tool",
            "Test hypothetical player packages and view before-and-after role coverage and game-plan fit.",
            "Lineup Lab role-scaled projections; no transaction, inventory, or payment data.",
            "It supports roster debate through transparent counterfactuals instead of a fake transaction workflow.",
            "Label every result hypothetical and provide reset controls.",
            GOLD,
        ),
        entry(
            "NBA Analytics Explorer",
            "Existing research-gated tool",
            "Explore shared-floor, on/off, context splits, adjusted impact, and RAPM after evidence validation.",
            "Licensed, validated play-by-play; dedicated analytics database; versioned derived model results.",
            "It can become a high-differentiation research surface once its data boundary is ready.",
            "Never present raw plus-minus, incomplete archives, or estimated half-court results as proven facts.",
            RISK,
        ),
        entry(
            "Slab-to-Stats",
            "Existing supporting storefront feature",
            "Show selected player-season context and verified statistics beside an eligible card after the product modal opens.",
            "Exact mapped product, buyer-safe fixed-shape stat response, and rights-confirmed media where used.",
            "It improves product discovery and sets up Season Signature, Card Clue Hunt, and player-card tools.",
            "Keep it modal-only; do not slow catalog-card or initial catalog loading.",
            BLUE,
        ),
        entry(
            "Shareable Run Cards",
            "Cross-cutting feature",
            "Create compact result cards such as a daily score, solved rotation, challenge outcome, or personal collection milestone.",
            "Local scenario state, deterministic scoring, source/version metadata, and optional share URL.",
            "It creates social value before the cost and moderation burden of accounts or global leaderboards.",
            "Start local-first; add accounts and leaderboards only after repeat usage proves the need.",
            BLUE,
        ),
        entry(
            "Daily and Seeded Challenges",
            "Cross-cutting feature",
            "Use fixed inputs, a versioned prompt, and a deterministic answer key for a daily, weekly, or evergreen challenge.",
            "Reviewed challenge definitions, reproducible scoring rules, source/version metadata, and a small publication workflow.",
            "It gives quizzes and rotation games a repeatable cadence without relying on opaque live simulations.",
            "Never silently change a prior challenge answer after its inputs or source data refresh.",
            BLUE,
        ),
        entry(
            "Source, Coverage, and Reliability Labels",
            "Cross-cutting trust feature",
            "Show the evidence type, source scope, sample/coverage, and model or reconstruction status beside analytical results.",
            "Versioned data provenance, available coverage fields, reliability rules, and a compact public explanation schema.",
            "It turns a technical safety boundary into a visible product advantage and makes results more useful to serious fans.",
            "Do not use a label as a substitute for missing validation; unavailable evidence must stay unavailable.",
            BLUE,
        ),
        entry(
            "Local-First Saved State and Watchlists",
            "Cross-cutting retention feature",
            "Let visitors save scenarios, selected cards, and watchlist ideas locally before any account system is justified.",
            "Browser local storage helpers, explicit opt-in labels, stable scenario IDs, and buyer-safe product references.",
            "It supports return visits and collector journeys without adding a new personal-data or moderation surface.",
            "Do not imply ownership, mutate inventory, or persist sensitive account data without separate approval.",
            BLUE,
        ),
        entry(
            "Explainable Scoring and Clear Endings",
            "Cross-cutting game-design feature",
            "Every game should end with a plainly stated score, constraint result, or evidence-backed explanation rather than an unexplained score alone.",
            "Deterministic scoring logic, a visible result breakdown, and fixed guardrail text for each challenge type.",
            "It makes the products replayable and trustworthy, especially where the optimizer or a model is involved.",
            "Avoid hidden win curves, unexplained grades, or claims of predictive certainty.",
            BLUE,
        ),
    ]
    add_catalog_table(doc, "1. Current foundation and established roadmap", "These are the existing fan-tool and collector concepts already defined in the project, plus the shared feature that should connect them.", core, new_page=False)

    build_now = [
        entry(
            "Lineup DNA",
            "Build now - Lineup Lab enhancement",
            "Explain why a chosen five works, which roles are covered, what is missing, and how a single substitution changes the profile.",
            "Current fan analytics helpers, role definitions, rate views, optimizer output, and source caveats.",
            "It makes the existing Lineup Lab understandable to casual users and strengthens every future game that links into it.",
            "Do not label box-score proxies as verified movement shooting, switching, or individual matchup data.",
            BLUE,
        ),
        entry(
            "Rotation Rescue",
            "Build now - first independent game",
            "A historical team-season gives the user a coaching brief and constraints; they build a five or full rotation and compare it with the exact optimizer result.",
            "Current optimizer, team-season player pool, role coverage, deterministic challenge definitions, and local save/share state.",
            "It turns the optimizer into a replayable daily or weekly game with an auditable finish rather than a black-box win forecast.",
            "Score distance from a defined optimum or constraint completion; do not invent season-win probabilities.",
            BLUE,
        ),
        entry(
            "Scout's Call",
            "Build now",
            "Show an historical opponent profile, let a user choose key priorities and a counter-lineup, then reveal the source-labeled game-plan recommendation.",
            "Current opponent game-plan module, historical team-season stats, and Lineup Lab priority controls.",
            "It gives fans a coaching-style decision loop and sends them directly into a meaningful Lineup Lab scenario.",
            "No player-to-player defensive assignments or live injury/schedule claims.",
            BLUE,
        ),
        entry(
            "Five-Role Draft",
            "Build now",
            "In a seeded draft, choose actual players to fill scarce roles such as creator, shooter, rebounder, connector, and defensive-activity proxy.",
            "Role definitions, historical player-season data, exact lineup feasibility, and deterministic draft pools.",
            "It borrows decision tension from build games while producing a real team profile instead of a fictional player simulation.",
            "Use roles and trade-offs, not a copied spin-wheel or 'steal a legend's skills' mechanic.",
            BLUE,
        ),
        entry(
            "Statline Sleuth",
            "Build now - NBA first",
            "A timed quiz asks the visitor to identify a player, team, season, or era from a reviewed stat clue.",
            "Approved public-safe NBA player-season data and a curated, versioned question bank.",
            "It is the fastest low-friction discovery loop and can later expand across sports without changing the core format.",
            "Do not generate claims from incomplete or ambiguous records; only publish reviewed questions.",
            BLUE,
        ),
        entry(
            "Evidence Court",
            "Build now",
            "A short analytics-literacy quiz asks whether a claim is a direct fact, reconstructed lineup result, model estimate, proxy, or unsupported.",
            "Curated scenarios based on documented evidence levels; no live private data required for the first version.",
            "It builds trust, differentiates the site from hot-take products, and teaches users why sample and source labels matter.",
            "Keep wording factual and explain the correct answer; do not reward overconfident causal conclusions.",
            BLUE,
        ),
        entry(
            "Optimizer Sensitivity Studio",
            "Build now",
            "Ask 'what must change for this answer to change?' by sweeping declared objective weights and user constraints, then showing stable alternatives versus fragile picks.",
            "Exact solver, stated objective weights, constraints, and historical player pool.",
            "It is a serious-fan differentiator: transparent model sensitivity instead of a single opaque lineup answer.",
            "Describe outputs as stability within the selected model and settings, never as player truth.",
            BLUE,
        ),
        entry(
            "Franchise Fingerprints",
            "Build now when team profiles are exposed",
            "An anonymized historical team-season card asks visitors to identify the franchise or era from style signals.",
            "Historical team-season totals and derived profile fields such as shooting mix, assists, turnovers, and rebounding.",
            "It broadens trivia beyond individual player stat lines and celebrates franchise history.",
            "Use transparent observed profiles; do not imply tactical labels unsupported by the source.",
            GOLD,
        ),
        entry(
            "Two Truths, One Box Score",
            "Build now for eligible mapped NBA seasons",
            "Show three exact claims about a player-season and ask users to identify the false one.",
            "Verified mapped player-season statistics and a fixed reviewed assertion bank.",
            "It is compact, card-adjacent, and easy to rotate as a daily challenge.",
            "Do not construct questions from unverified prose or uncertain identity matches.",
            BLUE,
        ),
        entry(
            "Phase Flip",
            "Build now where both phases exist",
            "Ask whether a selected player-season metric increased or decreased in the postseason, then show both samples and context.",
            "Verified regular-season and postseason rows, games, minutes, and selected rate metrics.",
            "It turns phase context into a quick learning game and makes the stat surface feel deeper.",
            "Omit the question when either phase is missing or the sample should not support a strong comparison.",
            BLUE,
        ),
        entry(
            "What Breaks This Five?",
            "Build now",
            "Present a high-profile lineup and ask the user to find its role weakness before the role model explains the gap.",
            "Current role coverage, objective metrics, and exact lineup feasibility.",
            "It makes role balance intuitive and creates an entertaining counterpart to the best-lineup optimizer.",
            "Call defensive labels proxies when the available inputs are box-score based.",
            BLUE,
        ),
        entry(
            "Role Evolution Reel",
            "Source-prep - public summary required",
            "A career timeline asks which season marked the greatest observable role shift in production, minutes, or playmaking.",
            "Multi-season player summaries with labeled source scope and reliable season denominators.",
            "It creates a compelling career-story layer without requiring a full fictional career simulation.",
            "Use observable role signals only; do not infer injuries, locker-room factors, or tracked movement.",
            GOLD,
        ),
        entry(
            "Era Translation Challenge",
            "Source-prep - full cohort required",
            "Compare players fairly within an era, position, and minimum-sample cohort; users predict who was more unusual in context.",
            "A labeled, complete league-season comparison cohort plus minutes/games thresholds and rate denominators.",
            "It replaces shallow cross-era arguments with transparent context and can power quizzes or a comparison tool.",
            "Never present a single team roster as a league-wide era baseline.",
            GOLD,
        ),
        entry(
            "Archetype Forge / Cohort Lab",
            "Source-prep - full cohort required",
            "Let a visitor set era-relative style sliders, then reveal the closest historical player-season or a small comparison cohort.",
            "A complete, labeled league-season cohort, position and era filters, rate denominators, and explicit similarity recipe.",
            "It creates a personalized discovery loop without pretending a single box-score comp is a scouting verdict.",
            "Show the exact recipe and comparison cohort; do not call it an objective player-equivalence ranking.",
            GOLD,
        ),
    ]
    add_catalog_table(doc, "2. Build-now games, quizzes, and analytics features", "These additions rely on current public-safe lineup/season information, deterministic rules, or a curated question bank. They are intentionally distinct from a generic analytics dashboard.", build_now)

    mapping = [
        entry(
            "Season Signature",
            "Mapping-dependent - highest collector priority",
            "Explain why an exact card season mattered within the athlete's own career, including regular/postseason context and meaningful selected markers.",
            "Verified NBA product mapping, depicted-season context, buyer-safe season rows, and selected advanced stats.",
            "It creates immediate card value and gives shoppers a reason to open the product modal beyond price and image.",
            "Only use exact mappings and actual depicted-season context; no title-based guesses.",
            GOLD,
        ),
        entry(
            "Card Clue Hunt",
            "Mapping-dependent",
            "Show a stat, era, team, or role clue and ask the user to identify a verified card/player match.",
            "Verified player-card mappings and a reviewed public clue bank.",
            "It converts player knowledge into a collector game and can live as a mode inside Player & Card Matchups.",
            "Never surface a product simply because its listing title resembles the answer.",
            GOLD,
        ),
        entry(
            "Card Timeline Tangle",
            "Mapping-dependent",
            "Arrange three to five verified mapped cards from one athlete in chronological or peak-to-valley order.",
            "Multiple exact mapped cards for an athlete plus reviewed season labels and selected stat context.",
            "It makes depth of collection and career history feel interactive instead of archival.",
            "Exclude any card without a fully resolved athlete and season context.",
            GOLD,
        ),
        entry(
            "Team-Year Reverse Lookup",
            "Mapping-dependent",
            "Choose a franchise and historical season, then see only cards whose exact athlete mapping and team-season context both resolve.",
            "Buyer-safe team-season projection and verified athlete-card-team-season relationships.",
            "It is a high-intent bridge from historical research into available catalog discovery.",
            "No inferred roster membership or title-search fallback.",
            GOLD,
        ),
        entry(
            "Multi-Subject Card Decoder",
            "Mapping-dependent",
            "Turn multi-player cards into a feature by showing ordered verified subjects and using a missing-subject microquiz.",
            "All ordered subjects verified through product mapping; subject order preserved.",
            "It makes a difficult catalog edge case distinctive rather than hiding it.",
            "If one subject is ambiguous, keep the item unresolved rather than showing a partial answer as complete.",
            GOLD,
        ),
        entry(
            "Stat-Profile Card Finder",
            "Mapping-dependent",
            "Filter exact mapped cards through transparent stat recipes such as scoring plus efficiency plus playmaking.",
            "Compact public-safe aggregate endpoint for requested mapped athletes and sport-specific stat groups.",
            "It improves discovery without implying that a statistical filter is a scouting verdict or price recommendation.",
            "Use sport-appropriate metrics; do not compare raw totals across NBA, MLB, and NFL.",
            GOLD,
        ),
        entry(
            "Franchise Passport",
            "Mapping-dependent / local-first",
            "A personal binder journey lets users save cards toward franchise or era milestones and earn factual progress badges.",
            "Verified mappings, local opt-in saved state, and clear milestone rules.",
            "It adds collection retention without asserting that a visitor owns any item or requiring an account on day one.",
            "Call the state saved or selected unless the user explicitly records ownership.",
            GOLD,
        ),
        entry(
            "Roster Blind Spot",
            "Mapping-dependent",
            "Identify a selected roster or Lineup Lab scenario's unfilled role, then surface eligible verified player-card matches.",
            "Role coverage, approved player-card mapping, and buyer-safe catalog projection.",
            "It joins the optimizer to card discovery in a useful, non-pushy way.",
            "It recommends role fit only; it never changes cart, inventory, pricing, or checkout state.",
            GOLD,
        ),
        entry(
            "Cross-Sport Stat Relay",
            "Mapping-dependent / sport-by-sport rollout",
            "A sport-appropriate stat clue leads to a verified NBA, MLB, or NFL player-card answer, with each sport retaining its own language and metrics.",
            "Reviewed sport-specific season rows, exact athlete-card mappings, public-safe clue definitions, and a separate adapter for each sport.",
            "It extends the quiz format into the wider catalog without flattening different sports into one fake universal ranking.",
            "Do not compare raw NBA, MLB, and NFL totals or reuse an NBA evidence label in another sport.",
            GOLD,
        ),
    ]
    add_catalog_table(doc, "3. Collector and card experiences", "These recommendations become safe only when the identity and product mapping path is exact. They are read-only discovery features, not inventory systems.", mapping)

    pbp = [
        entry(
            "Closing Five",
            "PBP validation-gated",
            "Let users choose who should close a historical situation, then compare against observed exact closing-lineup evidence and context.",
            "Validated compact derived exact-lineup outputs, starter/closer counts, possessions, context splits, and reliability labels.",
            "It makes advanced lineup evidence approachable through a concrete basketball decision.",
            "No true win probability, individual matchup assignment, or causal claim from a small sample.",
            RISK,
        ),
        entry(
            "Lineup Autopsy",
            "PBP validation-gated",
            "Present two historical five-man units, ask for a prediction in a defined context, then reveal observed results and projection context.",
            "Validated exact five-man rows, possessions, rating, reliability, and versioned projection data.",
            "It turns nuanced lineup research into a satisfying prediction-and-reveal format.",
            "Five players means an exact lineup; smaller player groups must never be displayed as one.",
            RISK,
        ),
        entry(
            "Pair Fit Lab",
            "PBP validation-gated",
            "Compare two player pairs and distinguish actual shared-floor evidence from a prospective fit hypothesis.",
            "Validated co-presence, on/off, WOWY cells, minutes/possessions, and reliability tiers.",
            "It gives fans a more rigorous way to debate complementary versus redundant players.",
            "Call results shared-floor evidence or fit hypothesis, not chemistry or causation.",
            RISK,
        ),
        entry(
            "Continuity Compass",
            "PBP validation-gated",
            "Visualize how stable a team's actual five-man usage was through concentration, starter/closer overlap, possessions, and sample size.",
            "Validated exact-lineup exposure, game use, possession share, and boundary-lineup evidence.",
            "It makes rotation stability a compelling historical story and teaches why some results are more trustworthy than others.",
            "Use only possession-observed boundary lineups; do not create dead-ball or missing-snapshot units.",
            RISK,
        ),
        entry(
            "Signal vs Noise",
            "PBP validation-gated",
            "Compare last-5, last-10, last-20, and season snapshots; users decide whether the difference merits a film question or is too thin to trust.",
            "Validated rolling splits, possession samples, contexts, and reliability flags.",
            "It teaches sound analytics judgment and prevents fans from overreacting to small samples.",
            "Reveal a recommendation to investigate, not a claim that a short window caused a change.",
            RISK,
        ),
        entry(
            "100-Possession Recipe",
            "PBP validation-gated",
            "Allocate tokens among threes, twos, free throws, turnovers, and offensive rebounds to recreate the closest observed team or unit profile.",
            "Validated direct event and four-factor aggregates with coverage status.",
            "It gamifies possession-level style and makes the four factors intuitive.",
            "Only use fully covered fields; do not turn absence of a fast-break qualifier into verified half-court offense.",
            RISK,
        ),
        entry(
            "Projection Ledger",
            "PBP validation-gated",
            "Show the transparent accounting for an unseen or observed exact five: summed RAPM baseline, shrunk residual if available, exposure, and reliability.",
            "Validated versioned RAPM, exact lineup projection, opponent/home context, and model provenance.",
            "It is a powerful advanced tool for serious users because it exposes the model rather than hiding it.",
            "Frame it as a hypothesis prioritizer, not a win forecast or causal result.",
            RISK,
        ),
        entry(
            "Stint Sleuth",
            "PBP validation-gated",
            "A daily puzzle asks users to reconstruct the order of a small set of observed lineup stints from a fully covered historical game.",
            "Validated game-level exact stint sequence and compact derived context only.",
            "It is an original way to turn roster history into a puzzle without pretending to recreate the game live.",
            "Ship no raw event payloads and avoid causal score-swing claims.",
            RISK,
        ),
        entry(
            "Possession Pathways",
            "PBP validation-gated",
            "Ask which unit converted better after a turnover, offensive rebound, or provider-confirmed fast break, then reveal the observed sample.",
            "Validated possession outcome aggregates and structured event-field coverage.",
            "It makes high-value possession contexts interactive and concrete.",
            "Treat unclassified or non-provider-fastbreak contexts honestly; never relabel them as half-court offense.",
            RISK,
        ),
        entry(
            "Impact Ledger",
            "PBP validation-gated",
            "A player case file compares role, exposure, direct on/off, teammate/opponent context, and offensive/defensive RAPM.",
            "Validated compact player profiles, on/off rows, context measures, and versioned RAPM outputs.",
            "It gives advanced users a more responsible alternative to one-number player rankings.",
            "End with experiment candidates and caveats, not an absolute 'best player' verdict.",
            RISK,
        ),
        entry(
            "Sample-Size Showdown",
            "PBP validation-gated",
            "Put two derived views with different exposure or possession counts side by side and ask which conclusion the evidence actually supports.",
            "Validated compact outputs, clear denominators, reliability tiers, and a reviewed answer explanation.",
            "It makes the site's evidence standard interactive and prevents a flashy small sample from being mistaken for a durable finding.",
            "Never use it to make a causal claim or hide a missing coverage warning.",
            RISK,
        ),
        entry(
            "Four-Factor Casefile",
            "PBP validation-gated",
            "Ask which observed factor best explains a stated scoring gap, then reveal the covered shooting, turnover, rebounding, and free-throw evidence.",
            "Validated four-factor aggregates, possession context, coverage labels, and fixed reviewed scenarios.",
            "It turns a core basketball-analytics framework into an approachable decision game.",
            "Describe the factor as observed context, not a complete causal explanation of a game result.",
            RISK,
        ),
        entry(
            "Leverage Map",
            "PBP validation-gated",
            "Map where a team or unit faced score-state, period, phase, or leverage-proxy pressure, then let the visitor explore the evidence.",
            "Validated context partitions, score-state definitions, possession samples, and a documented leverage proxy where used.",
            "It adds situational storytelling without inventing a proprietary win-probability product.",
            "Call every leverage measure a proxy unless a separately validated probability model is available.",
            RISK,
        ),
        entry(
            "Team Style Time Capsule",
            "PBP validation-gated / mapping-enhanced",
            "Place a historical team's covered style profile beside an eligible player-card season or team-year discovery result.",
            "Validated team aggregates, source-labeled context, and exact player-card-team-season mapping when tied to a card.",
            "It gives collectors a richer era and team story without turning a product modal into a raw analytics dashboard.",
            "Keep it a compact derived summary; no raw PBP, unverified tactics, or implied player causality.",
            RISK,
        ),
        entry(
            "Role-Change Replay",
            "PBP validation-gated / mapping-enhanced",
            "A reveal quiz asks the visitor to spot a player's observable starter, closer, minute, or production change across a documented span.",
            "Validated multi-season role outputs, sample labels, and exact card mapping only when shown beside a product.",
            "It creates a narrative continuation of Season Signature with stronger role evidence once the data gate is met.",
            "Use observable changes only; do not infer injuries, coaching intent, or private context.",
            RISK,
        ),
    ]
    add_catalog_table(doc, "4. Advanced research experiences - validated derived PBP only", "These are strong differentiators, but they must remain behind a successful validation and compact-derived-output gate. The raw archive and private tables are never browser assets.", pbp)

    services = [
        entry(
            "Exact Lineup Optimizer and Role Model",
            "Recommended core service/component",
            "The exact constrained engine behind best-five and full-rotation solutions, role coverage, alternatives, and sensitivity analysis.",
            "Historical player/team-season data, explicit constraints, and transparent objective definitions.",
            "It is the unique capability that makes DJHC experiences more credible than generic roster builders.",
            "Keep it as the feasibility source of truth and fail closed when proof cannot be established.",
            BLUE,
        ),
        entry(
            "Fan Analytics and Opponent Game-Plan Helpers",
            "Recommended browser-safe component",
            "Pure helpers for rate views, role explanations, comparison cohorts, and historical opponent priorities.",
            "Approved player/team-season summary inputs and explicitly labeled denominators or estimates.",
            "They make Lineup Lab explanations and light games possible without exposing private data.",
            "Hide per-100 rates when no valid possession denominator exists and keep proxy labels visible.",
            BLUE,
        ),
        entry(
            "Licensed Sportradar NBA v8",
            "Recommended private analytics source",
            "Licensed schedule, game-summary, and play-by-play source for reconstruction, lineup evidence, on/off, WOWY, and RAPM.",
            "Authorized license, private archive storage, validator, dedicated analytics target, and derived output pipeline.",
            "It is the appropriate source for reproducible advanced NBA PBP analytics, unlike descriptive-reference data.",
            "Credentials and raw data stay private; only compact validated derived results may ever support a product surface.",
            RISK,
        ),
        entry(
            "Dedicated Supabase Analytics Project",
            "Recommended isolated service",
            "A dedicated data target for analytics identities, derived outputs, migrations, and browser-safe read-only views.",
            "Separate project, explicit target checks, versioned migrations, and public-safe RPCs or projections where approved.",
            "It preserves a hard boundary between commerce data and sensitive sports analytics.",
            "Never import raw PBP, lineups, stints, or RAPM into the commerce Supabase project.",
            BLUE,
        ),
        entry(
            "Basketball-Reference",
            "Recommended descriptive reference only",
            "Historical player and season data source for descriptive profiles, positions, and public-summary support.",
            "Authorized descriptive imports/cache and clear source labeling.",
            "It is useful for player-season and career-style experiences that do not require play-by-play reconstruction.",
            "It is not an ingestion substitute for PBP, possessions, lineup reconstruction, WOWY, or RAPM.",
            GOLD,
        ),
        entry(
            "Baseball-Reference and Pro Football Reference",
            "Recommended private historical source for expansion",
            "Season-level MLB and NFL player tables for future cross-sport stat panels and verified identity work.",
            "Authorized local caches, stable provider identifiers, review-gated mapping, and sport-specific output adapters.",
            "They enable careful MLB/NFL expansion without pretending the NBA advanced stack transfers automatically.",
            "Keep source provenance private and do not compare raw cross-sport totals as equivalent.",
            GOLD,
        ),
        entry(
            "NFLverse and MLBAM/Chadwick",
            "Internal identity/media enrichment",
            "Potential supporting sources for stable identity or media coverage in expanded sports catalogs.",
            "Verified source policy, rights review, and the same cross-project identity gates as other sports data.",
            "They can improve matching and presentation where public-safe output is justified.",
            "They are not standalone public analytics experiences or a substitute for verified product mapping.",
            GOLD,
        ),
        entry(
            "FreeImage",
            "Adjacent media workflow only - not analytics",
            "An image rehosting or batching service that may support media operations independently of fan-tool logic.",
            "Separate media rights review, approved image handling policy, and an explicitly scoped visual workflow.",
            "It can help presentation work when needed, but it does not unlock a quiz, optimizer, mapping, or analytics feature.",
            "Keep it outside the analytics architecture and never treat media hosting as evidence for a sports-data claim.",
            GOLD,
        ),
        entry(
            "Slab-to-Stats Public Adapter",
            "Recommended buyer-safe service pattern",
            "A fixed-shape response that returns only requested active verified athlete data for eligible product surfaces.",
            "Verified product identity, active athlete status, rights-approved media, and aggregated season stats.",
            "It is the correct pattern for connecting sports data to product modals without exposing private tables.",
            "Withhold sold, hidden, deleted, disputed, unmapped, or unverified records.",
            BLUE,
        ),
    ]
    add_catalog_table(doc, "5. Recommended services and enabling components", "These are the data and technical services that support the catalog. They are not all public-facing products; several exist specifically to keep the public product safe and reproducible.", services)

    reviewed = [
        entry(
            "API-Sports API-NBA",
            "Considered prototype candidate - not cleared",
            "A potential rapid-prototype data provider considered during earlier exploration.",
            "Rights, coverage, quality, and terms must be reviewed before any product use.",
            "It could accelerate a limited prototype only if its legal and historical-data fit is confirmed.",
            "Do not treat it as approved or use it for the PBP/RAPM production path.",
            GOLD,
        ),
        entry(
            "NBA Stats API 2.0",
            "Early prototype source - not the production choice",
            "A source used in an earlier standalone prototype.",
            "Clear historical coverage and rights review would be required for continued product reliance.",
            "It may remain useful for narrow experimentation, but it does not replace a licensed reproducible PBP source.",
            "Do not present it as the authoritative historical analytics service.",
            GOLD,
        ),
        entry(
            "BALLDONTLIE",
            "Considered and not adopted",
            "A general sports-data API reviewed during source exploration.",
            "No new dependency unless future requirements and rights are explicitly reassessed.",
            "It is documented here to prevent accidental duplication of a rejected path.",
            "Not an approved analytics source for the current roadmap.",
            RISK,
        ),
        entry(
            "Cappers API",
            "Considered and not adopted",
            "A candidate source examined during earlier prototype research.",
            "No product dependency should be created without a fresh scoped evaluation.",
            "It remains a historical decision record, not a current recommendation.",
            "Not approved for implementation.",
            RISK,
        ),
        entry(
            "Synergy",
            "Optional, unadopted enrichment",
            "A licensed postgame play-type enrichment option considered as an addition, not a replacement for the core play-by-play source.",
            "Separate commercial entitlement, rights review, a specific product need, and provenance labeling.",
            "It could eventually add a carefully labeled play-type layer, but it is not selected for the current scope.",
            "Do not call it a chosen dependency; use it only after an explicit licensing and product decision.",
            GOLD,
        ),
    ]
    add_catalog_table(doc, "6. Considered, deferred, or rejected services", "These sources were discussed during the project but are not approved dependencies for the current product path. This distinction protects the roadmap from accidental source drift.", reviewed)

    doc.add_paragraph("7. Deliberate deferrals and capability gaps", style="Heading 1")
    add_body(doc, "Some inspiration should remain inspiration until the required data, systems, and governance exist. These are not recommended near-term builds.")
    for item in (
        "Full career simulator or commissioner mode: compelling long-term inspiration, but it requires its own progression, contracts, schedule, roster, multiplayer, and simulation systems. The current package should first power compact, truthful decision games.",
        "Real-time 1v1 or 3v3 multiplayer: defer until there is a proven reason to add authentication, anti-cheat, live state, moderation, and leaderboards.",
        "Verified defender matchups, screen/action labels, player/ball tracking, contest distance, expected shot quality, video clips, injuries, medical data, contracts, and true win probability: the current package does not establish these. They need separately licensed sources and a documented join strategy.",
        "Generic market-price, investment, or scarcity scoring: do not infer financial value from sports analytics or use fan tools to alter prices, inventory, sales state, checkout, or external commerce channels.",
    ):
        add_bullet(doc, item)

    doc.add_paragraph("8. Reference inspiration", style="Heading 1")
    add_body(doc, "The following sites shaped product-pattern thinking. They are inspiration, not implementation templates.")
    inspiration = [
        ("82-0", "Use the lesson of bounded choices, quick modes, trivia, challenges, and shareable results. Do not copy its opaque win-curve simulation. https://www.82-0.com/how-to-play"),
        ("Build-A-Bucket", "Use the lesson of scarce decisions, role paths, and a satisfying build-to-result loop. Keep DJHC distinct by drafting a real team profile or solving a constraint puzzle rather than copying a skill-stealing wheel. https://build-a-player.com/bucket/"),
        ("Hoopgoat / Goated", "Use the lesson of a saved personal story or run card. Defer full career and commissioner simulations until a much larger system and source set exists. https://goated.hoopgoat.com/"),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table)
    add_table_header(table.cell(0, 0), "Reference")
    add_table_header(table.cell(0, 1), "What to borrow - and what not to copy")
    for name, detail in inspiration:
        row_object = table.add_row()
        row = row_object.cells
        for i, cell in enumerate(row):
            set_cell_width(cell, (LABEL_DXA, DETAIL_DXA)[i])
            set_cell_margins(cell)
        p = row[0].paragraphs[0]
        set_paragraph(p, before=0, after=0, line=1.12)
        r = p.add_run(name)
        set_font(r, size=9.2, color=NAVY, bold=True)
        p = row[1].paragraphs[0]
        set_paragraph(p, before=0, after=0, line=1.12)
        r = p.add_run(detail)
        set_font(r, size=8.8, color=INK)
        prevent_row_split(row_object)

    doc.add_paragraph("Project evidence and interpretation notes", style="Heading 2")
    add_body(doc, "This catalog uses the current project registry, Lineup Lab architecture, player-card mapping rules, and private-analytics boundaries. 'Build now' means an experience can be developed from approved public-safe or local inputs; it does not claim that a feature is already live. 'PBP validation-gated' means the experience must wait for a current successful validation and a separately approved derived-output release path.", size=10.5, color=MUTED)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
