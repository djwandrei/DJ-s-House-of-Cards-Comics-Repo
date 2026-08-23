#!/usr/bin/env python3
"""Compare Excel-backed storefront listings with the current source workbooks."""

from __future__ import annotations

import json
import math
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from html import unescape
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PRODUCTS_PATH = ROOT / "products.json"
OUTPUT_DIR = ROOT / "outputs" / "nonlegacy-workbook-audit"

AUTHORITATIVE_WORKBOOK = Path(
    r"C:\Users\djwan\Downloads\Ebay Bulk Upload - 08-22-2026.xlsx"
)
SOURCE_SHEET = "Listings"
WORKBOOKS = [AUTHORITATIVE_WORKBOOK]

BULK_CANONICAL = AUTHORITATIVE_WORKBOOK.name
SOURCE_ALIASES = {
    "Ebay Bulk Upload (Final).xlsx": BULK_CANONICAL,
    "Ebay Bulk Upload (Final) - Photo Links Updated.xlsx": BULK_CANONICAL,
    "Ebay Bulk Upload with HTML.xlsx": BULK_CANONICAL,
    "Ebay Bulk Upload with HTML_league_fixed.xlsx": BULK_CANONICAL,
    BULK_CANONICAL: BULK_CANONICAL,
}

DETAIL_FIELDS = [
    "Title",
    "Start price",
    "Quantity",
    "Item photo URL",
    "HTML Full Link",
    "Condition ID",
    "CD:Professional Grader - (ID: 27501)",
    "CD:Grade - (ID: 27502)",
    "CDA:Certification Number - (ID: 27503)",
    "CD:Card Condition - (ID: 40001)",
    "C:Manufacturer",
    "C:Set",
    "C:Season",
    "C:Year Manufactured",
    "C:Player/Athlete",
    "C:Sport",
    "C:Features",
    "C:League",
    "C:Team",
    "C:Autographed",
    "C:Parallel/Variety",
]

ATTRIBUTE_ORDER = [
    "Rookie",
    "Autograph",
    "Serial Numbered",
    "One of One",
    "Short Print",
    "Memorabilia",
    "Parallel/Variety",
    "Insert",
    "Error",
]
ATTRIBUTE_ORDER_INDEX = {value: index for index, value in enumerate(ATTRIBUTE_ORDER)}
ATTRIBUTE_ALIASES = {
    "rookie": "Rookie",
    "rookies": "Rookie",
    "rc": "Rookie",
    "autograph": "Autograph",
    "autographed": "Autograph",
    "auto": "Autograph",
    "signed": "Autograph",
    "signature": "Autograph",
    "serial numbered": "Serial Numbered",
    "serial-numbered": "Serial Numbered",
    "numbered": "Serial Numbered",
    "one of one": "One of One",
    "1/1": "One of One",
    "short print": "Short Print",
    "short-print": "Short Print",
    "sp": "Short Print",
    "ssp": "Short Print",
    "memorabilia": "Memorabilia",
    "relic": "Memorabilia",
    "relics": "Memorabilia",
    "jersey": "Memorabilia",
    "patch": "Memorabilia",
    "parallel": "Parallel/Variety",
    "parallel/variety": "Parallel/Variety",
    "parallel / variety": "Parallel/Variety",
    "variation": "Parallel/Variety",
    "refractor": "Parallel/Variety",
    "prizm": "Parallel/Variety",
    "insert": "Insert",
    "error": "Error",
}


@dataclass
class WorkbookRow:
    source: str
    sheet: str
    row_number: int
    title: str
    key: str
    fields: dict[str, Any]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer():
            return str(int(value))
        return ("%.12g" % value).strip()
    return str(value).replace("\xa0", " ").strip()


def normalize_title(value: Any) -> str:
    text = clean_text(value).lower()
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = re.sub(r"[\u2010-\u2015]", "-", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_source(value: Any) -> str:
    raw = clean_text(value)
    return SOURCE_ALIASES.get(raw, raw)


def compare_value(value: Any) -> str:
    text = clean_text(value)
    text = unescape(text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", "\n", text)
    return text.strip()


def normalize_photo_urls(value: Any) -> list[str]:
    text = compare_value(value)
    if not text:
        return []
    parts = re.split(r"\s*\|\s*|\s*,\s*", text)
    return [part.strip() for part in parts if part.strip()]


def normalize_html(value: Any) -> str:
    text = compare_value(value)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_host_links(value: Any) -> list[str]:
    html = normalize_html(value)
    if not html:
        return []
    links = re.findall(r'href=["\']([^"\']+)["\']', html, flags=re.I)
    if not links and html.startswith(("http://", "https://")):
        links = [html]
    seen: set[str] = set()
    out: list[str] = []
    for link in links:
        link = link.strip()
        if link and link not in seen:
            seen.add(link)
            out.append(link)
    return out


def to_price(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and math.isfinite(value):
        return float(value)
    text = clean_text(value)
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def normalize_attribute(value: Any) -> str:
    raw = clean_text(value)
    if not raw:
        return ""
    if raw in ATTRIBUTE_ORDER_INDEX:
        return raw
    key = raw.lower()
    key = re.sub(r"\s*/\s*", "/", key)
    key = re.sub(r"\s+", " ", key).strip()
    return ATTRIBUTE_ALIASES.get(key, "")


def add_attribute(attributes: list[str], value: Any) -> None:
    attribute = normalize_attribute(value)
    if attribute and attribute not in attributes:
        attributes.append(attribute)


def split_features(value: Any) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    return [part.strip() for part in text.split("|") if part.strip()]


def has_serial_context(value: str) -> bool:
    return bool(
        re.search(
            r"\b(?:gold|silver|bronze|blue|red|green|purple|orange|pink|black|aqua|yellow|fuchsia|lime|cyan|white|emerald|sepia|tie-dye|rainbow|platinum|foil|border|parallel|refractor|prizm|holo|shimmer|wave|lava|pulsar|speckle|mojo|ice|glitter|chrome|optic|choice|cosmic|sapphire|x-?fractor|die[- ]cut|press proof|aspirations|status|mirror|prime|the finals|playoff ticket|premium stock|masterpieces|limited|numbered|serial|short print|sp|ssp|auto|autographs?|au|signatures?|sigs?|patch|relic|memorabilia|jersey|materials?|swatch|prospect)\b",
            value,
        )
    )


def has_serial_signal(title_text: str) -> bool:
    if re.search(
        r"\bserial[- ](?:ly[- ])?numbered\b|\bnumbered\s+(?:to|/)\s*\d+\b|\blimited\s+to\s+\d+\b|\bone\s+of\s+one\b|\b1\s*of\s*1\b|\b1/1\b",
        title_text,
    ):
        return True
    for segment in re.split(r"\s+\+\s+", title_text):
        for match in re.finditer(r"/\s*(\d{1,4})\b", segment):
            denominator = int(match.group(1))
            after = segment[match.end() : match.end() + 30]
            if re.match(r"\s*(?:cards?|pcs?|boxes?|packs?)\b", after):
                continue
            before = segment[max(0, match.start() - 90) : match.start()]
            if denominator >= 1900 and denominator <= 2035 and not has_serial_context(before):
                continue
            if not has_serial_context(before):
                continue
            return True
    return False


def derive_expected_attributes(fields: dict[str, Any]) -> list[str]:
    attributes: list[str] = []
    for feature in split_features(fields.get("C:Features")):
        add_attribute(attributes, feature)
    text = " ".join(
        clean_text(fields.get(name))
        for name in [
            "Title",
            "C:Features",
            "C:Parallel/Variety",
        ]
    ).lower()
    title_text = clean_text(fields.get("Title")).lower()
    autographed = clean_text(fields.get("C:Autographed")).lower()
    if re.search(r"\brookies?\b|\brc\b|\brookie related\b|\brated rookie\b|\bpre[- ]rookie\b", text):
        add_attribute(attributes, "Rookie")
    if autographed == "yes" or (
        not autographed
        and (
            re.search(
                r"\bauto(?:s|graph(?:ed|s)?|graphed|s)?\b|\bau\b|\bsigned\b|\bsignatures\b|\bsigs?\b|\bink\b|\bscript(?:s)?\b|\binscriptions?\b|\bpenmanship\b|\bsignature(?!\s+rookies)\b|\bpsa/dna certified authentic\b|\b(?:sticker|on-card|hard-signed)\s+auto\b",
                text,
            )
            or re.search(r"\bsignature\s+(?:series|shots|marks|materials|patch|jersey|memorabilia|autographs?)\b", text)
        )
    ):
        add_attribute(attributes, "Autograph")
    if has_serial_signal(title_text):
        add_attribute(attributes, "Serial Numbered")
    if re.search(r"\bone\s+of\s+one\b|\b1\s*of\s*1\b|\b1/1\b|\bprinting plate\b|\bpre[- ]production proof\b", text):
        add_attribute(attributes, "One of One")
    if re.search(r"\bshort[- ]print\b|\bssp\b", text):
        add_attribute(attributes, "Short Print")
    if (
        re.search(
            r"\bmemorabilia\b|\brelics?\b|\bjerseys?\b|\bjsy\b|\bpatch(?:es)?\b|\bswatches?\b|\bfabric\b|\bmaterials?\b|\bgame[- ](?:used|worn|bat)\b|\bpiece\s+of\s+the\s+game\b|\bplayer[- ]worn\b|\bclubhouse collection\b|\bby the letter\b",
            text,
        )
        or re.search(
            r"\b(?:black gold|throwback|rookie team|team|throwback)\s+threads\b|\bhot numbers game used\b|\bauthentic fabric\b|\bfabric of the future\b|\bsp game bat edition\b|\bbat kings\b|\bautograph-bat\b",
            text,
        )
    ):
        add_attribute(attributes, "Memorabilia")
    if (
        re.search(
            r"\bparallel\b|\bvariation\b|\bvariety\b|\brefractors?\b|\bfoils?\b|\bholo(?:foil)?\b|\bshimmer\b|\bwave\b|\blava\b|\bpulsar\b|\bspeckle\b|\bmojo\b|\bice\b|\bglitter\b|\bsapphire\b|\bx-?fractor\b|\bdie[- ]cut\b|\bpress proof\b",
            text,
        )
        or re.search(
            r"\b(?:gold|silver|bronze|blue|red|green|purple|orange|pink|black|aqua|yellow|fuchsia|lime|cyan|white|emerald|sepia)\s+(?:border|foil|parallel|refractor|prizm|holo|wave|shimmer|glitter|proof)\b",
            text,
        )
    ):
        add_attribute(attributes, "Parallel/Variety")
    if re.search(r"\binserts?\b|\bcase hit\b|\bvariation insert\b", text):
        add_attribute(attributes, "Insert")
    if re.search(r"\berrors?\b|\berr\b|\bwrong back\b|\bmisspell(?:ed|ing)\b|\bmisprint\b", text):
        add_attribute(attributes, "Error")
    return sorted(attributes, key=lambda value: (ATTRIBUTE_ORDER_INDEX.get(value, 999), value))


def active_product(product: dict[str, Any]) -> bool:
    return product.get("isDeleted") is not True


def product_title(product: dict[str, Any]) -> str:
    metadata = product.get("metadata") or {}
    fields = metadata.get("excelFields") or {}
    return clean_text(fields.get("Title") or product.get("name"))


def product_source(product: dict[str, Any], workbook_titles: dict[str, set[str]]) -> tuple[str, bool]:
    metadata = product.get("metadata") or {}
    raw = metadata.get("sourceWorkbook") or product.get("sourceWorkbook") or metadata.get("sourcePage") or product.get("sourcePage")
    source = normalize_source(raw)
    if source:
        return source, False
    key = normalize_title(product_title(product))
    possible = [source_name for source_name, titles in workbook_titles.items() if key in titles]
    if len(possible) == 1:
        return possible[0], True
    return "(unknown excel source)", False


def read_workbooks() -> list[WorkbookRow]:
    rows: list[WorkbookRow] = []
    for path in WORKBOOKS:
        if not path.exists():
            raise FileNotFoundError(path)
        wb = load_workbook(path, read_only=True, data_only=True)
        source = normalize_source(path.name)
        if SOURCE_SHEET not in wb.sheetnames:
            raise RuntimeError(f"{path} has no {SOURCE_SHEET!r} sheet")
        ws = wb[SOURCE_SHEET]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [clean_text(value) for value in header_row]
        for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            fields = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            title = clean_text(fields.get("Title"))
            if not title:
                continue
            rows.append(
                WorkbookRow(
                    source=source,
                    sheet=SOURCE_SHEET,
                    row_number=row_number,
                    title=title,
                    key=normalize_title(title),
                    fields=fields,
                )
            )
        wb.close()
    return rows


def read_products(workbook_titles: dict[str, set[str]]) -> list[dict[str, Any]]:
    products = json.loads(PRODUCTS_PATH.read_text(encoding="utf-8"))
    site_rows: list[dict[str, Any]] = []
    for product in products:
        metadata = product.get("metadata") or {}
        fields = metadata.get("excelFields")
        if not isinstance(fields, dict):
            continue
        source, inferred = product_source(product, workbook_titles)
        title = product_title(product)
        site_rows.append(
            {
                "id": product.get("id"),
                "source": source,
                "sourceInferred": inferred,
                "rawSource": clean_text(metadata.get("sourceWorkbook")),
                "title": title,
                "key": normalize_title(title),
                "active": active_product(product),
                "product": product,
                "metadata": metadata,
                "fields": fields,
            }
        )
    return site_rows


def summary_rows(summary: dict[str, Any]) -> list[dict[str, Any]]:
    return [{"metric": key, "value": value} for key, value in summary.items()]


def make_ref(row: WorkbookRow | dict[str, Any]) -> str:
    if isinstance(row, WorkbookRow):
        return f"{row.source} | {row.sheet}!{row.row_number}"
    return f"{row.get('source')} | product id {row.get('id')}"


def workbook_record(row: WorkbookRow, status: str, matched_site_ids: list[Any] | None = None) -> dict[str, Any]:
    return {
        "status": status,
        "workbook": row.source,
        "sheet": row.sheet,
        "row": row.row_number,
        "title": row.title,
        "startPrice": to_price(row.fields.get("Start price")),
        "team": clean_text(row.fields.get("C:Team")),
        "league": clean_text(row.fields.get("C:League")),
        "features": clean_text(row.fields.get("C:Features")),
        "itemPhotoUrl": compare_value(row.fields.get("Item photo URL")),
        "matchedSiteIds": ", ".join(str(value) for value in (matched_site_ids or [])),
    }


def site_record(row: dict[str, Any], status: str, workbook_refs: list[str] | None = None) -> dict[str, Any]:
    product = row["product"]
    fields = row["fields"]
    return {
        "status": status,
        "siteId": product.get("id"),
        "workbookSource": row.get("source"),
        "rawSource": row.get("rawSource"),
        "sourceInferred": bool(row.get("sourceInferred")),
        "active": bool(row.get("active")),
        "title": row.get("title"),
        "price": product.get("price"),
        "excelStartPrice": fields.get("Start price"),
        "team": product.get("team") or fields.get("C:Team"),
        "league": product.get("league") or fields.get("C:League"),
        "features": fields.get("C:Features"),
        "image": product.get("image"),
        "photoHostPageUrl": product.get("photoHostPageUrl"),
        "workbookRefs": ", ".join(workbook_refs or []),
    }


def is_placeholder_image(product: dict[str, Any]) -> bool:
    images = [product.get("image")]
    gallery = product.get("imageGallery")
    if isinstance(gallery, list):
        images.extend(gallery)
    images = [clean_text(image).lower() for image in images if clean_text(image)]
    if not images:
        return True
    return all("placeholder" in image for image in images)


def add_mismatch(mismatches: list[dict[str, Any]], severity: str, issue_type: str, site: dict[str, Any], workbook: WorkbookRow, field: str, site_value: Any, workbook_value: Any, action: str) -> None:
    product = site["product"]
    mismatches.append(
        {
            "severity": severity,
            "issueType": issue_type,
            "field": field,
            "siteId": product.get("id"),
            "workbook": workbook.source,
            "sheet": workbook.sheet,
            "row": workbook.row_number,
            "title": workbook.title,
            "siteValue": compare_value(site_value),
            "workbookValue": compare_value(workbook_value),
            "recommendedAction": action,
        }
    )


def compare_detail_fields(site: dict[str, Any], workbook: WorkbookRow, mismatches: list[dict[str, Any]]) -> None:
    product = site["product"]
    site_fields = site["fields"]
    for field in DETAIL_FIELDS:
        site_value = site_fields.get(field)
        workbook_value = workbook.fields.get(field)
        if field in {"Item photo URL"}:
            site_norm = normalize_photo_urls(site_value)
            wb_norm = normalize_photo_urls(workbook_value)
            if site_norm != wb_norm:
                add_mismatch(mismatches, "high", "metadata_field_mismatch", site, workbook, field, " | ".join(site_norm), " | ".join(wb_norm), "Refresh site metadata from workbook row.")
            continue
        if field == "HTML Full Link":
            site_norm = normalize_html(site_value)
            wb_norm = normalize_html(workbook_value)
        else:
            site_norm = compare_value(site_value)
            wb_norm = compare_value(workbook_value)
        if site_norm != wb_norm:
            severity = "high" if field in {"Title", "Start price", "Item photo URL", "HTML Full Link", "C:Features"} else "medium"
            add_mismatch(mismatches, severity, "metadata_field_mismatch", site, workbook, field, site_norm, wb_norm, "Refresh site metadata from workbook row.")

    site_price = to_price(product.get("price"))
    workbook_price = to_price(workbook.fields.get("Start price"))
    if site_price is not None and workbook_price is not None and abs(site_price - workbook_price) > 0.009:
        add_mismatch(mismatches, "high", "display_price_mismatch", site, workbook, "product.price", site_price, workbook_price, "Update storefront product price from Start price.")

    workbook_photo_urls = normalize_photo_urls(workbook.fields.get("Item photo URL"))
    if workbook_photo_urls and is_placeholder_image(product):
        add_mismatch(mismatches, "high", "placeholder_image_with_workbook_photos", site, workbook, "image", product.get("image"), " | ".join(workbook_photo_urls), "Rebuild image mapping from workbook photo URLs.")

    expected_host_links = extract_host_links(workbook.fields.get("HTML Full Link"))
    site_host = compare_value(product.get("photoHostPageUrl"))
    if expected_host_links and site_host and site_host not in expected_host_links:
        add_mismatch(mismatches, "medium", "photo_host_page_mismatch", site, workbook, "photoHostPageUrl", site_host, " | ".join(expected_host_links), "Refresh photo host page URL from workbook HTML link.")
    if expected_host_links and not site_host:
        add_mismatch(mismatches, "medium", "missing_photo_host_page", site, workbook, "photoHostPageUrl", site_host, " | ".join(expected_host_links), "Populate photo host page URL from workbook HTML link.")

    expected_attrs = derive_expected_attributes(workbook.fields)
    site_attrs = product.get("attributes")
    if not isinstance(site_attrs, list):
        site_attrs = []
    site_attrs = sorted(
        {normalize_attribute(value) for value in site_attrs if normalize_attribute(value)},
        key=lambda value: (ATTRIBUTE_ORDER_INDEX.get(value, 999), value),
    )
    missing_attrs = [value for value in expected_attrs if value not in site_attrs]
    unexpected_attrs = [value for value in site_attrs if value not in expected_attrs]
    if missing_attrs:
        add_mismatch(mismatches, "high", "missing_attribute_tags", site, workbook, "attributes", ", ".join(site_attrs), ", ".join(expected_attrs), "Regenerate product attributes from workbook features/title.")
    if unexpected_attrs:
        severity = "high" if "Serial Numbered" in unexpected_attrs else "medium"
        add_mismatch(mismatches, severity, "unexpected_attribute_tags", site, workbook, "attributes", ", ".join(site_attrs), ", ".join(expected_attrs), "Regenerate product attributes and remove stale unsupported tags.")


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook_rows = read_workbooks()
    workbook_titles: dict[str, set[str]] = defaultdict(set)
    for row in workbook_rows:
        workbook_titles[row.source].add(row.key)
    site_rows = read_products(workbook_titles)
    active_site_rows = [row for row in site_rows if row["active"]]

    workbook_by_key: dict[tuple[str, str], list[WorkbookRow]] = defaultdict(list)
    site_by_key: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    deleted_site_by_key: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    active_site_by_key: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in workbook_rows:
        workbook_by_key[(row.source, row.key)].append(row)
    for row in site_rows:
        site_by_key[(row["source"], row["key"])].append(row)
        if row["active"]:
            active_site_by_key[(row["source"], row["key"])].append(row)
        else:
            deleted_site_by_key[(row["source"], row["key"])].append(row)

    duplicates: list[dict[str, Any]] = []
    for key, rows in sorted(workbook_by_key.items()):
        if len(rows) > 1:
            duplicates.append(
                {
                    "location": "workbook",
                    "source": key[0],
                    "title": rows[0].title,
                    "count": len(rows),
                    "refs": "; ".join(make_ref(row) for row in rows),
                    "siteIds": "",
                }
            )
    for key, rows in sorted(active_site_by_key.items()):
        if len(rows) > 1:
            duplicates.append(
                {
                    "location": "active site",
                    "source": key[0],
                    "title": rows[0]["title"],
                    "count": len(rows),
                    "refs": "",
                    "siteIds": ", ".join(str(row["id"]) for row in rows),
                }
            )

    missing_from_active_site: list[dict[str, Any]] = []
    workbook_only_deleted: list[dict[str, Any]] = []
    for key, rows in sorted(workbook_by_key.items()):
        active_matches = active_site_by_key.get(key, [])
        if active_matches:
            continue
        deleted_matches = deleted_site_by_key.get(key, [])
        for row in rows:
            if deleted_matches:
                workbook_only_deleted.append(workbook_record(row, "Found only as deleted local product", [site["id"] for site in deleted_matches]))
            else:
                missing_from_active_site.append(workbook_record(row, "Missing from active site"))

    site_missing_from_workbooks: list[dict[str, Any]] = []
    for key, rows in sorted(active_site_by_key.items()):
        workbook_matches = workbook_by_key.get(key, [])
        if workbook_matches:
            continue
        for row in rows:
            site_missing_from_workbooks.append(site_record(row, "Active site row not found in any supplied workbook"))

    mismatches: list[dict[str, Any]] = []
    for key, site_matches in sorted(active_site_by_key.items()):
        workbook_matches = workbook_by_key.get(key, [])
        if len(site_matches) != 1 or len(workbook_matches) != 1:
            continue
        compare_detail_fields(site_matches[0], workbook_matches[0], mismatches)

    source_counts = []
    source_names = sorted(set([row.source for row in workbook_rows] + [row["source"] for row in site_rows]))
    for source in source_names:
        source_counts.append(
            {
                "source": source,
                "workbookRows": sum(1 for row in workbook_rows if row.source == source),
                "activeSiteRows": sum(1 for row in active_site_rows if row["source"] == source),
                "deletedSiteRows": sum(1 for row in site_rows if row["source"] == source and not row["active"]),
                "activeSiteNotInWorkbook": sum(1 for row in site_missing_from_workbooks if row["workbookSource"] == source),
                "workbookMissingFromActiveSite": sum(1 for row in missing_from_active_site if row["workbook"] == source),
                "workbookOnlyDeleted": sum(1 for row in workbook_only_deleted if row["workbook"] == source),
            }
        )

    mismatch_counter = Counter(row["issueType"] for row in mismatches)
    severity_counter = Counter(row["severity"] for row in mismatches)
    summary = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "catalogPath": str(PRODUCTS_PATH),
        "workbookRows": len(workbook_rows),
        "excelBackedSiteRows": len(site_rows),
        "activeExcelBackedSiteRows": len(active_site_rows),
        "deletedExcelBackedSiteRows": len(site_rows) - len(active_site_rows),
        "uniqueMatchedActiveRowsAudited": sum(1 for key in active_site_by_key if len(active_site_by_key[key]) == 1 and len(workbook_by_key.get(key, [])) == 1),
        "mismatchRows": len(mismatches),
        "highSeverityMismatchRows": severity_counter.get("high", 0),
        "mediumSeverityMismatchRows": severity_counter.get("medium", 0),
        "workbookRowsMissingFromActiveSite": len(missing_from_active_site),
        "workbookRowsFoundOnlyDeletedLocally": len(workbook_only_deleted),
        "activeSiteRowsNotFoundInWorkbooks": len(site_missing_from_workbooks),
        "duplicateGroups": len(duplicates),
        "metadataFieldMismatchRows": mismatch_counter.get("metadata_field_mismatch", 0),
        "displayPriceMismatchRows": mismatch_counter.get("display_price_mismatch", 0),
        "placeholderImageWithWorkbookPhotosRows": mismatch_counter.get("placeholder_image_with_workbook_photos", 0),
        "missingAttributeTagRows": mismatch_counter.get("missing_attribute_tags", 0),
        "unexpectedAttributeTagRows": mismatch_counter.get("unexpected_attribute_tags", 0),
    }

    payload = {
        "summary": summary_rows(summary),
        "sourceCounts": source_counts,
        "mismatches": mismatches,
        "workbookMissingFromActiveSite": missing_from_active_site,
        "workbookOnlyDeleted": workbook_only_deleted,
        "siteMissingFromWorkbooks": site_missing_from_workbooks,
        "duplicates": duplicates,
        "sampleMatchedRows": [
            {
                "siteId": site_matches[0]["id"],
                "workbook": key[0],
                "title": site_matches[0]["title"],
                "price": site_matches[0]["product"].get("price"),
                "sheet": workbook_by_key[key][0].sheet,
                "row": workbook_by_key[key][0].row_number,
            }
            for key, site_matches in sorted(active_site_by_key.items())
            if len(site_matches) == 1 and len(workbook_by_key.get(key, [])) == 1
        ][:250],
    }

    out_path = OUTPUT_DIR / "audit-data.json"
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(str(out_path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
