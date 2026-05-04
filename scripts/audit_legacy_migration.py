"""Compare legacy HTML inventory pages against the current products.json catalog.

The legacy site mostly lists a product description followed by one or more scan
images. This audit uses image basenames/alt text as the strongest migration key
and falls back to conservative fuzzy title matching when a legacy listing has no
scan or the image filename changed during cleanup.
"""

from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from html import unescape
from pathlib import Path
from typing import Iterable

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


WORKSPACE = Path(__file__).resolve().parents[1]
LEGACY_DIR = Path(r"C:\Users\djwan\Downloads\card_scans")
OUTPUT_DIR = WORKSPACE / "outputs" / "legacy-migration-audit"

LEGACY_FILES = [
    "baseball_cards_50s.htm",
    "baseball_cards_60s.htm",
    "baseball_cards_2000s.htm",
    "DJ Card Prices.htm",
    "football_cards.htm",
    "index.htm",
    "other_comics.htm",
    "marvel.htm",
    "comics.htm",
    "baseball_cards.htm",
    "baseball_cards_70s.htm",
    "baseball_cards_80s.htm",
    "baseball_cards_90s.htm",
    "baseball_cards_2010s.htm",
    "basketball_cards.htm",
    "cards.htm",
    "collectibles.htm",
]

PRODUCT_IMAGE_PATH_RE = re.compile(r"(baseball_card|basketball_cards|football_cards|comics/|collectibles/)", re.I)
NON_PRODUCT_IMAGE_RE = re.compile(r"(spacer|user_|facebook|logo|grass|bg|button|nav|header)", re.I)
PRICE_RE = re.compile(r"\$[\d,.]+(?:\s*[-–]\s*\$?[\d,.]+)?|contact for price", re.I)
YEAR_RE = re.compile(r"\b(19\d{2}|20\d{2})(?:[-–]\d{2,4})?\b")
CARD_NO_RE = re.compile(r"#\s*([A-Za-z0-9.-]+)")

STOP_WORDS = {
    "and",
    "the",
    "card",
    "cards",
    "rookie",
    "listed",
    "listing",
    "price",
    "pricing",
    "range",
    "good",
    "mint",
    "near",
    "condition",
    "contact",
    "please",
    "becket",
    "beckett",
    "baseball",
    "basketball",
    "football",
    "comics",
    "comic",
    "topps",
    "fleer",
    "upper",
    "deck",
    "bowman",
    "panini",
}


@dataclass
class LegacyRecord:
    source_file: str
    source_title: str
    category: str
    text: str
    image_src: str = ""
    image_alt: str = ""
    extraction_method: str = "text"

    @property
    def status(self) -> str:
        upper = self.text.upper()
        if "SOLD" in upper:
            return "SOLD"
        if "TRADED" in upper:
            return "TRADED"
        return "Available/Unknown"

    @property
    def price_text(self) -> str:
        matches = PRICE_RE.findall(self.text)
        return " | ".join(dict.fromkeys(clean_text(item) for item in matches))

    @property
    def image_file(self) -> str:
        return Path(self.image_src.replace("\\", "/")).name if self.image_src else ""

    @property
    def image_keys(self) -> set[str]:
        keys = set()
        for value in [self.image_file, self.image_alt]:
            key = normalize_key(Path(value).stem if "." in value else value)
            if key:
                keys.add(key)
        return keys


@dataclass
class ConsolidatedLegacyRecord:
    category: str
    text: str
    source_files: set[str] = field(default_factory=set)
    source_titles: set[str] = field(default_factory=set)
    statuses: set[str] = field(default_factory=set)
    price_texts: set[str] = field(default_factory=set)
    image_files: set[str] = field(default_factory=set)
    image_alts: set[str] = field(default_factory=set)
    image_keys: set[str] = field(default_factory=set)
    methods: set[str] = field(default_factory=set)


def clean_text(value: str | None) -> str:
    text = unescape(value or "").replace("\xa0", " ")
    text = text.replace("â€™", "'").replace("â€œ", '"').replace("â€", '"')
    text = text.replace("&#8217;", "'").replace("’", "'").replace("“", '"').replace("”", '"')
    return re.sub(r"\s+", " ", text).strip(" \t\r\n|")


def normalize_key(value: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean_text(value).lower())


def normalize_match_text(value: str | None) -> str:
    text = clean_text(value).lower()
    text = re.sub(r"\b(sold|traded)\s*!?", " ", text)
    text = re.sub(r"\b(nm[- ]?mt|nm|vg[- ]?ex|vg|ex[- ]?mt|ex|good|poor|pr|gem mint)\b", " ", text)
    text = re.sub(r"\bpsa\s*\d+(?:\.\d+)?\b", " ", text)
    text = PRICE_RE.sub(" ", text)
    text = re.sub(r"\b(name|card|beckett|price|range|condition|year|listed|pricing|contact|for|please)\b", " ", text)
    text = re.sub(r"[^a-z0-9#]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def significant_tokens(value: str | None) -> set[str]:
    return {
        token
        for token in normalize_match_text(value).replace("#", " #").split()
        if len(token) > 2 and token not in STOP_WORDS
    }


def extract_years(value: str | None) -> set[str]:
    return set(YEAR_RE.findall(clean_text(value)))


def extract_card_numbers(value: str | None) -> set[str]:
    return {normalize_key(item) for item in CARD_NO_RE.findall(clean_text(value)) if normalize_key(item)}


def guess_category(file_name: str, image_src: str = "", current_category: str = "") -> str:
    combined = f"{file_name} {image_src} {current_category}".lower().replace("\\", "/")
    if "basketball" in combined:
        return "Basketball"
    if "football" in combined:
        return "Football"
    if "collectibles" in combined:
        return "Collectibles"
    if "comic" in combined or "marvel" in combined:
        return "Comics"
    return "Baseball"


def is_product_image(src: str | None) -> bool:
    source = (src or "").replace("\\", "/")
    return bool(PRODUCT_IMAGE_PATH_RE.search(source)) and not NON_PRODUCT_IMAGE_RE.search(source)


def is_meaningful_paragraph(text: str) -> bool:
    lower = text.lower()
    if len(text) < 5:
        return False
    blocked = [
        "home |",
        "check us out",
        "name, card #",
        "name, comic #",
        "please contact me",
        "view marvel",
        "view other",
        "gd=good",
        "pricing:",
    ]
    if any(item in lower for item in blocked):
        return False
    if text in {"Sports Collectibles", "DC Comics", "Marvel Comics", "Pro Football", "Pro Basketball", "Pro Baseball"}:
        return False
    return True


def update_section_category(text: str, fallback: str) -> str:
    lower = text.lower()
    if "basketball" in lower:
        return "Basketball"
    if "football" in lower:
        return "Football"
    if "collectible" in lower:
        return "Collectibles"
    if "comic" in lower or "marvel" in lower:
        return "Comics"
    if "baseball" in lower:
        return "Baseball"
    return fallback


def split_text_lines(paragraph) -> list[str]:
    return [clean_text(line) for line in paragraph.get_text("\n", strip=True).splitlines() if clean_text(line)]


def looks_like_set_header(line: str) -> bool:
    if PRICE_RE.search(line) or "contact for price" in line.lower():
        return False
    if re.match(r"^(19\d{2}|20\d{2})(?:[-–]\d{2,4})?\s+", line):
        return True
    return bool(re.match(r"^[A-Z][A-Za-z0-9&'’ .-]{3,60}$", line)) and "#" not in line


def looks_like_product_line(line: str) -> bool:
    lower = line.lower()
    if not is_meaningful_paragraph(line):
        return False
    if "price:" in lower or "contact for price" in lower:
        return True
    if PRICE_RE.search(line) and (CARD_NO_RE.search(line) or len(line.split()) >= 3):
        return True
    if ("psa" in lower or "autograph" in lower or "game-used" in lower) and len(line.split()) >= 3:
        return True
    return False


def extract_legacy_records(path: Path) -> list[LegacyRecord]:
    raw = path.read_bytes()
    soup = BeautifulSoup(raw, "html.parser", from_encoding="iso-8859-1")
    source_title = clean_text(soup.title.get_text(" ", strip=True) if soup.title else path.stem)
    records: list[LegacyRecord] = []
    last_text = ""
    current_category = guess_category(path.name)
    current_set = ""

    for paragraph in soup.find_all("p"):
        text = clean_text(paragraph.get_text(" ", strip=True))
        current_category = update_section_category(text, current_category)
        images = [img for img in paragraph.find_all("img") if is_product_image(img.get("src"))]

        if images:
            for image in images:
                image_src = clean_text(image.get("src", ""))
                image_alt = clean_text(image.get("alt", "") or image.get("name", "") or image.get("id", ""))
                description = text if is_meaningful_paragraph(text) else last_text
                if not description:
                    description = image_alt or Path(image_src).stem
                records.append(
                    LegacyRecord(
                        source_file=path.name,
                        source_title=source_title,
                        category=guess_category(path.name, image_src, current_category),
                        text=description,
                        image_src=image_src,
                        image_alt=image_alt,
                        extraction_method="image",
                    )
                )

        lines = split_text_lines(paragraph)
        if lines:
            first = lines[0]
            if update_section_category(first, current_category) != current_category:
                current_category = update_section_category(first, current_category)
            if looks_like_set_header(first):
                current_set = first
                candidate_lines = lines[1:]
            else:
                candidate_lines = lines

            for line in candidate_lines:
                if looks_like_set_header(line):
                    current_set = line
                    continue
                if not looks_like_product_line(line):
                    continue
                title = line
                if current_set and not normalize_match_text(line).startswith(normalize_match_text(current_set)[:10]):
                    title = f"{current_set} {line}"
                records.append(
                    LegacyRecord(
                        source_file=path.name,
                        source_title=source_title,
                        category=current_category,
                        text=title,
                        extraction_method="text",
                    )
                )

        if is_meaningful_paragraph(text) and not images:
            last_text = text

    return records


def canonical_legacy_key(record: LegacyRecord) -> str:
    normalized_text = normalize_match_text(record.text)
    if normalized_text:
        return f"{record.category}|text|{normalized_text[:180]}"
    image_key = next(iter(sorted(record.image_keys)), "")
    return f"{record.category}|image|{image_key}"


def consolidate_records(records: Iterable[LegacyRecord]) -> list[ConsolidatedLegacyRecord]:
    grouped: dict[str, ConsolidatedLegacyRecord] = {}
    for record in records:
        key = canonical_legacy_key(record)
        if key not in grouped:
            grouped[key] = ConsolidatedLegacyRecord(category=record.category, text=record.text)
        item = grouped[key]
        if len(record.text) > len(item.text):
            item.text = record.text
        item.source_files.add(record.source_file)
        item.source_titles.add(record.source_title)
        item.statuses.add(record.status)
        if record.price_text:
            item.price_texts.add(record.price_text)
        if record.image_file:
            item.image_files.add(record.image_file)
        if record.image_alt:
            item.image_alts.add(record.image_alt)
        item.image_keys.update(record.image_keys)
        item.methods.add(record.extraction_method)

    return sorted(grouped.values(), key=lambda item: (item.category, sorted(item.source_files)[0], item.text.lower()))


def current_image_keys(product: dict) -> set[str]:
    values: list[str] = []
    for field in ["legacyImageLabel", "image"]:
        if product.get(field):
            values.append(str(product[field]))
    for field in ["imageGallery", "itemPhotoUrls", "htmlImageUrls"]:
        if isinstance(product.get(field), list):
            values.extend(str(item) for item in product[field])

    keys = set()
    for value in values:
        path = value.replace("\\", "/")
        basename = Path(path).name
        keys.add(normalize_key(Path(basename).stem))
        if field_value := normalize_key(value):
            keys.add(field_value)
    return {key for key in keys if key}


def current_match_text(product: dict) -> str:
    parts = [
        product.get("name", ""),
        product.get("legacyImageLabel", ""),
        product.get("sourcePage", ""),
        product.get("team", ""),
        product.get("playerAthlete", ""),
        product.get("condition", ""),
    ]
    metadata = product.get("metadata")
    if isinstance(metadata, dict):
        parts.extend(str(metadata.get(key, "")) for key in ["playerAthlete", "sourceWorkbook"])
    return " ".join(str(part) for part in parts if part)


def current_name_text(product: dict) -> str:
    """Return the clean product title used for the strongest title comparison."""
    return str(product.get("name", ""))


def overlap_coefficient(left: set[str], right: set[str]) -> float:
    """Score shared meaningful tokens without punishing a longer modern title."""
    denominator = min(len(left), len(right))
    return len(left & right) / denominator if denominator else 0.0


def score_title_match(legacy: ConsolidatedLegacyRecord, product: dict) -> float:
    old_text = getattr(legacy, "_audit_match_text", "") or normalize_match_text(legacy.text)
    new_name_text = product.get("_audit_name_match_text") or normalize_match_text(current_name_text(product))
    new_full_text = product.get("_audit_full_match_text") or normalize_match_text(current_match_text(product))
    if not old_text or not new_full_text:
        return 0.0

    old_tokens = getattr(legacy, "_audit_tokens", None) or significant_tokens(legacy.text)
    name_tokens = product.get("_audit_name_tokens") or significant_tokens(current_name_text(product))
    full_tokens = product.get("_audit_full_tokens") or significant_tokens(current_match_text(product))
    name_union = old_tokens | name_tokens
    full_union = old_tokens | full_tokens
    name_jaccard = len(old_tokens & name_tokens) / len(name_union) if name_union else 0.0
    full_jaccard = len(old_tokens & full_tokens) / len(full_union) if full_union else 0.0
    name_overlap = overlap_coefficient(old_tokens, name_tokens)
    full_overlap = overlap_coefficient(old_tokens, full_tokens)
    name_ratio = SequenceMatcher(None, old_text, new_name_text).ratio() if new_name_text else 0.0
    full_ratio = SequenceMatcher(None, old_text, new_full_text).ratio()

    years_old = getattr(legacy, "_audit_years", None) or extract_years(legacy.text)
    years_new = product.get("_audit_years") or extract_years(current_match_text(product))
    cards_old = getattr(legacy, "_audit_card_numbers", None) or extract_card_numbers(legacy.text)
    cards_new = product.get("_audit_card_numbers") or extract_card_numbers(current_match_text(product))

    # Old price-list rows are often terse, while current titles are normalized.
    # Combining title Jaccard, title overlap, and sequence similarity avoids
    # marking obvious year/card-number matches as missing.
    score = max(
        (0.42 * name_jaccard) + (0.32 * name_ratio) + (0.18 * name_overlap),
        (0.36 * full_jaccard) + (0.28 * full_ratio) + (0.18 * full_overlap),
    )
    if years_old and years_new and years_old & years_new:
        score += 0.08
    if cards_old and cards_new and cards_old & cards_new:
        score += 0.12
    if str(product.get("category", "")).lower() == legacy.category.lower():
        score += 0.03

    # Strong structured evidence beats weaker prose similarity. This catches
    # entries such as "1968 Topps #220" where the old site omitted the player.
    has_year_match = bool(years_old and years_new and years_old & years_new)
    has_card_match = bool(cards_old and cards_new and cards_old & cards_new)
    strong_title_overlap = max(name_overlap, full_overlap)
    if has_year_match and has_card_match and strong_title_overlap >= 0.50:
        score = max(score, 0.86)
    elif has_year_match and strong_title_overlap >= 0.72:
        score = max(score, 0.82)
    elif has_card_match and strong_title_overlap >= 0.72:
        score = max(score, 0.78)
    return min(score, 0.99)


def build_current_indexes(products: list[dict]) -> tuple[dict[str, list[dict]], list[dict]]:
    image_index: dict[str, list[dict]] = {}
    for product in products:
        product["_audit_image_keys"] = current_image_keys(product)
        product["_audit_match_text"] = current_match_text(product)
        product["_audit_name_match_text"] = normalize_match_text(current_name_text(product))
        product["_audit_full_match_text"] = normalize_match_text(product["_audit_match_text"])
        product["_audit_name_tokens"] = significant_tokens(current_name_text(product))
        product["_audit_full_tokens"] = significant_tokens(product["_audit_match_text"])
        product["_audit_years"] = extract_years(product["_audit_match_text"])
        product["_audit_card_numbers"] = extract_card_numbers(product["_audit_match_text"])
        product["_audit_category_lower"] = str(product.get("category", "")).lower()
        for key in product["_audit_image_keys"]:
            image_index.setdefault(key, []).append(product)
    return image_index, products


def prepare_legacy_record(legacy: ConsolidatedLegacyRecord) -> None:
    """Cache expensive matching tokens on the consolidated legacy record."""
    legacy._audit_match_text = normalize_match_text(legacy.text)  # type: ignore[attr-defined]
    legacy._audit_tokens = significant_tokens(legacy.text)  # type: ignore[attr-defined]
    legacy._audit_years = extract_years(legacy.text)  # type: ignore[attr-defined]
    legacy._audit_card_numbers = extract_card_numbers(legacy.text)  # type: ignore[attr-defined]


def best_match(legacy: ConsolidatedLegacyRecord, image_index: dict[str, list[dict]], products: list[dict]) -> dict:
    prepare_legacy_record(legacy)
    for key in sorted(legacy.image_keys):
        if key in image_index:
            product = image_index[key][0]
            return {
                "match_status": "Matched",
                "match_method": f"exact image/label key: {key}",
                "score": 1.0,
                "product": product,
            }

    category_products = [p for p in products if p.get("_audit_category_lower") == legacy.category.lower()]
    fallback_candidates = category_products or products
    candidates = narrow_candidates(legacy, fallback_candidates)
    best_product, best_score = find_best_title_match(legacy, candidates)
    if candidates is not fallback_candidates and best_score < 0.62:
        best_product, best_score = find_best_title_match(legacy, fallback_candidates)

    if best_score >= 0.78:
        status = "Matched"
    elif best_score >= 0.62:
        status = "Possible Match - Review"
    else:
        status = "Missing Candidate"

    return {
        "match_status": status,
        "match_method": "title fuzzy",
        "score": best_score,
        "product": best_product,
    }


def find_best_title_match(legacy: ConsolidatedLegacyRecord, candidates: list[dict]) -> tuple[dict | None, float]:
    shortlisted = shortlist_candidates(legacy, candidates)
    if not shortlisted:
        return None, 0.0

    best_product = None
    best_score = 0.0
    for product in shortlisted:
        score = score_title_match(legacy, product)
        if score > best_score:
            best_score = score
            best_product = product
    return best_product, best_score


def quick_candidate_signal(legacy: ConsolidatedLegacyRecord, product: dict) -> float:
    """Cheap pre-score used to avoid expensive fuzzy checks on every product."""
    old_tokens = getattr(legacy, "_audit_tokens", set())
    name_tokens = product.get("_audit_name_tokens", set())
    full_tokens = product.get("_audit_full_tokens", set())
    years_old = getattr(legacy, "_audit_years", set())
    cards_old = getattr(legacy, "_audit_card_numbers", set())
    years_new = product.get("_audit_years", set())
    cards_new = product.get("_audit_card_numbers", set())

    name_overlap = overlap_coefficient(old_tokens, name_tokens)
    full_overlap = overlap_coefficient(old_tokens, full_tokens)
    year_bonus = 0.28 if years_old and years_new and years_old & years_new else 0.0
    card_bonus = 0.34 if cards_old and cards_new and cards_old & cards_new else 0.0
    return max(name_overlap, full_overlap) + year_bonus + card_bonus


def shortlist_candidates(legacy: ConsolidatedLegacyRecord, candidates: list[dict]) -> list[dict]:
    scored = [(quick_candidate_signal(legacy, product), product) for product in candidates]
    positive = [(score, product) for score, product in scored if score > 0]
    if not positive:
        return []
    positive.sort(key=lambda item: item[0], reverse=True)
    # Keeping a modest shortlist retains close title matches while making reruns
    # practical on a large, image-heavy catalog.
    return [product for _, product in positive[:120]]


def narrow_candidates(legacy: ConsolidatedLegacyRecord, products: list[dict]) -> list[dict]:
    """Prefer products sharing structured keys before falling back to category."""
    legacy_years = getattr(legacy, "_audit_years", set())
    legacy_cards = getattr(legacy, "_audit_card_numbers", set())
    if legacy_years and legacy_cards:
        year_card = [
            product
            for product in products
            if legacy_years & product.get("_audit_years", set())
            and legacy_cards & product.get("_audit_card_numbers", set())
        ]
        if year_card:
            return year_card
    if legacy_years:
        year_matches = [product for product in products if legacy_years & product.get("_audit_years", set())]
        if year_matches:
            return year_matches
    if legacy_cards:
        card_matches = [product for product in products if legacy_cards & product.get("_audit_card_numbers", set())]
        if card_matches:
            return card_matches
    return products


def row_for_report(legacy: ConsolidatedLegacyRecord, match: dict) -> dict:
    product = match.get("product") or {}
    return {
        "Match Status": match["match_status"],
        "Migration Bucket": migration_bucket(legacy, match),
        "Match Method": match["match_method"],
        "Score": round(float(match["score"]), 3),
        "Legacy Category": legacy.category,
        "Legacy Text": legacy.text,
        "Legacy Status": " | ".join(sorted(legacy.statuses)),
        "Legacy Price Text": " | ".join(sorted(legacy.price_texts)),
        "Legacy Image Files": " | ".join(sorted(legacy.image_files)),
        "Legacy Image Alts": " | ".join(sorted(legacy.image_alts)),
        "Source Files": " | ".join(sorted(legacy.source_files)),
        "Source Page Titles": " | ".join(sorted(legacy.source_titles)),
        "Extraction Method": " | ".join(sorted(legacy.methods)),
        "Best Current ID": product.get("id", ""),
        "Best Current Name": product.get("name", ""),
        "Best Current Category": product.get("category", ""),
        "Best Current Price": product.get("priceLabel") or product.get("price", ""),
        "Best Current Image": product.get("image", ""),
        "Best Current Source Page": product.get("sourcePage", ""),
        "Review Note": review_note(legacy, match),
    }


def migration_bucket(legacy: ConsolidatedLegacyRecord, match: dict) -> str:
    if match["match_status"] == "Matched":
        return "Matched to current site"
    if match["match_status"].startswith("Possible"):
        return "Possible match needing manual review"
    if "SOLD" in legacy.statuses or "TRADED" in legacy.statuses:
        return "Missing, but old site marked sold/traded"
    if "image" in legacy.methods:
        return "High-confidence missing scanned listing"
    return "Missing text-only old price-list entry"


def review_note(legacy: ConsolidatedLegacyRecord, match: dict) -> str:
    if match["match_status"] == "Missing Candidate" and ("SOLD" in legacy.statuses or "TRADED" in legacy.statuses):
        return "Missing from current catalog; old status was SOLD/TRADED, so absence may be intentional."
    if match["match_status"] == "Missing Candidate":
        return "No exact image key or strong title match found in current products.json."
    if match["match_status"].startswith("Possible"):
        return "Possible title match only; review before treating as migrated."
    return "Matched by image/key or strong title score."


def write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        return
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path: Path, all_rows: list[dict]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    sections = [
        ("Summary", []),
        (
            "Missing Scanned Listings",
            [row for row in all_rows if row["Migration Bucket"] == "High-confidence missing scanned listing"],
        ),
        (
            "Missing Text Only",
            [row for row in all_rows if row["Migration Bucket"] == "Missing text-only old price-list entry"],
        ),
        (
            "Missing Sold Traded",
            [row for row in all_rows if row["Migration Bucket"] == "Missing, but old site marked sold/traded"],
        ),
        ("Possible Matches", [row for row in all_rows if row["Match Status"] == "Possible Match - Review"]),
        ("Matched Legacy Products", [row for row in all_rows if row["Match Status"] == "Matched"]),
        ("All Extracted Legacy Products", all_rows),
    ]

    header_fill = PatternFill("solid", fgColor="1F5FFF")
    missing_fill = PatternFill("solid", fgColor="FDE9EC")
    review_fill = PatternFill("solid", fgColor="FFF3D2")
    matched_fill = PatternFill("solid", fgColor="E8F5ED")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    thin = Side(style="thin", color="D7DEEF")

    for sheet_name, rows in sections:
        ws = wb.create_sheet(sheet_name)
        ws.sheet_properties.tabColor = {
            "Summary": "0B1020",
            "Missing Scanned Listings": "EF1823",
            "Missing Text Only": "B65300",
            "Missing Sold Traded": "7A869A",
            "Possible Matches": "F2C94C",
            "Matched Legacy Products": "1B8D3D",
            "All Extracted Legacy Products": "1F5FFF",
        }.get(sheet_name, "1F5FFF")

        if sheet_name == "Summary":
            missing = sum(1 for row in all_rows if row["Match Status"] == "Missing Candidate")
            possible = sum(1 for row in all_rows if row["Match Status"] == "Possible Match - Review")
            matched = sum(1 for row in all_rows if row["Match Status"] == "Matched")
            sold_missing = sum(1 for row in all_rows if row["Match Status"] == "Missing Candidate" and ("SOLD" in row["Legacy Status"] or "TRADED" in row["Legacy Status"]))
            available_missing = missing - sold_missing
            scanned_missing = sum(1 for row in all_rows if row["Migration Bucket"] == "High-confidence missing scanned listing")
            text_missing = sum(1 for row in all_rows if row["Migration Bucket"] == "Missing text-only old price-list entry")
            summary_rows = [
                ("Audit created", datetime.now().strftime("%Y-%m-%d %I:%M %p")),
                ("Current catalog source", "products.json"),
                ("Legacy HTML folder", str(LEGACY_DIR)),
                ("Unique legacy products extracted", len(all_rows)),
                ("Matched legacy products", matched),
                ("Possible matches needing review", possible),
                ("Missing candidates", missing),
                ("High-confidence missing scanned listings", scanned_missing),
                ("Missing text-only old price-list entries", text_missing),
                ("Missing candidates marked SOLD/TRADED on old site", sold_missing),
                ("Missing candidates not marked sold/traded", available_missing),
                ("Match rule", "Exact old image/alt/legacy label first; structured year/card/title matching second."),
            ]
            ws.append(["Metric", "Value"])
            for metric, value in summary_rows:
                ws.append([metric, value])
            ws.column_dimensions["A"].width = 44
            ws.column_dimensions["B"].width = 110
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = white_font
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                row[0].font = bold_font
                row[1].alignment = Alignment(wrap_text=True, vertical="top")
            continue

        if not rows:
            ws.append(["No rows"])
            continue

        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header, "") for header in headers])

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = white_font
            cell.alignment = Alignment(wrap_text=True, horizontal="center")
            cell.border = Border(top=thin, right=thin, bottom=thin, left=thin)

        fill = missing_fill if "Missing" in sheet_name else review_fill if "Possible" in sheet_name else matched_fill
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.fill = fill
                cell.border = Border(top=thin, right=thin, bottom=thin, left=thin)
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        widths = {
            "A": 22,
            "B": 34,
            "C": 24,
            "D": 10,
            "E": 18,
            "F": 60,
            "G": 20,
            "H": 24,
            "I": 34,
            "J": 26,
            "K": 35,
            "L": 28,
            "M": 18,
            "N": 14,
            "O": 52,
            "P": 18,
            "Q": 18,
            "R": 34,
            "S": 24,
            "T": 48,
        }
        for column, width in widths.items():
            ws.column_dimensions[column].width = width
        ws.freeze_panes = "A2"
        table_ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=re.sub(r"[^A-Za-z0-9]", "", sheet_name)[:25], ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False, showColumnStripes=False)
        ws.add_table(table)

    wb.save(path)


def main() -> None:
    products = json.loads((WORKSPACE / "products.json").read_text(encoding="utf-8"))
    image_index, current_products = build_current_indexes(products)

    legacy_records: list[LegacyRecord] = []
    missing_files: list[str] = []
    for file_name in LEGACY_FILES:
        path = LEGACY_DIR / file_name
        if not path.exists():
            missing_files.append(file_name)
            continue
        legacy_records.extend(extract_legacy_records(path))

    consolidated = consolidate_records(legacy_records)
    rows = [row_for_report(record, best_match(record, image_index, current_products)) for record in consolidated]
    rows.sort(key=lambda row: (row["Match Status"], row["Legacy Category"], row["Source Files"], row["Legacy Text"]))

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    csv_path = OUTPUT_DIR / f"legacy-products-missing-audit-{timestamp}.csv"
    xlsx_path = OUTPUT_DIR / f"legacy-products-missing-audit-{timestamp}.xlsx"
    md_path = OUTPUT_DIR / f"legacy-products-missing-summary-{timestamp}.md"

    write_csv(csv_path, rows)
    write_workbook(xlsx_path, rows)

    missing_rows = [row for row in rows if row["Match Status"] == "Missing Candidate"]
    possible_rows = [row for row in rows if row["Match Status"] == "Possible Match - Review"]
    matched_rows = [row for row in rows if row["Match Status"] == "Matched"]
    available_missing = [row for row in missing_rows if "SOLD" not in row["Legacy Status"] and "TRADED" not in row["Legacy Status"]]
    scanned_missing = [row for row in rows if row["Migration Bucket"] == "High-confidence missing scanned listing"]
    text_missing = [row for row in rows if row["Migration Bucket"] == "Missing text-only old price-list entry"]
    sold_missing = [row for row in rows if row["Migration Bucket"] == "Missing, but old site marked sold/traded"]

    summary = [
        "# Legacy Migration Audit",
        "",
        f"- Created: {datetime.now():%Y-%m-%d %I:%M %p}",
        f"- Legacy folder: `{LEGACY_DIR}`",
        f"- Current catalog: `{WORKSPACE / 'products.json'}`",
        f"- Unique legacy product candidates extracted: {len(rows)}",
        f"- Matched to current catalog: {len(matched_rows)}",
        f"- Possible matches needing review: {len(possible_rows)}",
        f"- Missing candidates: {len(missing_rows)}",
        f"- High-confidence missing scanned listings: {len(scanned_missing)}",
        f"- Missing text-only old price-list entries: {len(text_missing)}",
        f"- Missing candidates not marked SOLD/TRADED: {len(available_missing)}",
        f"- Missing candidates marked SOLD/TRADED: {len(sold_missing)}",
    ]
    if missing_files:
        summary.append(f"- Missing legacy source files: {', '.join(missing_files)}")
    summary.extend(["", "## Top Missing Candidates Not Marked Sold/Traded", ""])
    for row in available_missing[:25]:
        summary.append(f"- **{row['Legacy Category']}**: {row['Legacy Text']}  ")
        summary.append(f"  Source: `{row['Source Files']}` | Images: `{row['Legacy Image Files'] or 'none'}` | Best current guess: `{row['Best Current Name']}` ({row['Score']})")
    summary.extend(["", "## Output Files", "", f"- `{xlsx_path}`", f"- `{csv_path}`"])
    md_path.write_text("\n".join(summary), encoding="utf-8")

    print(json.dumps({
        "legacy_records_raw": len(legacy_records),
        "legacy_records_unique": len(rows),
        "matched": len(matched_rows),
        "possible_review": len(possible_rows),
        "missing": len(missing_rows),
        "high_confidence_missing_scanned": len(scanned_missing),
        "missing_text_only": len(text_missing),
        "missing_not_sold_or_traded": len(available_missing),
        "missing_sold_or_traded": len(sold_missing),
        "xlsx": str(xlsx_path),
        "csv": str(csv_path),
        "summary": str(md_path),
        "missing_source_files": missing_files,
    }, indent=2))


if __name__ == "__main__":
    main()
