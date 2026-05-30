from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


TARGET_HEADERS = [
    "Title",
    "Item photo URL",
    "HTML Full Link",
    "Description",
    "Condition ID",
    "CD:Professional Grader - (ID: 27501)",
    "CD:Grade - (ID: 27502)",
    "CD:Card Condition - (ID: 40001)",
    "C:Player/Athlete",
    "C:Sport",
    "C:Features",
    "C:League",
    "C:Team",
    "C:Autographed",
    "C:Set",
    "C:Season",
    "C:Year Manufactured",
]

PRODUCT_FILES = [
    "products-baseball.json",
    "products-basketball.json",
    "products-football.json",
    "products-collectibles.json",
]

GRADER_RE = re.compile(
    r"\b(PSA/DNA|PSA|BGS|BVG|BCCG|SGC|CGC|CSG|HGA|GMA|ISA|TAG|Beckett)\b",
    re.IGNORECASE,
)

AUTO_RE = re.compile(r"\b(auto(graph)?|signed|signature|on-card auto|sticker auto)\b", re.IGNORECASE)


def normalize_title(value: str | None) -> str:
    if not value:
        return ""
    normalized = (
        str(value)
        .replace("\ufeff", "")
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
    )
    return re.sub(r"\s+", " ", normalized).strip().lower()


def extract_description_field(description: str | None, label: str) -> str:
    if not description:
        return ""
    match = re.search(re.escape(label) + r":\s*([^;\r\n]+)", description, re.IGNORECASE)
    return re.sub(r"\s+", " ", match.group(1)).strip() if match else ""


def unique_text_list(value) -> list[str]:
    results: list[str] = []
    seen: set[str] = set()

    def append_text(text: str | None) -> None:
        if not text:
            return
        for part in re.split(r"\s*\|\s*", str(text)):
            clean = re.sub(r"\s+", " ", part).strip()
            if clean and clean not in seen:
                seen.add(clean)
                results.append(clean)

    if value is None:
        return results
    if isinstance(value, str):
        append_text(value)
        return results
    if isinstance(value, list):
        for entry in value:
            if isinstance(entry, dict):
                for field_value in entry.values():
                    if isinstance(field_value, str):
                        append_text(field_value)
            else:
                append_text(str(entry))
        return results
    if isinstance(value, dict):
        for field_value in value.values():
            if isinstance(field_value, str):
                append_text(field_value)
        return results
    append_text(str(value))
    return results


def join_unique_text_list(*values) -> str:
    results: list[str] = []
    seen: set[str] = set()
    for value in values:
        for item in unique_text_list(value):
            if item not in seen:
                seen.add(item)
                results.append(item)
    return " | ".join(results)


def preferred_description(product: dict, excel_fields: dict, existing_row: dict) -> str:
    description = str(product.get("description") or "")
    if description and not description.startswith("Imported from the legacy"):
        return description
    if excel_fields.get("Description"):
        return str(excel_fields["Description"])
    return str(existing_row.get("Description") or "")


def feature_value(product: dict, excel_fields: dict, description: str) -> str:
    detail = extract_description_field(description, "Features")
    if detail:
        detail = re.sub(r"\.\s*(Please review|Great for)\b.*$", "", detail, flags=re.IGNORECASE).strip()
        detail = detail.rstrip(".; ")
        return "|".join(unique_text_list(detail))
    existing = excel_fields.get("C:Features")
    if existing:
        return "|".join(unique_text_list(existing))
    return ""


def autograph_value(product: dict, excel_fields: dict) -> str:
    haystack = " ".join(
        [
            str(product.get("name") or ""),
            str(product.get("description") or ""),
            str(excel_fields.get("C:Features") or ""),
            str(excel_fields.get("C:Autographed") or ""),
        ]
    )
    return "Yes" if AUTO_RE.search(haystack) else "No"


def year_manufactured_value(product: dict, description: str, excel_fields: dict, existing_row: dict):
    year_value = product.get("year")
    if year_value not in (None, ""):
        try:
            return int(str(year_value))
        except ValueError:
            return str(year_value)

    detail_year = extract_description_field(description, "Year")
    if detail_year:
        try:
            return int(detail_year)
        except ValueError:
            return detail_year

    if excel_fields.get("C:Year Manufactured") not in (None, ""):
        return excel_fields.get("C:Year Manufactured")
    return str(existing_row.get("C:Year Manufactured") or "")


def grading_fields(product: dict, excel_fields: dict, existing_row: dict) -> dict[str, str]:
    condition = re.sub(r"\s+", " ", str(product.get("condition") or "")).strip()
    grading_match = GRADER_RE.search(condition)
    if grading_match:
        grader = grading_match.group(1)
        grader = "Beckett" if grader.upper() == "BECKETT" else grader.upper()
        grade = condition[grading_match.end() :].strip()
        return {
            "Condition ID": "2750-Graded",
            "CD:Professional Grader - (ID: 27501)": grader,
            "CD:Grade - (ID: 27502)": grade,
            "CD:Card Condition - (ID: 40001)": "",
        }

    card_condition = str(excel_fields.get("CD:Card Condition - (ID: 40001)") or "")
    if not card_condition:
        card_condition = str(existing_row.get("CD:Card Condition - (ID: 40001)") or "")
    if not card_condition and condition.lower() == "near mint or better":
        card_condition = "Near mint or better - (ID: 400010)"

    return {
        "Condition ID": "4000-Ungraded",
        "CD:Professional Grader - (ID: 27501)": "",
        "CD:Grade - (ID: 27502)": "",
        "CD:Card Condition - (ID: 40001)": card_condition,
    }


def product_update_map(product: dict, existing_row: dict) -> dict:
    metadata = product.get("metadata") or {}
    excel_fields = (metadata.get("excelFields") or {}).copy()

    description = preferred_description(product, excel_fields, existing_row)
    item_photo_url = join_unique_text_list(
        product.get("itemPhotoUrls"),
        product.get("itemPhotoUrl"),
        excel_fields.get("Item photo URL"),
        existing_row.get("Item photo URL"),
    )
    html_full_link = join_unique_text_list(
        product.get("htmlFullLink"),
        excel_fields.get("HTML Full Link"),
        existing_row.get("HTML Full Link"),
    )

    sport = str(product.get("sport") or product.get("category") or excel_fields.get("C:Sport") or "")
    team = str(product.get("team") or excel_fields.get("C:Team") or "")
    league = str(product.get("league") or excel_fields.get("C:League") or "")
    player = str(product.get("playerAthlete") or excel_fields.get("C:Player/Athlete") or "")
    set_value = extract_description_field(description, "Set") or str(excel_fields.get("C:Set") or "")
    season_value = extract_description_field(description, "Season") or str(excel_fields.get("C:Season") or "")

    updates = {
        "Title": str(product.get("name") or ""),
        "Item photo URL": item_photo_url,
        "HTML Full Link": html_full_link,
        "Description": description,
        "C:Player/Athlete": player,
        "C:Sport": sport,
        "C:Features": feature_value(product, excel_fields, description),
        "C:League": league,
        "C:Team": team,
        "C:Autographed": autograph_value(product, excel_fields),
        "C:Set": set_value,
        "C:Season": season_value,
        "C:Year Manufactured": year_manufactured_value(product, description, excel_fields, existing_row),
    }
    updates.update(grading_fields(product, excel_fields, existing_row))
    return updates


def update_signature(update_map: dict) -> str:
    return ";".join(f"{key}={update_map[key]}" for key in sorted(update_map.keys()))


def load_products(root: Path) -> list[dict]:
    all_products: list[dict] = []
    for product_file in PRODUCT_FILES:
        all_products.extend(json.loads((root / product_file).read_text(encoding="utf-8")))
    return all_products


def build_product_map(products: list[dict]) -> dict[str, list[dict]]:
    product_map: dict[str, list[dict]] = defaultdict(list)
    seen_per_key: dict[str, set[int]] = defaultdict(set)

    for product in products:
        excel_fields = ((product.get("metadata") or {}).get("excelFields") or {})
        candidate_titles = [product.get("name"), excel_fields.get("Title")]
        for title in candidate_titles:
            key = normalize_title(title)
            if not key:
                continue
            product_id = int(product.get("id"))
            if product_id in seen_per_key[key]:
                continue
            seen_per_key[key].add(product_id)
            product_map[key].append(product)

    return product_map


def backup_workbook(workbook_path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_path = workbook_path.with_name(f"{workbook_path.stem}.backup-{timestamp}{workbook_path.suffix}")
    shutil.copy2(workbook_path, backup_path)
    return backup_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook-path", required=True)
    parser.add_argument("--root", default=str(Path(__file__).resolve().parent.parent))
    parser.add_argument("--output-path", default="")
    parser.add_argument("--no-backup", action="store_true")
    args = parser.parse_args()

    workbook_path = Path(args.workbook_path).resolve()
    root = Path(args.root).resolve()
    output_path = Path(args.output_path).resolve() if args.output_path else workbook_path

    backup_path = None if args.no_backup else backup_workbook(workbook_path)
    products = load_products(root)
    product_map = build_product_map(products)

    workbook = load_workbook(workbook_path)
    worksheet = workbook.worksheets[0]

    header_map: dict[str, int] = {}
    for column_index in range(1, worksheet.max_column + 1):
        header = worksheet.cell(row=1, column=column_index).value
        if header:
            header_map[str(header)] = column_index

    rows_matched = 0
    rows_updated = 0
    rows_no_match = 0
    rows_ambiguous = 0
    cells_updated = 0
    touched_headers: dict[str, int] = defaultdict(int)

    for row_index in range(2, worksheet.max_row + 1):
        title = str(worksheet.cell(row=row_index, column=header_map["Title"]).value or "")
        if not title.strip():
            continue

        key = normalize_title(title)
        candidates = product_map.get(key)
        if not candidates:
            rows_no_match += 1
            continue

        existing_row = {}
        for header in TARGET_HEADERS:
            column_index = header_map.get(header)
            if column_index:
                existing_row[header] = worksheet.cell(row=row_index, column=column_index).value

        candidate_updates = [product_update_map(candidate, existing_row) for candidate in candidates]
        signatures = {update_signature(candidate_update) for candidate_update in candidate_updates}
        if len(signatures) > 1:
            rows_ambiguous += 1
            continue

        update_map = candidate_updates[0]
        row_changed = False

        for header in TARGET_HEADERS:
            column_index = header_map.get(header)
            if not column_index or header not in update_map:
                continue

            old_value = worksheet.cell(row=row_index, column=column_index).value
            new_value = update_map[header]
            if old_value is None:
                old_text = ""
            else:
                old_text = str(old_value)
            new_text = "" if new_value is None else str(new_value)

            if old_text == new_text:
                continue

            worksheet.cell(row=row_index, column=column_index, value=copy.copy(new_value))
            cells_updated += 1
            row_changed = True
            touched_headers[header] += 1

        rows_matched += 1
        if row_changed:
            rows_updated += 1

    workbook.save(output_path)

    print("Workbook update complete.")
    print(f"Workbook: {output_path}")
    if backup_path:
        print(f"Backup: {backup_path}")
    print(f"Rows matched: {rows_matched}")
    print(f"Rows updated: {rows_updated}")
    print(f"Rows skipped (no match): {rows_no_match}")
    print(f"Rows skipped (ambiguous): {rows_ambiguous}")
    print(f"Cells updated: {cells_updated}")
    for header in sorted(touched_headers.keys()):
        print(f"  {header}: {touched_headers[header]}")


if __name__ == "__main__":
    main()
