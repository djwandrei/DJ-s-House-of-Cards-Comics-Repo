"""Build an exhaustive product-to-image match audit workbook.

This intentionally scores every storefront product against every source image
in the live asset folders. The workbook does not write every raw pair because
Excel cannot hold the ~16 million possible rows. Instead, it records the scan
totals and keeps the highest-scoring candidates for each product, which gives a
usable review file while still making the matching process exhaustive.
"""

from __future__ import annotations

import heapq
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs" / "placeholder-image-audit-20260501"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_WORKBOOK = OUTPUT_DIR / "placeholder-image-full-crossmatch-2026-05-01.xlsx"
PROMOTIONS_REPORT = OUTPUT_DIR / "gallery-primary-promotions-2026-05-01.json"
SAFE_PROMOTIONS_REPORT = OUTPUT_DIR / "safe-full-crossmatch-image-promotions-2026-05-01.json"

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".svg"}
EXCLUDED_IMAGE_PARTS = {"thumbnails", "fonts"}
PRODUCT_IMAGE_KEYS = ("image", "imageGallery", "itemPhotoUrl", "itemPhotoUrls", "htmlImageUrls")
TOP_MATCHES_PER_PRODUCT = 12
TOP_MATCHES_PER_PLACEHOLDER = 20

STOPWORDS = {
    "a",
    "an",
    "and",
    "auto",
    "autos",
    "autograph",
    "autographs",
    "baseball",
    "basketball",
    "better",
    "bowman",
    "card",
    "cards",
    "chrome",
    "club",
    "comic",
    "comics",
    "collectible",
    "collectibles",
    "deck",
    "donruss",
    "fleer",
    "football",
    "for",
    "from",
    "image",
    "jersey",
    "leaf",
    "lot",
    "main",
    "memorabilia",
    "mint",
    "near",
    "number",
    "numbered",
    "of",
    "panini",
    "parallel",
    "patch",
    "photo",
    "plus",
    "prizm",
    "rc",
    "refractor",
    "relic",
    "rookie",
    "score",
    "select",
    "serial",
    "set",
    "sets",
    "signed",
    "signature",
    "skybox",
    "sp",
    "sports",
    "stadium",
    "the",
    "topps",
    "trading",
    "ultra",
    "upper",
    "with",
}

CATEGORY_HINTS = (
    ("baseball-cards", "Baseball"),
    ("baseball", "Baseball"),
    ("basketball-cards", "Basketball"),
    ("basketball", "Basketball"),
    ("football-cards", "Football"),
    ("football", "Football"),
    ("comics", "Comics"),
    ("comic", "Comics"),
    ("collectibles", "Collectibles"),
    ("collectible", "Collectibles"),
)


def normalize_path(value: str | None) -> str:
    """Normalize local asset references so product and inventory paths compare."""

    if not value:
        return ""
    cleaned = unquote(str(value).strip()).replace("\\", "/")
    if cleaned.startswith("./"):
        cleaned = cleaned[2:]
    return cleaned.lower()


def as_list(value):
    if isinstance(value, list):
        return value
    return [value] if value else []


def is_placeholder(value: object) -> bool:
    return isinstance(value, str) and "placeholder" in value.lower()


def is_external_url(value: object) -> bool:
    return isinstance(value, str) and urlparse(value).scheme in {"http", "https"}


def split_camel_and_digits(text: object) -> str:
    value = str(text or "")
    value = re.sub(r"([a-z])([A-Z])", r"\1 \2", value)
    value = re.sub(r"([A-Za-z])(\d)", r"\1 \2", value)
    value = re.sub(r"(\d)([A-Za-z])", r"\1 \2", value)
    return value


def tokens(text: object) -> set[str]:
    expanded = split_camel_and_digits(str(text or "").replace("&", " and "))
    raw_tokens = re.findall(r"[a-z0-9]+", expanded.lower())
    return {token for token in raw_tokens if len(token) > 1 and token not in STOPWORDS}


def squashed(text: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def category_from_path(relative_path: str) -> str:
    lowered = relative_path.lower().replace("\\", "/")
    for hint, category in CATEGORY_HINTS:
        if hint in lowered:
            return category
    return ""


def product_player(product: dict) -> str:
    metadata = product.get("metadata") if isinstance(product.get("metadata"), dict) else {}
    return str(product.get("playerAthlete") or metadata.get("playerAthlete") or "")


def product_token_set(product: dict) -> set[str]:
    metadata = product.get("metadata") if isinstance(product.get("metadata"), dict) else {}
    parts = [
        product.get("name"),
        product.get("team"),
        product.get("sourcePage"),
        product.get("legacyImageLabel"),
        product_player(product),
        metadata.get("playerAthlete"),
        metadata.get("team"),
        metadata.get("set"),
        metadata.get("sport"),
    ]
    return tokens(" ".join(str(part or "") for part in parts))


def current_image_refs(product: dict) -> list[str]:
    refs: list[str] = []
    for key in PRODUCT_IMAGE_KEYS:
        for ref in as_list(product.get(key)):
            if isinstance(ref, str) and ref.strip():
                refs.append(ref.strip())
    return list(dict.fromkeys(refs))


def build_product_records(products: list[dict]) -> list[dict]:
    records = []
    for product in products:
        player = product_player(product)
        refs = current_image_refs(product)
        local_refs = [normalize_path(ref) for ref in refs if normalize_path(ref).startswith("assets/")]
        external_refs = [ref for ref in refs if is_external_url(ref)]
        records.append(
            {
                "product": product,
                "id": product.get("id"),
                "category": product.get("category") or "",
                "title": product.get("name") or "",
                "year": str(product.get("year") or ""),
                "team": product.get("team") or "",
                "condition": product.get("condition") or "",
                "price": product.get("priceLabel") or product.get("displayPrice") or product.get("price") or "",
                "player": player,
                "tokens": product_token_set(product),
                "playerTokens": tokens(player),
                "legacyLabel": squashed(product.get("legacyImageLabel")),
                "titleSquashed": squashed(product.get("name")),
                "isPlaceholderPrimary": is_placeholder(product.get("image")),
                "localImageRefs": set(local_refs),
                "externalImageRefs": external_refs,
                "primaryImage": product.get("image") or "",
            }
        )
    return records


def build_image_inventory(products: list[dict]) -> list[dict]:
    used_paths = Counter()
    for product in products:
        for ref in current_image_refs(product):
            normalized = normalize_path(ref)
            if normalized.startswith("assets/"):
                used_paths[normalized] += 1

    images = []
    for path in (ROOT / "assets").rglob("*"):
        if not path.is_file() or path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue
        relative_parts = {part.lower() for part in path.relative_to(ROOT).parts}
        if relative_parts & EXCLUDED_IMAGE_PARTS:
            continue
        # Placeholder SVGs are valid site assets, but they should never outrank
        # real card/comic photos in a replacement-candidate audit.
        if "placeholder" in path.name.lower():
            continue
        relative_path = path.relative_to(ROOT).as_posix()
        normalized = relative_path.lower()
        parent_name = path.parent.name
        grandparent_name = path.parent.parent.name if path.parent.parent != ROOT else ""
        alias_text = " ".join([path.stem, parent_name, grandparent_name, relative_path])
        images.append(
            {
                "relativePath": relative_path,
                "normalizedPath": normalized,
                "absolutePath": str(path),
                "filename": path.name,
                "folder": path.parent.relative_to(ROOT).as_posix(),
                "category": category_from_path(relative_path),
                "tokens": tokens(alias_text),
                "years": set(re.findall(r"\d{4}", relative_path)),
                "pathSquashed": squashed(relative_path),
                "usedReason": f"Product catalog image ({used_paths[normalized]} ref)"
                if normalized in used_paths
                else "",
                "catalogRefCount": used_paths[normalized],
                "bytes": path.stat().st_size,
            }
        )
    return images


def score_image_match(product_record: dict, image_record: dict) -> tuple[float, str]:
    """Score a single product/image pair using title, player, year, and path cues."""

    score = 0.0
    reasons: list[str] = []
    product_category = product_record["category"]
    image_category = image_record["category"]

    if image_record["normalizedPath"] in product_record["localImageRefs"]:
        score += 26
        reasons.append("already attached to this product")

    if image_category and product_category and image_category == product_category:
        score += 16
        reasons.append("category folder matches")
    elif image_category and product_category and product_category != "Collectibles":
        score -= 8

    if product_record["year"] and product_record["year"] in image_record["years"]:
        score += 12
        reasons.append("year appears in image path")

    product_tokens = product_record["tokens"]
    image_tokens = image_record["tokens"]
    overlap = product_tokens & image_tokens
    if product_tokens:
        overlap_ratio = len(overlap) / max(1, len(product_tokens))
        score += 46 * overlap_ratio
        if overlap:
            reasons.append("token overlap: " + ", ".join(sorted(overlap)[:10]))

    player_tokens = product_record["playerTokens"]
    if player_tokens:
        if player_tokens <= image_tokens:
            score += 34
            reasons.append("full player name appears in image path")
        elif player_tokens & image_tokens:
            score += 16
            reasons.append("partial player name appears in image path")

    legacy_label = product_record["legacyLabel"]
    if legacy_label and legacy_label in image_record["pathSquashed"]:
        score += 42
        reasons.append("legacy image label appears in image path")

    title_squashed = product_record["titleSquashed"]
    if title_squashed and title_squashed in image_record["pathSquashed"]:
        score += 35
        reasons.append("full title appears in image path")

    return round(score, 2), "; ".join(reasons)


def confidence(score: float) -> str:
    if score >= 78:
        return "High"
    if score >= 60:
        return "Medium"
    if score >= 42:
        return "Low"
    return "Very Low"


def score_every_image_against_every_product(product_records: list[dict], image_records: list[dict]):
    all_rows = []
    placeholder_rows = []
    mismatch_rows = []

    for product_index, product_record in enumerate(product_records, 1):
        top_heap: list[tuple[float, int, dict]] = []
        placeholder_heap: list[tuple[float, int, dict]] = []
        current_scores: list[float] = []

        for image_index, image_record in enumerate(image_records):
            score, reasons = score_image_match(product_record, image_record)
            is_current = image_record["normalizedPath"] in product_record["localImageRefs"]
            if is_current:
                current_scores.append(score)

            row = {
                "Product ID": product_record["id"],
                "Category": product_record["category"],
                "Title": product_record["title"],
                "Year": product_record["year"],
                "Team / Publisher": product_record["team"],
                "Condition": product_record["condition"],
                "Price": product_record["price"],
                "Primary Image": product_record["primaryImage"],
                "Is Placeholder Primary": "Yes" if product_record["isPlaceholderPrimary"] else "No",
                "Player / Athlete": product_record["player"],
                "Score": score,
                "Confidence": confidence(score),
                "Is Current Product Image": "Yes" if is_current else "No",
                "Candidate Relative Path": image_record["relativePath"],
                "Candidate Folder": image_record["folder"],
                "Detected Candidate Category": image_record["category"],
                "Catalog Ref Count": image_record["catalogRefCount"],
                "Already Used Reason": image_record["usedReason"],
                "Match Reasons": reasons,
                "Candidate Absolute Path": image_record["absolutePath"],
            }

            heap_item = (score, -image_index, row)
            if len(top_heap) < TOP_MATCHES_PER_PRODUCT:
                heapq.heappush(top_heap, heap_item)
            elif score > top_heap[0][0]:
                heapq.heapreplace(top_heap, heap_item)

            if product_record["isPlaceholderPrimary"]:
                if len(placeholder_heap) < TOP_MATCHES_PER_PLACEHOLDER:
                    heapq.heappush(placeholder_heap, heap_item)
                elif score > placeholder_heap[0][0]:
                    heapq.heapreplace(placeholder_heap, heap_item)

        ranked_rows = [item[2] for item in sorted(top_heap, key=lambda item: item[0], reverse=True)]
        for rank, row in enumerate(ranked_rows, 1):
            row["Rank"] = rank
            all_rows.append(row)

        if product_record["isPlaceholderPrimary"]:
            ranked_placeholder = [item[2] for item in sorted(placeholder_heap, key=lambda item: item[0], reverse=True)]
            for rank, row in enumerate(ranked_placeholder, 1):
                placeholder_copy = dict(row)
                placeholder_copy["Rank"] = rank
                placeholder_rows.append(placeholder_copy)

        best_score = ranked_rows[0]["Score"] if ranked_rows else 0
        best_path = ranked_rows[0]["Candidate Relative Path"] if ranked_rows else ""
        current_best = max(current_scores) if current_scores else None
        should_flag = (
            current_best is not None
            and best_score >= 60
            and best_score - current_best >= 24
            and best_path not in product_record["localImageRefs"]
        )
        if should_flag:
            mismatch_rows.append(
                {
                    "Product ID": product_record["id"],
                    "Category": product_record["category"],
                    "Title": product_record["title"],
                    "Primary Image": product_record["primaryImage"],
                    "Best Current Image Score": current_best,
                    "Best Candidate Score": best_score,
                    "Best Candidate Path": best_path,
                    "Recommended Action": "Review current image against higher-scoring candidate",
                }
            )

        if product_index % 250 == 0:
            print(f"Scored {product_index}/{len(product_records)} products...", flush=True)

    return all_rows, placeholder_rows, mismatch_rows


def write_table(sheet, headers: list[str], rows: list[dict]) -> None:
    header_fill = PatternFill("solid", fgColor="17213A")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in rows:
        sheet.append([row.get(header, "") for header in headers])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, header in enumerate(headers, 1):
        sample_values = [str(header)] + [str(row.get(header, "")) for row in rows[:150]]
        width = min(70, max(10, max(len(value) for value in sample_values) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    for sheet_row in sheet.iter_rows():
        for cell in sheet_row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border


def add_confidence_formatting(sheet) -> None:
    headers = [cell.value for cell in sheet[1]]
    if "Confidence" not in headers:
        return
    column_index = headers.index("Confidence") + 1
    fills = {
        "High": PatternFill("solid", fgColor="D8F3DC"),
        "Medium": PatternFill("solid", fgColor="DCE7FF"),
        "Low": PatternFill("solid", fgColor="FFF3CD"),
        "Very Low": PatternFill("solid", fgColor="F8D7DA"),
    }
    for row_index in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row_index, column=column_index)
        if cell.value in fills:
            cell.fill = fills[cell.value]


def write_workbook(
    products: list[dict],
    images: list[dict],
    all_rows: list[dict],
    placeholder_rows: list[dict],
    mismatch_rows: list[dict],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    placeholder_count = sum(1 for product in products if is_placeholder(product.get("image")))
    pair_count = len(products) * len(images)
    promotion_count = 0
    if PROMOTIONS_REPORT.exists():
        promotion_count = len(json.loads(PROMOTIONS_REPORT.read_text(encoding="utf-8")).get("promotedProducts", []))
    safe_promotion_count = 0
    if SAFE_PROMOTIONS_REPORT.exists():
        safe_promotion_count = len(json.loads(SAFE_PROMOTIONS_REPORT.read_text(encoding="utf-8")).get("applied", []))

    summary_rows = [
        ["Full Product/Image Crossmatch Audit", ""],
        ["Generated", datetime.now().isoformat(timespec="seconds")],
        ["Products scanned", len(products)],
        ["Source images scanned", len(images)],
        ["Raw pair comparisons scored", pair_count],
        ["Top candidates retained per product", TOP_MATCHES_PER_PRODUCT],
        ["Top candidates retained per placeholder", TOP_MATCHES_PER_PLACEHOLDER],
        ["Placeholder primary images remaining", placeholder_count],
        ["Safe gallery promotions already applied", promotion_count],
        ["Safe full-crossmatch promotions applied", safe_promotion_count],
        ["Potential current-photo mismatch flags", len(mismatch_rows)],
        ["", ""],
        ["Product counts by category", ""],
    ]
    for category, count in sorted(Counter(product.get("category") or "Unknown" for product in products).items()):
        summary_rows.append([category, count])
    summary_rows.extend(
        [
            ["", ""],
            [
                "Method",
                "Every source image was scored against every product. The workbook keeps top-ranked pairs so Excel stays usable while preserving exhaustive scan coverage.",
            ],
        ]
    )
    for row in summary_rows:
        summary.append(row)
    summary["A1"].fill = PatternFill("solid", fgColor="17213A")
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 90
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    match_headers = [
        "Product ID",
        "Category",
        "Title",
        "Year",
        "Team / Publisher",
        "Condition",
        "Price",
        "Primary Image",
        "Is Placeholder Primary",
        "Player / Athlete",
        "Rank",
        "Score",
        "Confidence",
        "Is Current Product Image",
        "Candidate Relative Path",
        "Candidate Folder",
        "Detected Candidate Category",
        "Catalog Ref Count",
        "Already Used Reason",
        "Match Reasons",
        "Candidate Absolute Path",
    ]
    all_sheet = workbook.create_sheet("All Product Top Matches")
    write_table(all_sheet, match_headers, all_rows)
    add_confidence_formatting(all_sheet)

    placeholder_sheet = workbook.create_sheet("Placeholder Top Matches")
    write_table(placeholder_sheet, match_headers, placeholder_rows)
    add_confidence_formatting(placeholder_sheet)

    mismatch_headers = [
        "Product ID",
        "Category",
        "Title",
        "Primary Image",
        "Best Current Image Score",
        "Best Candidate Score",
        "Best Candidate Path",
        "Recommended Action",
    ]
    write_table(workbook.create_sheet("Potential Photo Mismatches"), mismatch_headers, mismatch_rows)

    image_rows = [
        {
            "Relative Path": image["relativePath"],
            "Folder": image["folder"],
            "Filename": image["filename"],
            "Detected Category": image["category"],
            "Catalog Ref Count": image["catalogRefCount"],
            "Already Used Reason": image["usedReason"],
            "Bytes": image["bytes"],
            "Absolute Path": image["absolutePath"],
        }
        for image in images
    ]
    write_table(
        workbook.create_sheet("Image Inventory"),
        [
            "Relative Path",
            "Folder",
            "Filename",
            "Detected Category",
            "Catalog Ref Count",
            "Already Used Reason",
            "Bytes",
            "Absolute Path",
        ],
        image_rows,
    )

    workbook.save(OUTPUT_WORKBOOK)


def main() -> int:
    products = json.loads((ROOT / "products.json").read_text(encoding="utf-8"))
    product_records = build_product_records(products)
    image_records = build_image_inventory(products)
    print(
        json.dumps(
            {
                "products": len(product_records),
                "images": len(image_records),
                "pairComparisons": len(product_records) * len(image_records),
                "output": str(OUTPUT_WORKBOOK),
            },
            indent=2,
        ),
        flush=True,
    )
    all_rows, placeholder_rows, mismatch_rows = score_every_image_against_every_product(product_records, image_records)
    write_workbook(products, image_records, all_rows, placeholder_rows, mismatch_rows)
    print(
        json.dumps(
            {
                "workbook": str(OUTPUT_WORKBOOK),
                "allProductTopMatchRows": len(all_rows),
                "placeholderTopMatchRows": len(placeholder_rows),
                "potentialPhotoMismatches": len(mismatch_rows),
            },
            indent=2,
        ),
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
