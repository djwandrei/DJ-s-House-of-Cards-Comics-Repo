#!/usr/bin/env python3
"""Apply Beckett Legacy workbook corrections to the storefront product data.

The workbook is the human-reviewed source of truth for legacy Beckett matches.
This script keeps the update repeatable: it backs up the workbook, adds/fills a
"Beckett Real Time Pricing" column, updates selected product IDs, and rebuilds
the static JSON/JS product bundles used as storefront fallbacks.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import shutil
import sys
import time
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = Path.home() / "Documents" / "eBay Docs" / "Listing Automation" / "Beckett Legacy Pricing Review.xlsx"
CACHE_PATH = Path.home() / "Documents" / "eBay Docs" / "Listing Automation" / "beckett_legacy_pricing_cache.json"
BACKUP_ROOT = Path.home() / "Documents" / "eBay Docs" / "Listing Automation" / "backups"
REPORT_PATH = ROOT / "outputs" / "beckett-legacy-update-report.json"
REAL_TIME_HEADERS = ("Beckett Live Pricing", "Beckett Real Time Pricing")

PRODUCT_FILES = {
    "products.json": lambda items: items,
    "products-baseball.json": lambda items: [item for item in items if item.get("category") == "Baseball"],
    "products-basketball.json": lambda items: [item for item in items if item.get("category") == "Basketball"],
    "products-football.json": lambda items: [item for item in items if item.get("category") == "Football"],
    "products-comics.json": lambda items: [item for item in items if item.get("category") == "Comics"],
    "products-collectibles.json": lambda items: [item for item in items if item.get("category") == "Collectibles"],
    "products-sports.json": lambda items: [
        item for item in items if item.get("category") in {"Baseball", "Basketball", "Football"}
    ],
    "products-featured.json": lambda items: [item for item in items if item.get("isFeatured")],
}

DATA_BUNDLE_FILES = {
    "products.json": "products-data-full.js",
    "products-baseball.json": "products-data-baseball.js",
    "products-basketball.json": "products-data-basketball.js",
    "products-football.json": "products-data-football.js",
    "products-comics.json": "products-data-comics.js",
    "products-collectibles.json": "products-data-collectibles.js",
    "products-sports.json": "products-data-sports.js",
    "products-featured.json": "products-data-featured.js",
}

PRICING_SHEETS = [
    "Legacy Beckett Pricing",
    "Review Needed",
    "No Beckett Match",
    "Action Queue",
    "Pricing Review",
]

MONEY_RE = re.compile(r"\$?\s*(?:\d[\d,]*(?:\.\d+)?|\.\d+)")
GRADE_RE = re.compile(r"\b(PSA|BGS|SGC|CGC|HGA|BCCG|GAI)\s*(10|[1-9](?:\.\d)?)\b", re.I)
REMOVE_LISTING_RE = re.compile(r"^\s*(?:remove(?:\s+listing)?|removed|delete(?:\s+listing)?)\s*$", re.I)
REAL_TIME_RE = re.compile(
    r"Beckett\s+Real\s+Time\s+Pricing\s*(\$?\s*\d[\d,]*(?:\.\d+)?)\s*to\s*(\$?\s*\d[\d,]*(?:\.\d+)?)",
    re.I,
)


def load_cache() -> dict[str, Any]:
    if CACHE_PATH.exists():
        return json.loads(CACHE_PATH.read_text(encoding="utf-8"))
    return {"searches": {}, "cards": {}, "completed": {}}


def save_cache(cache: dict[str, Any]) -> None:
    CACHE_PATH.write_text(json.dumps(cache, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def money_to_float(value: Any) -> float | None:
    text = str(value or "").replace(",", "").replace("$", "").strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def product_id_from_value(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def is_remove_listing(value: Any) -> bool:
    return bool(REMOVE_LISTING_RE.match(str(value or "")))


def row_requests_removal(row: dict[str, Any]) -> bool:
    """Treat any match-oriented Remove Listing value as a storefront removal."""
    for header, value in row.items():
        if "match" in str(header).lower() and is_remove_listing(value):
            return True
    return False


def remove_listing_rows(wb) -> tuple[list[int], list[dict[str, Any]]]:
    """Delete workbook rows explicitly marked Remove Listing.

    Product IDs are returned so the matching storefront records can also be
    removed and the backend can be soft-deleted during the sync step.
    """
    removed_ids: set[int] = set()
    removed_rows: list[dict[str, Any]] = []

    for ws in wb.worksheets:
        headers = get_headers(ws)
        product_col = headers.get("Product ID")
        match_columns = [column for header, column in headers.items() if "match" in header.lower()]
        if not product_col or not match_columns:
            continue

        rows_to_delete: list[int] = []
        for row_index in range(2, ws.max_row + 1):
            row = row_to_dict(ws, row_index)
            if not row_requests_removal(row):
                continue
            product_id = product_id_from_value(ws.cell(row_index, product_col).value)
            rows_to_delete.append(row_index)
            removed_rows.append(
                {
                    "sheet": ws.title,
                    "row": row_index,
                    "productId": product_id,
                    "title": str(row.get("Title") or row.get("Beckett Matched Title") or "").strip(),
                }
            )
            if product_id is not None:
                removed_ids.add(product_id)

        for row_index in reversed(rows_to_delete):
            ws.delete_rows(row_index, 1)

    return sorted(removed_ids), removed_rows


def extract_money_values(value: Any) -> list[float]:
    values: list[float] = []
    for match in MONEY_RE.findall(str(value or "")):
        number = money_to_float(match)
        if number is not None:
            values.append(number)
    return values


def format_money(value: float) -> str:
    return f"${value:,.2f}"


def normalize_price_label(value: Any) -> str:
    values = extract_money_values(value)
    if not values:
        return ""
    if len(values) == 1:
        return format_money(values[0])
    low, high = min(values[:2]), max(values[:2])
    return f"{format_money(low)}-{format_money(high)}"


def normalize_realtime_label(value: Any) -> str:
    text = str(value or "").strip()
    if re.search(r"\bN/A\b", text, flags=re.I):
        return "N/A"
    return normalize_price_label(text)


def has_realtime_price(value: Any) -> bool:
    return bool(normalize_price_label(value))


def midpoint_from_label(label: str) -> float | None:
    values = extract_money_values(label)
    if not values:
        return None
    if len(values) == 1:
        return round(values[0], 2)
    return round((min(values[:2]) + max(values[:2])) / 2, 2)


def grade_suffix(condition: Any) -> tuple[str, int | None]:
    match = GRADE_RE.search(str(condition or ""))
    if not match:
        return "", None
    company = match.group(1).upper()
    grade = float(match.group(2))
    label_grade = str(int(grade)) if grade.is_integer() else str(grade).rstrip("0").rstrip(".")
    return f"{company} {label_grade}", int(math.ceil(grade))


def is_no_match(status: Any) -> bool:
    normalized = str(status or "").strip().lower()
    return normalized.startswith("no ") or normalized.startswith("no-") or normalized.startswith("no_")


def has_confirmed_match(row: dict[str, Any]) -> bool:
    status = str(row.get("Match Status") or "").strip().lower()
    if not (
        status == "matched"
        or status == "now matched"
        or status.startswith("confirmed")
        or (status.startswith("match") and "no " not in status)
    ):
        return False
    return bool(row.get("Beckett URL") and row.get("Beckett Matched Title"))


def get_headers(ws) -> dict[str, int]:
    return {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}


def ensure_realtime_column(ws) -> int:
    headers = get_headers(ws)
    for header in REAL_TIME_HEADERS:
        if header in headers:
            return headers[header]

    target_col = ws.max_column + 1
    ws.cell(1, target_col).value = "Beckett Real Time Pricing"
    ws.cell(1, target_col).fill = PatternFill("solid", fgColor="1F2937")
    ws.cell(1, target_col).font = Font(color="FFFFFF", bold=True)
    return target_col


def realtime_value_from_row(row: dict[str, Any]) -> Any:
    for header in REAL_TIME_HEADERS:
        value = row.get(header)
        if str(value or "").strip():
            return value
    return row.get("Beckett Real Time Pricing")


def row_to_dict(ws, row_index: int) -> dict[str, Any]:
    headers = get_headers(ws)
    return {header: ws.cell(row_index, column).value for header, column in headers.items()}


def parse_realtime_pricing(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    rt_heading = soup.select_one(".rtHad")
    if rt_heading:
        text = rt_heading.parent.get_text(" ", strip=True)
        match = REAL_TIME_RE.search(text)
        if match:
            return normalize_price_label(f"{match.group(1)}-{match.group(2)}")
        if re.search(r"Beckett\s+Real\s+Time\s+Pricing\s+N/A\b", text, flags=re.I):
            return "N/A"
    text = soup.get_text(" ", strip=True)
    match = REAL_TIME_RE.search(text)
    if match:
        return normalize_price_label(f"{match.group(1)}-{match.group(2)}")
    if re.search(r"Beckett\s+Real\s+Time\s+Pricing\s+N/A\b", text, flags=re.I):
        return "N/A"
    return ""


def fetch_realtime_pricing(url: str, session: Any, cache: dict[str, Any], delay: float) -> str:
    card_cache = cache.setdefault("cards", {}).setdefault(url, {})
    cached = str(card_cache.get("real_time_pricing") or "").strip()
    if cached:
        return normalize_realtime_label(cached)

    time.sleep(max(0, delay))
    response = session.get(url, timeout=45)
    response.raise_for_status()
    label = parse_realtime_pricing(response.text)
    if label:
        card_cache["real_time_pricing"] = label
        card_cache["real_time_pricing_fetched_at"] = datetime.now().isoformat(timespec="seconds")
    return label


def cached_realtime_pricing(url: str, cache: dict[str, Any]) -> str:
    return normalize_realtime_label(cache.get("cards", {}).get(url, {}).get("real_time_pricing"))


def grade_price_from_cache(url: str, grade_number: int | None, cache: dict[str, Any]) -> str:
    if not url or grade_number is None:
        return ""
    graded_prices = cache.get("cards", {}).get(url, {}).get("graded_prices") or {}
    for basis, price in graded_prices.items():
        match = re.search(r"\((\d+(?:\.\d+)?)\)", str(basis))
        if not match:
            continue
        if int(math.ceil(float(match.group(1)))) == grade_number:
            return normalize_price_label(price)
    return ""


def build_product_name(row: dict[str, Any], current_name: str) -> str:
    title = str(row.get("Beckett Matched Title") or row.get("Title") or current_name or "").strip()
    suffix, _ = grade_suffix(row.get("Current Grade / Condition"))
    if suffix and suffix.lower() not in title.lower():
        title = f"{title} {suffix}"
    return title


def best_price_for_row(row: dict[str, Any], cache: dict[str, Any]) -> tuple[str, str]:
    url = str(row.get("Beckett URL") or "").strip()
    suffix, grade_number = grade_suffix(row.get("Current Grade / Condition"))
    if suffix:
        label = normalize_price_label(row.get("Beckett Matched Grade Price"))
        if not label:
            label = grade_price_from_cache(url, grade_number, cache)
        if not label:
            label = normalize_price_label(row.get("Beckett Graded Range"))
        return label, "graded"

    # For ungraded legacy rows, the workbook's Beckett Real Time Pricing column
    # is the source of truth. Preserve explicit "N/A" values instead of silently
    # falling back to older guide/original pricing on the storefront.
    label = normalize_realtime_label(realtime_value_from_row(row))
    if label == "N/A":
        return label, "real_time_na"
    return label, "real_time"


def update_product(product: dict[str, Any], row: dict[str, Any], cache: dict[str, Any]) -> dict[str, Any]:
    updated = deepcopy(product)
    if row.get("Current Grade / Condition"):
        updated["condition"] = str(row["Current Grade / Condition"]).strip()
    if row.get("Site Team / Publisher"):
        updated["team"] = str(row["Site Team / Publisher"]).strip()
    if row.get("Site Year"):
        try:
            updated["year"] = int(row["Site Year"])
        except (TypeError, ValueError):
            pass
    if row.get("Site Player / Athlete"):
        updated["playerAthlete"] = str(row["Site Player / Athlete"]).strip()

    # Keep the storefront anchored to the workbook even when a row remains unconfirmed.
    fallback_name = str(row.get("Title") or updated.get("name", "")).strip()
    fallback_price_label = normalize_price_label(row.get("Original Price / Range"))
    fallback_price_value = midpoint_from_label(fallback_price_label)
    if fallback_name:
        updated["name"] = fallback_name
    if fallback_price_label and fallback_price_value is not None:
        updated["priceLabel"] = fallback_price_label
        updated["price"] = fallback_price_value

    if has_confirmed_match(row):
        updated["name"] = build_product_name(row, updated.get("name", ""))
        price_label, price_basis = best_price_for_row(row, cache)
        price_value = midpoint_from_label(price_label)
        if price_basis == "real_time_na":
            updated["priceLabel"] = "N/A"
            updated["price"] = None
        elif price_label and price_value is not None:
            updated["priceLabel"] = price_label
            updated["price"] = price_value

    return updated


def write_json(path: Path, data: list[dict[str, Any]]) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")


def write_data_bundle(source_name: str, bundle_path: Path, data: list[dict[str, Any]]) -> None:
    serialized = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    bundle = (
        f'window.DJ_PRELOADED_SOURCE = "{source_name}";\n'
        f"window.DJ_PRELOADED_PRODUCTS = {serialized}\n"
        ";\n"
    )
    bundle_path.write_text(bundle, encoding="utf-8")


def rebuild_product_files(products: list[dict[str, Any]]) -> None:
    for source_name, selector in PRODUCT_FILES.items():
        subset = selector(products)
        write_json(ROOT / source_name, subset)
        write_data_bundle(source_name, ROOT / DATA_BUNDLE_FILES[source_name], subset)


def backup_workbook() -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-beckett-realtime")
    backup_dir = BACKUP_ROOT / stamp
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / WORKBOOK_PATH.name
    shutil.copy2(WORKBOOK_PATH, backup_path)
    return backup_path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--min-id", type=int, default=1)
    parser.add_argument("--max-id", type=int, default=180)
    parser.add_argument("--fetch-min-id", type=int, default=None)
    parser.add_argument("--fetch-max-id", type=int, default=None)
    parser.add_argument("--update-min-id", type=int, default=None)
    parser.add_argument("--update-max-id", type=int, default=None)
    parser.add_argument("--fetch-realtime", action="store_true")
    parser.add_argument("--clear-realtime", action="store_true")
    parser.add_argument("--delay", type=float, default=0.15)
    parser.add_argument("--beckett-email", default=os.environ.get("BECKETT_EMAIL", ""))
    parser.add_argument("--beckett-password", default=os.environ.get("BECKETT_PASSWORD", ""))
    args = parser.parse_args()
    fetch_min_id = args.fetch_min_id if args.fetch_min_id is not None else args.min_id
    fetch_max_id = args.fetch_max_id if args.fetch_max_id is not None else args.max_id
    update_min_id = args.update_min_id if args.update_min_id is not None else args.min_id
    update_max_id = args.update_max_id if args.update_max_id is not None else args.max_id

    cache = load_cache()
    backup_path = backup_workbook()
    wb = load_workbook(WORKBOOK_PATH)
    remove_listing_ids, removed_workbook_rows = remove_listing_rows(wb)

    session = None
    if args.fetch_realtime:
        sys.path.insert(0, str(ROOT / "scripts"))
        from beckett_legacy_pricing import create_session  # type: ignore

        if not args.beckett_email or not args.beckett_password:
            raise SystemExit("--beckett-email and --beckett-password are required with --fetch-realtime")
        session = create_session(args.beckett_email, args.beckett_password)

    realtime_by_url: dict[str, str] = {}
    workbook_rows_by_id: dict[int, dict[str, Any]] = {}
    realtime_fetch_errors: dict[str, str] = {}

    for sheet_name in PRICING_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        realtime_col = ensure_realtime_column(ws)
        headers = get_headers(ws)
        pid_col = headers.get("Product ID")
        url_col = headers.get("Beckett URL")
        raw_col = headers.get("Beckett Raw Range")
        price_col = headers.get("Beckett Price Range")
        if not pid_col:
            continue

        for row_index in range(2, ws.max_row + 1):
            row = row_to_dict(ws, row_index)
            url = str(row.get("Beckett URL") or "").strip() if url_col else ""
            product_id = product_id_from_value(ws.cell(row_index, pid_col).value)
            should_fetch_realtime = (
                session is not None
                and sheet_name == "Legacy Beckett Pricing"
                and product_id is not None
                and fetch_min_id <= product_id <= fetch_max_id
            )
            if args.clear_realtime:
                ws.cell(row_index, realtime_col).value = None
                row["Beckett Real Time Pricing"] = None
            existing_rt = str(realtime_value_from_row(row) or "").strip()
            realtime_label = normalize_realtime_label(existing_rt)

            if not realtime_label and url:
                realtime_label = cached_realtime_pricing(url, cache)

            if not realtime_label and url and should_fetch_realtime:
                if url in realtime_by_url:
                    realtime_label = realtime_by_url[url]
                else:
                    try:
                        realtime_label = fetch_realtime_pricing(url, session, cache, args.delay)
                        save_cache(cache)
                    except Exception as error:  # noqa: BLE001 - report per URL and continue.
                        realtime_fetch_errors[url] = str(error)
                        realtime_label = ""
                    realtime_by_url[url] = realtime_label

            if realtime_label:
                ws.cell(row_index, realtime_col).value = realtime_label
                row["Beckett Real Time Pricing"] = realtime_label

            if (
                product_id is not None
                and sheet_name == "Legacy Beckett Pricing"
                and product_id not in remove_listing_ids
                and update_min_id <= product_id <= update_max_id
            ):
                workbook_rows_by_id[product_id] = row

    products_path = ROOT / "products.json"
    products = json.loads(products_path.read_text(encoding="utf-8"))
    product_by_id = {int(product["id"]): product for product in products}

    removed_products: list[dict[str, Any]] = []
    for product_id in remove_listing_ids:
        removed = product_by_id.pop(product_id, None)
        if removed:
            removed_products.append(
                {
                    "id": product_id,
                    "name": removed.get("name"),
                    "category": removed.get("category"),
                    "priceLabel": removed.get("priceLabel"),
                }
            )

    changed_products: list[dict[str, Any]] = []
    skipped_ids: list[int] = []
    for product_id in range(update_min_id, update_max_id + 1):
        if product_id in remove_listing_ids:
            continue
        product = product_by_id.get(product_id)
        row = workbook_rows_by_id.get(product_id)
        if not product or not row:
            skipped_ids.append(product_id)
            continue
        updated = update_product(product, row, cache)
        if updated != product:
            product_by_id[product_id] = updated
            changed_products.append(
                {
                    "id": product_id,
                    "oldName": product.get("name"),
                    "newName": updated.get("name"),
                    "oldPriceLabel": product.get("priceLabel"),
                    "newPriceLabel": updated.get("priceLabel"),
                    "oldCondition": product.get("condition"),
                    "newCondition": updated.get("condition"),
                    "matchStatus": row.get("Match Status"),
                }
            )

    updated_products = [product_by_id[int(product["id"])] for product in products if int(product["id"]) in product_by_id]
    rebuild_product_files(updated_products)

    for ws in wb.worksheets:
        headers = get_headers(ws)
        realtime_col = next((headers[header] for header in REAL_TIME_HEADERS if header in headers), None)
        if realtime_col is not None:
            ws.column_dimensions[ws.cell(1, realtime_col).column_letter].width = 24

    wb.save(WORKBOOK_PATH)
    save_cache(cache)

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "workbook": str(WORKBOOK_PATH),
        "backup": str(backup_path),
        "fetchMinId": fetch_min_id,
        "fetchMaxId": fetch_max_id,
        "updateMinId": update_min_id,
        "updateMaxId": update_max_id,
        "removeListingIds": remove_listing_ids,
        "removedWorkbookRows": removed_workbook_rows,
        "removedProductCount": len(removed_products),
        "removedProducts": removed_products,
        "changedProductCount": len(changed_products),
        "changedProducts": changed_products,
        "skippedIds": skipped_ids,
        "realTimeFetchCount": len([value for value in realtime_by_url.values() if value]),
        "realTimeFetchErrors": realtime_fetch_errors,
    }
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
