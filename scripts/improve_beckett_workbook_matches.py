#!/usr/bin/env python3
"""Improve Beckett workbook matches without echoing account credentials.

The normal Beckett generator can search with a logged-in account, but this
repair pass is intentionally credential-free. It builds a candidate index from
the existing Beckett search caches, then publicly refreshes only the selected
card pages. This gives us a safer second pass over unresolved rows while
preserving manually confirmed matches already present in the workbooks.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
import json
import re
import shutil
import sys
import time
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from beckett_legacy_pricing import (  # noqa: E402
    SPORT_CATEGORIES,
    build_workbook,
    fetch_card_page,
    format_money,
    format_percent,
    is_graded,
    load_cache,
    matched_grade_price,
    midpoint_from_range,
    normalize_card_code,
    normalize_key,
    normalize_spaces,
    price_range,
    price_signal,
    save_cache,
    selected_products,
)
from improve_beckett_nonlegacy_matches import (  # noqa: E402
    enhanced_query_variants,
    hard_conflicts,
    deep_score,
)
from update_beckett_realtime_column import (  # noqa: E402
    normalize_realtime_label,
    parse_realtime_pricing,
)


DOCS = Path.home() / "Documents" / "eBay Docs" / "Listing Automation"
BACKUP_ROOT = DOCS / "backups"
DEFAULT_LEGACY_WORKBOOK = DOCS / "Beckett Legacy Pricing Review.xlsx"
DEFAULT_NONLEGACY_WORKBOOK = DOCS / "Beckett Non-Legacy Pricing Review.xlsx"
DEFAULT_ALL_CACHE = DOCS / "beckett_pricing_cache_all.json"
DEFAULT_LEGACY_CACHE = DOCS / "beckett_legacy_pricing_cache.json"
REPORT_PATH = ROOT / "outputs" / "beckett-improved-cache-match-report.json"

MATCHABLE_STATUSES = {
    "Low confidence match - review",
    "Uncertain match - review",
    "No Beckett result found",
    "No confident Beckett match",
}
SKIP_STATUSES = {
    "Matched",
    "Not searched - non sports-card row",
    "Not searched - non sports-card legacy row",
    "REMOVE LISTING",
    "REMOVE LISTING ",
    "Sold",
    "NOT AVAILABLE",
}
REALTIME_HEADER = "Beckett Real Time Pricing"
YEAR_RE = re.compile(r"\b((?:18|19|20)\d{2})(?:[-/](?:\d{2}|\d{4}))?\b")
CARD_CODE_RE = re.compile(r"(?:#|no\.?\s*)([A-Za-z0-9]+(?:[-/][A-Za-z0-9]+)*)", re.I)
CRITICAL_SET_TOKENS = {
    "absolute",
    "artistry",
    "authentic",
    "best",
    "black",
    "bleachers",
    "bowman",
    "certified",
    "champions",
    "chrome",
    "classics",
    "contenders",
    "crown",
    "deck",
    "elite",
    "edition",
    "event",
    "finest",
    "flair",
    "fleer",
    "friday",
    "gameday",
    "gold",
    "golden",
    "head",
    "heritage",
    "holidays",
    "hit",
    "hoops",
    "hot",
    "illusions",
    "impact",
    "impulse",
    "instant",
    "karat",
    "leaf",
    "logo",
    "materials",
    "metal",
    "numbers",
    "origins",
    "optic",
    "paint",
    "panini",
    "paper",
    "patches",
    "pacific",
    "playbook",
    "plays",
    "prime",
    "prizm",
    "royale",
    "score",
    "select",
    "signature",
    "spectra",
    "sportkings",
    "star",
    "sterling",
    "tie",
    "trading",
    "threads",
    "topps",
    "tribute",
    "ultra",
    "upper",
    "volume",
    "warriors",
    "white",
}
TARGET_ONLY_BLOCK_TOKENS = {
    "artistry",
    "booklet",
    "commemorative",
    "crowd",
    "dye",
    "event",
    "golden",
    "head",
    "illusions",
    "impact",
    "logo",
    "mini",
    "pins",
    "tie",
    "warriors",
    "white",
}
BENIGN_TARGET_ONLY_TOKENS = {
    # Beckett often expands abbreviated autograph/relic inserts into these
    # words. They are okay when the source already has the same feature family.
    "signature",
    "signatures",
    "swatches",
    "materials",
    "memorabilia",
}
TOKEN_EQUIVALENTS = {
    "autograph": {"auto", "autographs", "autos"},
    "autographs": {"auto", "autograph", "autos"},
    "auto": {"autograph", "autographs", "autos", "signature", "signatures"},
    "autos": {"autograph", "autographs", "auto", "signature", "signatures"},
    "jersey": {"jsy", "relic", "swatch", "material", "materials"},
    "patch": {"patches", "prime", "swatch", "material", "materials"},
    "patches": {"patch", "prime", "swatch", "material", "materials"},
    "relic": {"relics", "patch", "patches", "swatch", "material", "materials"},
    "relics": {"relic", "patch", "patches", "swatch", "material", "materials"},
}


@dataclass(frozen=True)
class Candidate:
    title: str
    url: str
    sport: str
    years: frozenset[str]
    tokens: frozenset[str]
    compact: str
    cached_page: bool


def load_rows(workbook_path: Path, main_sheet: str) -> tuple[list[str], list[dict[str, Any]]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    ws = wb[main_sheet]
    headers = [normalize_spaces(cell.value) for cell in ws[1]]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if values[0] is None:
            continue
        row = {header: values[index] for index, header in enumerate(headers)}
        for header in headers:
            row.setdefault(header, "")
        rows.append(row)
    wb.close()
    return headers, rows


def backup_workbook(workbook_path: Path, scope: str) -> Path:
    stamp = datetime.now().strftime(f"%Y%m%d-%H%M%S-{scope}-improved-match")
    backup_dir = BACKUP_ROOT / stamp
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / workbook_path.name
    shutil.copy2(workbook_path, backup_path)
    return backup_path


def sport_from_url(url: str) -> str:
    lowered = url.lower()
    for sport in SPORT_CATEGORIES:
        if f"/{sport.lower()}/" in lowered:
            return sport
    if "/multisport/" in lowered:
        return "Multisport"
    return ""


def row_product(row: dict[str, Any], products_by_id: dict[int, dict[str, Any]]) -> dict[str, Any] | None:
    try:
        return products_by_id.get(int(row.get("Product ID")))
    except (TypeError, ValueError):
        return None


def row_category(row: dict[str, Any], product: dict[str, Any] | None) -> str:
    return normalize_spaces(row.get("Category") or (product or {}).get("category"))


def row_year(row: dict[str, Any], product: dict[str, Any] | None) -> str:
    raw = normalize_spaces(row.get("Site Year") or (product or {}).get("year"))
    if raw:
        return raw.split("-")[0].split("/")[0]
    match = YEAR_RE.search(normalize_spaces(row.get("Title") or (product or {}).get("name")))
    return match.group(1) if match else ""


def card_codes(value: Any) -> set[str]:
    return {normalize_card_code(match.group(1)) for match in CARD_CODE_RE.finditer(str(value or ""))}


def explicit_card_codes(value: Any) -> set[str]:
    return {code for code in card_codes(value) if code}


def candidate_payload(candidate: Candidate) -> dict[str, str]:
    return {"title": candidate.title, "url": candidate.url}


def missing_critical_set_terms(row: dict[str, Any], candidate: Candidate) -> list[str]:
    """Reject fuzzy matches that jump to a different Beckett set/parallel.

    Player, year, and card number can make unrelated card pages look deceptively
    strong. These terms are the collector-facing set/brand words that should not
    silently change during an automated match.
    """

    source_tokens = set(normalize_key(row.get("Title")).split())
    target_tokens = set(candidate.tokens)
    missing: list[str] = []
    for token in sorted(source_tokens & CRITICAL_SET_TOKENS):
        equivalents = TOKEN_EQUIVALENTS.get(token, set())
        if token not in target_tokens and not (equivalents & target_tokens):
            missing.append(token)
    # A named color/parallel in both title and Beckett candidate must agree.
    source_colors = source_tokens & {"black", "blue", "bronze", "gold", "green", "orange", "pink", "purple", "red", "silver"}
    target_colors = target_tokens & {"black", "blue", "bronze", "gold", "green", "orange", "pink", "purple", "red", "silver"}
    if source_colors and target_colors and not (source_colors & target_colors):
        missing.append(f"color mismatch ({', '.join(sorted(source_colors))} vs {', '.join(sorted(target_colors))})")
    source_has_auto = bool({"autograph", "autographs", "auto", "autos", "signed", "signature", "signatures"} & source_tokens)
    source_has_mem = bool(
        {"jersey", "jsy", "patch", "patches", "relic", "relics", "swatch", "swatches", "material", "materials", "memorabilia"}
        & source_tokens
    )
    target_only_block_tokens = set(TARGET_ONLY_BLOCK_TOKENS)
    if source_has_auto:
        target_only_block_tokens -= {"signature"}
    if source_has_mem:
        target_only_block_tokens -= {"materials", "memorabilia", "swatches"}
    target_only = sorted((target_tokens & target_only_block_tokens) - source_tokens - BENIGN_TARGET_ONLY_TOKENS)
    if target_only:
        missing.append("target-only variant " + ", ".join(target_only))
    target_auto = {"autograph", "autographs", "auto", "autos", "signature", "signatures"} & target_tokens
    if target_auto and not source_has_auto:
        missing.append("target-only autograph")
    return missing


def is_multi_item_listing(title: str) -> bool:
    """Avoid promoting lots/sets to one-card Beckett pages automatically."""

    text = normalize_spaces(title)
    return bool(
        re.search(r"\s\+\s", text)
        or re.search(r"\b(?:lot)\b", text, flags=re.I)
        or re.search(r"\bset\b", text, flags=re.I)
        or re.search(r"\(\s*x\s*\d+\s*\)", text, flags=re.I)
        or re.search(r"\b(?:base|insert|rookie)\s+set\b", text, flags=re.I)
        or re.search(r"(?:^|\s)w/\s|\bwith\b|\bin[-\s]?person\b", text, flags=re.I)
    )


def is_single_card_exact_match(row: dict[str, Any], candidate: Candidate) -> bool:
    source_codes = explicit_card_codes(row.get("Title"))
    target_codes = explicit_card_codes(candidate.title)
    if source_codes and target_codes and source_codes & target_codes:
        return True
    return False


def exact_card_code_conflicts(row: dict[str, Any], candidate: Candidate) -> list[str]:
    source_codes = explicit_card_codes(row.get("Title"))
    target_codes = explicit_card_codes(candidate.title)
    if source_codes and target_codes and not (source_codes & target_codes):
        return [f"card number mismatch: {', '.join(sorted(source_codes))} vs {', '.join(sorted(target_codes))}"]
    return []


def material_specific_conflicts(row: dict[str, Any], candidate: Candidate) -> list[str]:
    """Keep patch/prime relic listings from degrading into generic jersey hits."""

    source_tokens = set(normalize_key(row.get("Title")).split())
    target_tokens = set(candidate.tokens)
    conflicts: list[str] = []
    if {"patch", "patches"} & source_tokens and not ({"patch", "patches", "prime"} & target_tokens):
        conflicts.append("patch listing matched to non-patch material page")
    if {"jersey", "jerseys", "jsy"} & source_tokens and {"patch", "patches", "prime"} & target_tokens:
        conflicts.append("jersey listing matched to patch/prime page")
    return conflicts


def unique_candidate_records(caches: list[dict[str, Any]]) -> list[Candidate]:
    """Flatten cached searches/cards into one deduplicated candidate universe."""

    records: dict[str, str] = {}
    cached_pages: set[str] = set()
    for cache in caches:
        for url, card in cache.get("cards", {}).items():
            title = normalize_spaces(card.get("title"))
            if title and sport_from_url(url):
                records[url] = title
                cached_pages.add(url)
        for results in cache.get("searches", {}).values():
            if not isinstance(results, list):
                continue
            for item in results:
                url = normalize_spaces(item.get("url"))
                title = normalize_spaces(item.get("title"))
                if title and sport_from_url(url) and url not in records:
                    records[url] = title

    candidates: list[Candidate] = []
    for url, title in records.items():
        blob = f"{title} {url}"
        candidates.append(
            Candidate(
                title=title,
                url=url,
                sport=sport_from_url(url),
                years=frozenset(YEAR_RE.findall(blob)),
                tokens=frozenset(normalize_key(blob).split()),
                compact=normalize_card_code(blob),
                cached_page=url in cached_pages,
            )
        )
    return candidates


def build_indexes(candidates: list[Candidate], caches: list[dict[str, Any]]) -> dict[str, Any]:
    by_year: dict[tuple[str, str], list[Candidate]] = defaultdict(list)
    by_token: dict[tuple[str, str], list[Candidate]] = defaultdict(list)
    by_code: dict[str, list[Candidate]] = defaultdict(list)
    by_url = {candidate.url: candidate for candidate in candidates}
    search_by_norm: dict[str, list[Candidate]] = defaultdict(list)
    for candidate in candidates:
        # Only page-cached cards join the fuzzy indexes. Direct cached search
        # hits may still use the larger candidate universe, but the expensive
        # fallback stays small and price-ready.
        if not candidate.cached_page:
            continue
        sports = [candidate.sport]
        if candidate.sport == "Multisport":
            sports = list(SPORT_CATEGORIES)
        for sport in sports:
            for year in candidate.years:
                by_year[(sport, year)].append(candidate)
            for token in candidate.tokens:
                if len(token) >= 4:
                    by_token[(sport, token)].append(candidate)
        for match in CARD_CODE_RE.finditer(candidate.title):
            code = normalize_card_code(match.group(1))
            if len(code) >= 2:
                by_code[code].append(candidate)
    for cache in caches:
        for query, results in cache.get("searches", {}).items():
            if not isinstance(results, list):
                continue
            norm_query = normalize_key(query)
            for item in results:
                candidate = by_url.get(normalize_spaces(item.get("url")))
                if candidate:
                    search_by_norm[norm_query].append(candidate)

    return {"by_year": by_year, "by_token": by_token, "by_code": by_code, "by_url": by_url, "search_by_norm": search_by_norm}


def cached_search_results(query: str, indexes: dict[str, Any]) -> list[Candidate]:
    norm_query = normalize_key(query)
    results: list[Candidate] = []
    seen: set[str] = set()
    for candidate in indexes["search_by_norm"].get(norm_query, []):
        if candidate.url not in seen:
            seen.add(candidate.url)
            results.append(candidate)
    return results


def distinctive_tokens(row: dict[str, Any], product: dict[str, Any] | None) -> list[str]:
    text = " ".join(
        normalize_spaces(value)
        for value in [
            row.get("Title"),
            row.get("Site Player / Athlete"),
            row.get("Site Team / Publisher"),
            (product or {}).get("playerAthlete"),
            (product or {}).get("team"),
        ]
        if value
    )
    stop = {
        "card",
        "cards",
        "rookie",
        "auto",
        "autograph",
        "serial",
        "numbered",
        "near",
        "mint",
        "baseball",
        "basketball",
        "football",
        "topps",
        "upper",
        "deck",
        "panini",
        "fleer",
        "donruss",
        "bowman",
    }
    tokens = [token for token in normalize_key(text).split() if len(token) >= 4 and token not in stop]
    counts = Counter(tokens)
    return [token for token, _ in counts.most_common(10)]


def candidate_pool(
    row: dict[str, Any],
    product: dict[str, Any],
    indexes: dict[str, Any],
) -> list[Candidate]:
    category = row_category(row, product)
    year = row_year(row, product)
    by_year = indexes["by_year"]
    by_token = indexes["by_token"]
    by_code = indexes["by_code"]

    pool: dict[str, Candidate] = {}
    for query in enhanced_query_variants(product, row):
        for candidate in cached_search_results(query, indexes):
            pool[candidate.url] = candidate

    if year:
        for candidate in by_year.get((category, year), []):
            pool[candidate.url] = candidate

    for code in card_codes(row.get("Title")):
        for candidate in by_code.get(code, []):
            if candidate.sport in {category, "Multisport"}:
                pool[candidate.url] = candidate
        if len(code) >= 3:
            for candidate in indexes["by_year"].get((category, year), []) if year else []:
                if code in candidate.compact:
                    pool[candidate.url] = candidate

    # Add candidates sharing multiple meaningful player/set/team tokens. This
    # helps when a cached Beckett search was created under a different query.
    token_hits: Counter[str] = Counter()
    token_candidates: dict[str, Candidate] = {}
    for token in distinctive_tokens(row, product):
        for candidate in by_token.get((category, token), [])[:3500]:
            token_hits[candidate.url] += 1
            token_candidates[candidate.url] = candidate
    for url, hit_count in token_hits.items():
        candidate = token_candidates[url]
        if hit_count >= 2 or (year and year in candidate.years):
            pool[url] = candidate

    return list(pool.values())


def best_candidate_from_cache(
    row: dict[str, Any],
    product: dict[str, Any],
    indexes: dict[str, Any],
) -> tuple[Candidate | None, str, float, list[str], int]:
    pool = candidate_pool(row, product, indexes)
    if not pool:
        return None, "", 0.0, [], 0

    best: tuple[float, str, Candidate, list[str]] | None = None
    direct_queries = enhanced_query_variants(product, row)
    query_for_score = direct_queries[0] if direct_queries else normalize_spaces(row.get("Title"))
    for candidate in pool:
        payload = candidate_payload(candidate)
        conflicts = hard_conflicts(product, row, payload)
        if is_multi_item_listing(normalize_spaces(row.get("Title"))):
            conflicts = [*conflicts, "multi-card lot/set requires manual review"]
        code_conflicts = exact_card_code_conflicts(row, candidate)
        if code_conflicts:
            conflicts = [*conflicts, *code_conflicts]
        conflicts = [*conflicts, *material_specific_conflicts(row, candidate)]
        critical_missing = missing_critical_set_terms(row, candidate)
        if critical_missing:
            conflicts = [*conflicts, "set/parallel term mismatch: " + ", ".join(critical_missing)]
        score = deep_score(product, row, query_for_score, payload)
        if best is None or (len(conflicts) == 0, score) > (len(best[3]) == 0, best[0]):
            best = (score, query_for_score, candidate, conflicts)
    if best is None:
        return None, "", 0.0, [], len(pool)
    score, query, candidate, conflicts = best
    return candidate, query, score, conflicts, len(pool)


def realtime_for_url(session: requests.Session, url: str, cache: dict[str, Any], delay: float) -> str:
    card_cache = cache.setdefault("cards", {}).setdefault(url, {})
    cached = normalize_realtime_label(card_cache.get("real_time_pricing"))
    if cached:
        return cached
    response = session.get(url, timeout=45)
    response.raise_for_status()
    label = parse_realtime_pricing(response.text)
    if label:
        card_cache["real_time_pricing"] = label
        card_cache["real_time_pricing_fetched_at"] = datetime.now().isoformat(timespec="seconds")
    if delay:
        time.sleep(delay)
    return label


def set_if_present(row: dict[str, Any], header: str, value: Any) -> None:
    if header in row:
        row[header] = value


def apply_match(
    row: dict[str, Any],
    product: dict[str, Any],
    candidate: Candidate,
    query: str,
    score: float,
    session: requests.Session,
    cache: dict[str, Any],
    delay: float,
) -> None:
    card = fetch_card_page(session, candidate.url, cache, delay)
    raw_prices = card.get("raw_prices", {})
    graded_prices = card.get("graded_prices", {})
    raw_range = price_range(raw_prices)
    graded_range = price_range(graded_prices)
    matched_price, matched_basis, grade_note = matched_grade_price(
        str(product.get("condition", "") or row.get("Current Grade / Condition", "")),
        raw_prices,
        graded_prices,
    )
    realtime = realtime_for_url(session, candidate.url, cache, delay)
    graded = is_graded(str(product.get("condition", "") or row.get("Current Grade / Condition", "")))
    price_choice = graded_range if graded and graded_range else raw_range
    if not graded and realtime:
        price_choice = realtime

    row["Lookup Query"] = query
    row["Match Status"] = "Matched"
    row["Match Confidence"] = round(score, 4)
    row["Beckett Matched Title"] = normalize_spaces(card.get("title") or candidate.title)
    row["Beckett URL"] = candidate.url
    row["Beckett Price Range"] = price_choice
    row["Beckett Raw Range"] = raw_range
    row["Beckett Graded Range"] = graded_range
    set_if_present(row, "Review Bucket", "Enhanced confirmed match")
    set_if_present(row, "Beckett Raw Market Range", card.get("raw_market", {}).get("range", ""))
    set_if_present(row, "Beckett Graded Market Range", card.get("graded_market", {}).get("range", ""))
    set_if_present(row, "Beckett Availability", "Guide pricing" if raw_range or graded_range else "No visible pricing")
    set_if_present(row, "Beckett Page Source", card.get("source_page_title", ""))
    set_if_present(row, "Beckett Matched Grade Price", matched_price)
    set_if_present(row, "Beckett Matched Grade Basis", matched_basis)
    set_if_present(row, REALTIME_HEADER, realtime)

    site_mid = midpoint_from_range(row.get("Original Price / Range") or "")
    beckett_mid = midpoint_from_range(row.get("Beckett Price Range") or "")
    delta_money, delta_pct, signal = price_signal(site_mid, beckett_mid)
    set_if_present(row, "Site Price Midpoint", format_money(site_mid))
    set_if_present(row, "Beckett Price Midpoint", format_money(beckett_mid))
    set_if_present(row, "Midpoint Delta", delta_money)
    set_if_present(row, "Midpoint Delta %", format_percent(delta_pct) if isinstance(delta_pct, float) else delta_pct)
    set_if_present(row, "Price Signal", signal)

    notes = ["Improved cache/public Beckett pass confirmed this match."]
    if realtime and not graded:
        notes.append("Ungraded price range uses Beckett Real Time Pricing.")
    if grade_note:
        notes.append(grade_note)
    set_if_present(row, "Notes", " ".join(notes))


def process_workbook(
    workbook_path: Path,
    scope: str,
    main_sheet: str,
    workbook_cache_path: Path,
    all_caches: list[dict[str, Any]],
    indexes: dict[str, Any],
    args: argparse.Namespace,
) -> dict[str, Any]:
    headers, rows = load_rows(workbook_path, main_sheet)
    if REALTIME_HEADER not in headers:
        headers.append(REALTIME_HEADER)
        for row in rows:
            row[REALTIME_HEADER] = ""
    products = selected_products(ROOT / "products.json", scope)
    products_by_id = {int(product["id"]): product for product in products}
    workbook_cache = load_cache(workbook_cache_path)

    stats_before = Counter(normalize_spaces(row.get("Match Status")) for row in rows)
    candidates = [
        row
        for row in rows
        if normalize_spaces(row.get("Match Status")) in MATCHABLE_STATUSES
        and row_category(row, row_product(row, products_by_id)) in SPORT_CATEGORIES
    ]
    if args.limit:
        candidates = candidates[: args.limit]

    session = requests.Session()
    session.headers.update({"User-Agent": "Mozilla/5.0 (compatible; DJ-House-of-Cards-Beckett-Review/1.0)"})
    backup_path = "" if args.dry_run else str(backup_workbook(workbook_path, scope))
    changed: list[dict[str, Any]] = []
    skipped: Counter[str] = Counter()

    for index, row in enumerate(candidates, start=1):
        product = row_product(row, products_by_id)
        if not product:
            skipped["missing product"] += 1
            continue
        original_status = normalize_spaces(row.get("Match Status"))
        candidate, query, score, conflicts, pool_size = best_candidate_from_cache(row, product, indexes)
        if not candidate:
            skipped["no cached candidate pool"] += 1
            continue
        if conflicts:
            skipped["conflicted best candidate"] += 1
            continue
        if score < args.threshold:
            skipped["below confidence threshold"] += 1
            continue
        try:
            apply_match(row, product, candidate, query, score, session, workbook_cache, args.delay)
        except Exception as exc:  # noqa: BLE001 - keep processing the workbook.
            skipped["page refresh failed"] += 1
            set_if_present(row, "Notes", f"Improved candidate found but Beckett page refresh failed: {exc}")
            continue
        changed.append(
            {
                "productId": row.get("Product ID"),
                "fromStatus": original_status,
                "score": round(score, 4),
                "candidatePoolSize": pool_size,
                "title": row.get("Title"),
                "beckettTitle": row.get("Beckett Matched Title"),
                "beckettUrl": row.get("Beckett URL"),
            }
        )
        if index % args.save_every == 0:
            print(f"{scope}: processed {index}/{len(candidates)} candidates; promoted {len(changed)}", flush=True)

    stats_after = Counter(normalize_spaces(row.get("Match Status")) for row in rows)
    if not args.dry_run:
        build_workbook(rows, workbook_path, scope)
        save_cache(workbook_cache_path, workbook_cache)

    return {
        "workbook": str(workbook_path),
        "backup": backup_path,
        "mainSheet": main_sheet,
        "scope": scope,
        "candidateRows": len(candidates),
        "changedRows": len(changed),
        "skipped": dict(skipped),
        "before": dict(stats_before),
        "after": dict(stats_after),
        "changes": changed,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--legacy-workbook", type=Path, default=DEFAULT_LEGACY_WORKBOOK)
    parser.add_argument("--nonlegacy-workbook", type=Path, default=DEFAULT_NONLEGACY_WORKBOOK)
    parser.add_argument("--legacy-cache", type=Path, default=DEFAULT_LEGACY_CACHE)
    parser.add_argument("--all-cache", type=Path, default=DEFAULT_ALL_CACHE)
    parser.add_argument("--scope", choices=["legacy", "nonlegacy", "both"], default="both")
    parser.add_argument("--threshold", type=float, default=0.86)
    parser.add_argument("--delay", type=float, default=0.08)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--save-every", type=int, default=50)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    cache_paths = [args.all_cache, args.legacy_cache]
    caches = [load_cache(path) for path in cache_paths if path.exists()]
    if not caches:
        raise SystemExit("No Beckett cache files found; cannot run credential-free match pass.")

    candidates = unique_candidate_records(caches)
    indexes = build_indexes(candidates, caches)
    report: dict[str, Any] = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "dryRun": args.dry_run,
        "threshold": args.threshold,
        "candidateUniverse": len(candidates),
        "workbooks": [],
    }

    if args.scope in {"legacy", "both"}:
        report["workbooks"].append(
            process_workbook(
                args.legacy_workbook,
                "legacy",
                "Legacy Beckett Pricing",
                args.legacy_cache,
                caches,
                indexes,
                args,
            )
        )
    if args.scope in {"nonlegacy", "both"}:
        report["workbooks"].append(
            process_workbook(
                args.nonlegacy_workbook,
                "nonlegacy",
                "Non-Legacy Beckett Pricing",
                args.all_cache,
                caches,
                indexes,
                args,
            )
        )

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
