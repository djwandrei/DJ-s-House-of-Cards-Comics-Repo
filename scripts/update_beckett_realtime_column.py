#!/usr/bin/env python3
"""Add/fill Beckett Real Time Pricing columns in Beckett review workbooks.

The pricing-review workbooks have several sheets that repeat the same product
rows for review queues. This utility keeps the "Beckett Real Time Pricing"
column consistent everywhere by treating Beckett URLs as the source key and by
recording only values shown in the Beckett Real Time Pricing section itself.
It intentionally does not fall back to guide/raw/graded pricing fields.
"""

from __future__ import annotations

import argparse
from collections import Counter
import json
import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = (
    Path.home()
    / "Documents"
    / "eBay Docs"
    / "Listing Automation"
    / "Beckett Non-Legacy Pricing Review.xlsx"
)
DEFAULT_CACHE = (
    Path.home()
    / "Documents"
    / "eBay Docs"
    / "Listing Automation"
    / "beckett_pricing_cache_all.json"
)
BACKUP_ROOT = Path.home() / "Documents" / "eBay Docs" / "Listing Automation" / "backups"
REPORT_PATH = ROOT / "outputs" / "beckett-realtime-column-report.json"

REAL_TIME_HEADER = "Beckett Real Time Pricing"
URL_HEADER = "Beckett URL"
REAL_TIME_RE = re.compile(
    r"Beckett\s+Real\s+Time\s+Pricing\s*(\$?\s*\d[\d,]*(?:\.\d+)?)\s*to\s*(\$?\s*\d[\d,]*(?:\.\d+)?)",
    re.I,
)
MONEY_RE = re.compile(r"\$?\s*(?:\d[\d,]*(?:\.\d+)?|\.\d+)")


def load_cache(cache_path: Path) -> dict[str, Any]:
    if cache_path.exists():
        return json.loads(cache_path.read_text(encoding="utf-8"))
    return {"searches": {}, "cards": {}, "completed": {}}


def save_cache(cache_path: Path, cache: dict[str, Any]) -> None:
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(json.dumps(cache, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def money_to_float(value: Any) -> float | None:
    text = str(value or "").replace(",", "").replace("$", "").strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def format_money(value: float) -> str:
    return f"${value:,.2f}"


def normalize_price_label(value: Any) -> str:
    values: list[float] = []
    for match in MONEY_RE.findall(str(value or "")):
        number = money_to_float(match)
        if number is not None:
            values.append(number)
    if not values:
        return ""
    if len(values) == 1:
        return format_money(values[0])
    low, high = min(values[:2]), max(values[:2])
    return f"{format_money(low)}-{format_money(high)}"


def normalize_realtime_label(value: Any) -> str:
    """Return a normalized Beckett real-time label, preserving explicit N/A."""

    text = str(value or "").strip()
    if re.search(r"\bN/A\b", text, flags=re.I):
        return "N/A"
    return normalize_price_label(text)


def get_headers(ws) -> dict[str, int]:
    return {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}


def ensure_realtime_column(ws) -> int:
    headers = get_headers(ws)
    if REAL_TIME_HEADER in headers:
        return headers[REAL_TIME_HEADER]

    target_col = ws.max_column + 1
    ws.cell(1, target_col).value = REAL_TIME_HEADER
    ws.cell(1, target_col).fill = PatternFill("solid", fgColor="1F2937")
    ws.cell(1, target_col).font = Font(color="FFFFFF", bold=True)
    return target_col


def parse_realtime_pricing(html: str) -> str:
    """Extract the exact Beckett Real Time Pricing range from a card page."""

    soup = BeautifulSoup(html, "html.parser")
    rt_heading = soup.select_one(".rtHad")
    if rt_heading:
        text = rt_heading.parent.get_text(" ", strip=True)
        match = REAL_TIME_RE.search(text)
        if match:
            return normalize_price_label(f"{match.group(1)}-{match.group(2)}")
        if re.search(r"Beckett\s+Real\s+Time\s+Pricing\s+N/A\b", text, flags=re.I):
            return "N/A"

    text = soup.get_text(" ", strip=True)
    match = REAL_TIME_RE.search(text)
    if match:
        return normalize_price_label(f"{match.group(1)}-{match.group(2)}")
    if re.search(r"Beckett\s+Real\s+Time\s+Pricing\s+N/A\b", text, flags=re.I):
        return "N/A"
    return ""


def cached_realtime_pricing(url: str, cache: dict[str, Any]) -> str:
    return normalize_realtime_label(cache.get("cards", {}).get(url, {}).get("real_time_pricing"))


def fetch_realtime_pricing(
    url: str,
    session: Any,
    cache: dict[str, Any],
    delay: float,
    max_retries: int,
    retry_delay: float,
) -> str:
    """Fetch one card page and cache only the real-time section value."""

    card_cache = cache.setdefault("cards", {}).setdefault(url, {})
    cached = normalize_realtime_label(card_cache.get("real_time_pricing"))
    if cached:
        return cached

    last_error: Exception | None = None
    for attempt in range(1, max(1, max_retries) + 1):
        if attempt == 1:
            time.sleep(max(0, delay))
        else:
            time.sleep(max(0, retry_delay * attempt))
        try:
            response = session.get(url, timeout=60)
            response.raise_for_status()
            break
        except Exception as error:  # noqa: BLE001 - retried and reported by URL.
            last_error = error
    else:
        if last_error:
            raise last_error
        return ""

    label = parse_realtime_pricing(response.text)
    if label:
        card_cache["real_time_pricing"] = label
        card_cache["real_time_pricing_fetched_at"] = datetime.now().isoformat(timespec="seconds")
    return label


def backup_workbook(workbook_path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-beckett-realtime")
    backup_dir = BACKUP_ROOT / stamp
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / workbook_path.name
    shutil.copy2(workbook_path, backup_path)
    return backup_path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--cache", type=Path, default=DEFAULT_CACHE)
    parser.add_argument("--fetch-realtime", action="store_true")
    parser.add_argument("--clear-realtime", action="store_true")
    parser.add_argument("--delay", type=float, default=0.15)
    parser.add_argument("--max-retries", type=int, default=3)
    parser.add_argument("--retry-delay", type=float, default=1.5)
    parser.add_argument("--abort-after-consecutive-errors", type=int, default=35)
    parser.add_argument("--beckett-email", default="")
    parser.add_argument("--beckett-password", default="")
    args = parser.parse_args()

    workbook_path = args.workbook
    cache_path = args.cache
    if not workbook_path.exists():
        raise SystemExit(f"Workbook does not exist: {workbook_path}")

    cache = load_cache(cache_path)
    backup_path = backup_workbook(workbook_path)
    wb = load_workbook(workbook_path)

    unique_urls: list[str] = []
    seen_urls: set[str] = set()
    for ws in wb.worksheets:
        headers = get_headers(ws)
        url_col = headers.get(URL_HEADER)
        if not url_col:
            continue
        for row_index in range(2, ws.max_row + 1):
            url = str(ws.cell(row_index, url_col).value or "").strip()
            if url and url not in seen_urls:
                seen_urls.add(url)
                unique_urls.append(url)
    unique_urls.sort()
    realtime_by_url = {url: cached_realtime_pricing(url, cache) for url in unique_urls}

    session = None
    if args.fetch_realtime and any(not value for value in realtime_by_url.values()):
        sys.path.insert(0, str(ROOT / "scripts"))
        from beckett_legacy_pricing import create_session  # type: ignore

        if not args.beckett_email or not args.beckett_password:
            raise SystemExit("--beckett-email and --beckett-password are required with --fetch-realtime")
        session = create_session(args.beckett_email, args.beckett_password)

    fetch_errors: dict[str, str] = {}
    error_types: Counter[str] = Counter()
    consecutive_errors = 0
    aborted_reason = ""
    fetched_count = 0
    if session is not None:
        remaining = [url for url, label in realtime_by_url.items() if not label]
        for index, url in enumerate(remaining, start=1):
            try:
                label = fetch_realtime_pricing(
                    url,
                    session,
                    cache,
                    args.delay,
                    args.max_retries,
                    args.retry_delay,
                )
                if label:
                    fetched_count += 1
                consecutive_errors = 0
                realtime_by_url[url] = label
                save_cache(cache_path, cache)
            except Exception as error:  # noqa: BLE001 - keep processing other review rows.
                error_text = str(error)
                fetch_errors[url] = error_text
                if "Failed to resolve" in error_text:
                    error_types["dns"] += 1
                elif "Read timed out" in error_text or "timed out" in error_text:
                    error_types["timeout"] += 1
                elif "403" in error_text:
                    error_types["forbidden"] += 1
                elif "404" in error_text:
                    error_types["not_found"] += 1
                else:
                    error_types[type(error).__name__] += 1
                consecutive_errors += 1
                realtime_by_url[url] = ""
                if consecutive_errors >= args.abort_after_consecutive_errors:
                    aborted_reason = (
                        f"Stopped after {consecutive_errors} consecutive Beckett fetch errors; "
                        "rerun later to continue from the cache."
                    )
                    print(aborted_reason, flush=True)
                    break
            if index % 50 == 0:
                print(f"Fetched real-time pricing for {index}/{len(remaining)} uncached URLs...", flush=True)

    sheet_stats: dict[str, dict[str, int]] = {}
    url_row_count = 0

    for ws in wb.worksheets:
        headers = get_headers(ws)
        url_col = headers.get(URL_HEADER)
        if not url_col:
            continue
        realtime_col = ensure_realtime_column(ws)
        stats = {
            "rowsWithUrl": 0,
            "dollarRangeRows": 0,
            "naRows": 0,
            "missingRows": 0,
            "clearedRows": 0,
        }
        for row_index in range(2, ws.max_row + 1):
            url = str(ws.cell(row_index, url_col).value or "").strip()
            if args.clear_realtime:
                ws.cell(row_index, realtime_col).value = None
                stats["clearedRows"] += 1
            if not url:
                continue
            stats["rowsWithUrl"] += 1
            url_row_count += 1
            label = realtime_by_url.get(url) or cached_realtime_pricing(url, cache)
            if label:
                ws.cell(row_index, realtime_col).value = label
                if label == "N/A":
                    stats["naRows"] += 1
                else:
                    stats["dollarRangeRows"] += 1
            else:
                stats["missingRows"] += 1
        ws.column_dimensions[ws.cell(1, realtime_col).column_letter].width = 24
        sheet_stats[ws.title] = stats

    wb.save(workbook_path)
    save_cache(cache_path, cache)

    report = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "workbook": str(workbook_path),
        "backup": str(backup_path),
        "cache": str(cache_path),
        "sheets": sheet_stats,
        "uniqueUrlCount": len(unique_urls),
        "urlRowCount": url_row_count,
        "realtimeCachedOrFetchedCount": len([value for value in realtime_by_url.values() if value]),
        "fetchedCount": fetched_count,
        "fetchErrorCount": len(fetch_errors),
        "fetchErrorTypes": dict(error_types),
        "fetchErrorsSample": dict(list(fetch_errors.items())[:25]),
        "abortedReason": aborted_reason,
    }
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
