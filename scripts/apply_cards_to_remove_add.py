#!/usr/bin/env python3
"""Apply DJ's add/remove workbook to the storefront catalog.

This migration is intentionally repeatable. It reads the workbook supplied by
DJ, removes the requested listings, adds the requested card records, copies
matched photos into the live assets tree, and rebuilds the static product
bundles used by the storefront fallback.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import shutil
import sys
import time
from copy import deepcopy
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

import beckett_legacy_pricing as beckett  # noqa: E402
from apply_beckett_legacy_updates import (  # noqa: E402
    fetch_realtime_pricing,
    midpoint_from_label,
    normalize_price_label,
    rebuild_product_files,
)

WORKBOOK_PATH = Path.home() / "Documents" / "Cards to remove and add.xlsx"
PHOTO_ROOT = Path("H:/My Drive/Unused Assets/unused-legacy-photos/baseball_cardx")
NEW_PHOTO_DIR = PHOTO_ROOT / "New Site Listings"
REPLACE_PHOTO_DIR = PHOTO_ROOT / "Replace Photos"
PRODUCTS_PATH = ROOT / "products.json"
OUTPUT_DIR = ROOT / "outputs"
REPORT_PATH = OUTPUT_DIR / "cards-remove-add-report.json"
CACHE_PATH = Path.home() / "Documents" / "eBay Docs" / "Listing Automation" / "beckett_cards_add_cache.json"

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".svg"}
AUTOGRAPH_RE = re.compile(r"\b(auto(?:graph)?|autographed|signed|signature)\b", re.I)
GRADE_RE = re.compile(r"\b(PSA|BGS|SGC|CGC|HGA|BCCG|GAI)\s*(10|[1-9](?:\.\d)?)\b", re.I)
YEAR_RE = re.compile(r"\b(19\d{2}|20\d{2})\b")
CARD_NO_RE = re.compile(r"#\s*([A-Za-z0-9]+(?:[-/][A-Za-z0-9]+)*)", re.I)
MONEY_RE = re.compile(r"\$(?:\d[\d,]*(?:\.\d{1,2})?|\.\d{1,2})")

NOISE_TOKENS = {
    "topps",
    "kelloggs",
    "kellogg",
    "donruss",
    "baseball",
    "best",
    "bronze",
    "rookies",
    "rookie",
    "card",
    "cards",
    "psa",
    "bgs",
    "sgc",
    "gem",
    "mint",
    "nm",
    "ex",
    "vg",
    "mc",
    "upd",
    "all",
    "star",
    "as",
    "leaders",
    "leader",
    "strikeout",
    "record",
    "breaker",
    "special",
    "mvp",
    "mvps",
    "world",
    "series",
    "sereis",
    "first",
    "baseman",
    "action",
    "in",
    "nl",
    "al",
    "mini",
}

VARIANT_COLOR_TOKENS = {
    "black",
    "blue",
    "cyan",
    "gold",
    "green",
    "magenta",
    "orange",
    "plate",
    "plates",
    "printing",
    "purple",
    "red",
    "silver",
    "superfractor",
    "superfractors",
    "yellow",
}

TEAM_OVERRIDES: dict[tuple[str, str], str] = {
    ("1970", "453"): "Minnesota Twins",
    ("1970", "500"): "Atlanta Braves",
    ("1970", "600"): "San Francisco Giants",
    ("1971", "130"): "Pittsburgh Pirates",
    ("1971", "260"): "Milwaukee Brewers",
    ("1971", "341"): "Los Angeles Dodgers",
    ("1971", "360"): "California Angels",
    ("1971", "400"): "Atlanta Braves",
    ("1971", "470"): "Kansas City Royals",
    ("1971", "513"): "New York Mets",
    ("1971", "600"): "San Francisco Giants",
    ("1971", "630"): "Pittsburgh Pirates",
    ("1972", "49"): "San Francisco Giants",
    ("1972", "50"): "San Francisco Giants",
    ("1972", "595"): "California Angels",
    ("1972", "777"): "Los Angeles Dodgers",
    ("1973", "23"): "San Francisco Giants",
    ("1973", "25"): "New York Yankees",
    ("1973", "50"): "Pittsburgh Pirates",
    ("1973", "51"): "Minnesota Twins",
    ("1973", "67"): "Philadelphia Phillies | California Angels",
    ("1973", "90"): "Baltimore Orioles",
    ("1973", "100"): "Atlanta Braves",
    ("1973", "142"): "New York Yankees",
    ("1973", "193"): "Boston Red Sox",
    ("1973", "213"): "Los Angeles Dodgers",
    ("1973", "220"): "California Angels",
    ("1973", "232"): "Montreal Expos",
    ("1973", "245"): "Boston Red Sox",
    ("1973", "255"): "Oakland Athletics",
    ("1973", "300"): "Philadelphia Phillies",
    ("1973", "305"): "New York Mets",
    ("1973", "325"): "Baltimore Orioles",
    ("1973", "434"): "San Francisco Giants",
    ("1973", "498"): "New York Yankees",
    ("1974", "1"): "Atlanta Braves | Milwaukee Braves",
    ("1974", "2"): "Atlanta Braves | Milwaukee Braves",
    ("1974", "3"): "Atlanta Braves | Milwaukee Braves",
    ("1974", "4"): "Atlanta Braves | Milwaukee Braves",
    ("1974", "5"): "Atlanta Braves | Milwaukee Braves",
    ("1974", "6"): "Atlanta Braves | Milwaukee Braves",
    ("1974", "60"): "St. Louis Cardinals",
    ("1974", "100"): "Pittsburgh Pirates",
    ("1974", "252"): "Pittsburgh Pirates",
    ("1974", "332"): "Chicago White Sox | Atlanta Braves",
    ("1974", "473"): "New York Mets",
    ("1974", "646"): "Cincinnati Reds",
    ("1975", "50"): "Baltimore Orioles",
    ("1975", "192"): "New York Yankees | New York Giants",
    ("1975", "201"): "New York Yankees | Los Angeles Dodgers",
    ("1975", "280"): "Boston Red Sox",
    ("1975", "312"): "California Angels | Philadelphia Phillies",
    ("1975", "530"): "Cleveland Indians",
    ("1976", "1"): "Milwaukee Brewers",
    ("1976", "45"): "Cleveland Indians",
    ("1976", "300"): "Cincinnati Reds",
    ("1976", "550"): "Milwaukee Brewers",
    ("1977", "130"): "Pittsburgh Pirates",
    ("1979", "1"): "Minnesota Twins | Pittsburgh Pirates",
    ("1979", "3"): "Boston Red Sox | Cincinnati Reds",
    ("1979", "25"): "Philadelphia Phillies",
    ("1979", "55"): "Pittsburgh Pirates",
    ("1979", "115"): "California Angels",
    ("1979", "170"): "Los Angeles Dodgers",
    ("1979", "200"): "Cincinnati Reds",
    ("1979", "204"): "Cincinnati Reds",
    ("1979", "321"): "San Diego Padres",
    ("1979", "413"): "New York Yankees | Milwaukee Braves | Atlanta Braves",
    ("1979", "430"): "Pittsburgh Pirates",
    ("1979", "590"): "Houston Astros",
    ("1979", "610"): "Philadelphia Phillies",
    ("2001", "R97"): "St. Louis Cardinals",
}

MANUAL_BECKETT_OVERRIDES: dict[str, str] = {
    # These rows use old-site/workbook card numbers, while Beckett uses a
    # different checklist number for the same visible card/player/set. Keeping
    # the corrections explicit avoids loosening matching for the whole catalog.
    "2011 bowman chrome prospects purple refractors 87 eric thames":
        "https://www.beckett.com/baseball/2011/bowman-chrome-prospects-purple-refractors/bcp102-eric-thames-8140564",
    "2011 topps 60 52 sandy koufax":
        "https://www.beckett.com/baseball/2011/topps-60/57-sandy-koufax-8177432",
    "2010 bowman platinum prospect autograph bpapg paul goldschmidt psa 9":
        "https://www.beckett.com/baseball/2010/bowman-platinum-prospect-autographs-refractors/pg-paul-goldschmidt-7927083",
}

PLAYER_OVERRIDES: dict[tuple[str, str], str] = {
    ("1970", "453"): "Rod Carew",
    ("1970", "500"): "Hank Aaron",
    ("1970", "600"): "Willie Mays",
    ("1971", "341"): "Steve Garvey",
    ("1973", "67"): "Steve Carlton | Nolan Ryan",
    ("1974", "332"): "Dick Allen | Hank Aaron",
    ("1975", "192"): "Yogi Berra | Willie Mays",
    ("1975", "201"): "Elston Howard | Sandy Koufax",
    ("1975", "312"): "Nolan Ryan | Steve Carlton",
    ("1979", "1"): "Rod Carew | Dave Parker",
    ("1979", "3"): "Jim Rice | George Foster",
    ("1979", "413"): "Roger Maris | Hank Aaron",
}


def normalize_spaces(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize_key(value: Any) -> str:
    return normalize_spaces(re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()))


def parse_year(value: Any) -> str:
    text = re.sub(r"\b23011\b", "2011", str(value or ""))
    match = YEAR_RE.search(text)
    return match.group(1) if match else ""


def parse_card_number(value: Any) -> str:
    match = CARD_NO_RE.search(str(value or ""))
    return match.group(1).upper() if match else ""


def lookup_title(value: str) -> str:
    """Repair small workbook typos for lookup while preserving the raw note."""
    return normalize_spaces(re.sub(r"\b23011\b", "2011", value))


def base_copy_title(value: str) -> str:
    return normalize_spaces(re.sub(r"\(\s*#?\d+\s*\)\s*$", "", lookup_title(value)))


def parse_grade(value: Any) -> tuple[str, float | None]:
    match = GRADE_RE.search(str(value or ""))
    if not match:
        return "", None
    grade = float(match.group(2))
    grade_label = str(int(grade)) if grade.is_integer() else str(grade).rstrip("0").rstrip(".")
    return f"{match.group(1).upper()} {grade_label}", grade


def parse_money_values(value: Any) -> list[float]:
    values: list[float] = []
    for match in MONEY_RE.finditer(str(value or "")):
        values.append(float(match.group(0).replace("$", "").replace(",", "")))
    return values


def safe_asset_name(source: Path) -> str:
    stem = source.stem.replace("#", "")
    stem = stem.replace("&", " and ").replace("/", "-").replace("\\", "-")
    stem = re.sub(r"[^A-Za-z0-9 ._()'-]+", " ", stem)
    stem = re.sub(r"\s+", " ", stem).strip(" .")
    stem = stem.replace("'", "")
    return f"{stem}{source.suffix.lower()}"


def card_tokens(value: Any) -> set[str]:
    tokens = set(normalize_key(value).split())
    return {token for token in tokens if token not in NOISE_TOKENS and not token.isdigit()}


def workbook_rows() -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    adds: list[dict[str, str]] = []
    removes: list[dict[str, str]] = []
    seen_adds: set[str] = set()
    seen_removes: set[str] = set()

    for ws in wb.worksheets:
        add_header = normalize_key(ws.cell(1, 1).value)
        remove_header = normalize_key(ws.cell(1, 2).value)
        if "cards to add" not in add_header or "cards to remove" not in remove_header:
            continue
        for row_index in range(2, ws.max_row + 1):
            add_value = normalize_spaces(ws.cell(row_index, 1).value)
            remove_value = normalize_spaces(ws.cell(row_index, 2).value)
            if add_value and normalize_key(add_value) not in seen_adds:
                seen_adds.add(normalize_key(add_value))
                adds.append({"sheet": ws.title, "row": str(row_index), "title": add_value})
            if remove_value and normalize_key(remove_value) not in seen_removes:
                seen_removes.add(normalize_key(remove_value))
                removes.append({"sheet": ws.title, "row": str(row_index), "title": remove_value})
    return adds, removes


def image_files() -> list[Path]:
    files: list[Path] = []
    for folder in [NEW_PHOTO_DIR, PHOTO_ROOT]:
        if not folder.exists():
            continue
        for file in folder.iterdir():
            if file.is_file() and file.suffix.lower() in IMAGE_EXTENSIONS:
                files.append(file)
    return files


def image_match_score(title: str, image: Path) -> tuple[float, str]:
    title_year = parse_year(title)
    image_year = parse_year(image.stem)
    title_number = parse_card_number(title)
    image_number = parse_card_number(image.stem)
    title_grade, title_grade_value = parse_grade(title)
    image_grade, image_grade_value = parse_grade(image.stem)

    if title_year and image_year and title_year != image_year:
        return -50.0, "year mismatch"
    if title_number and image_number and title_number != image_number:
        return -50.0, "card number mismatch"
    if title_grade_value is not None and image_grade_value is not None and title_grade_value != image_grade_value:
        return -50.0, "grade mismatch"

    score = 0.0
    if title_year and title_year == image_year:
        score += 4
    if title_number and title_number == image_number:
        score += 6
    if title_grade and title_grade == image_grade:
        score += 2
    elif title_grade and not image_grade:
        score -= 0.5

    title_tokens = card_tokens(title)
    image_tokens = card_tokens(image.stem)
    overlap = title_tokens & image_tokens
    score += len(overlap) * 1.2
    if title_tokens and image_tokens and not overlap:
        score -= 6

    if normalize_key(title).replace(" ", "") in normalize_key(image.stem).replace(" ", ""):
        score += 8
    if image.parent == NEW_PHOTO_DIR:
        score += 0.5
    return score, ""


def best_image_for_add(title: str, files: list[Path]) -> tuple[Path | None, float, str]:
    scored = sorted(
        ((image_match_score(title, file)[0], image_match_score(title, file)[1], file) for file in files),
        key=lambda item: item[0],
        reverse=True,
    )
    if not scored:
        return None, 0.0, "no image files found"
    score, reason, file = scored[0]
    threshold = 8.0 if parse_card_number(title) else 7.0
    if score < threshold:
        return None, score, reason or "below image confidence threshold"
    return file, score, ""


def candidate_card_number(candidate: dict[str, str]) -> str:
    """Extract Beckett's card number from a search result title or URL."""
    title_number = parse_card_number(candidate.get("title", ""))
    if title_number:
        return title_number
    slug = str(candidate.get("url") or "").rstrip("/").rsplit("/", 1)[-1]
    slug = re.sub(r"-\d+$", "", slug)
    first = slug.split("-", 1)[0]
    return first.upper() if first else ""


def card_numbers_compatible(requested: str, candidate: str) -> bool:
    """Keep exact numeric cards exact while allowing Beckett set-code trimming.

    Beckett often lists modern autographs as #KB while source listings include
    the checklist prefix (#BCA-KB). Numeric vintage cards do not get that
    flexibility, because #513 silently becoming #355 is a real catalog error.
    """
    requested = re.sub(r"[^A-Za-z0-9]+", "", requested or "").upper()
    candidate = re.sub(r"[^A-Za-z0-9]+", "", candidate or "").upper()
    if not requested:
        return True
    if not candidate:
        return False
    if requested.isdigit() or candidate.isdigit():
        return requested == candidate
    return requested == candidate or requested.endswith(candidate) or candidate.endswith(requested) or candidate.startswith(requested)


def candidate_title_compatible(product: dict[str, Any], candidate: dict[str, str]) -> bool:
    requested_tokens = card_tokens(product.get("name", ""))
    candidate_tokens = card_tokens(candidate.get("title", ""))
    if not requested_tokens:
        return True
    # Require at least one real subject/name/variant token to survive the match.
    # This catches workbook typos such as "Al Oliver #130" matching Beckett's
    # true #130 card, "Denis Menke".
    return bool(requested_tokens & candidate_tokens)


def candidate_variant_compatible(product: dict[str, Any], candidate: dict[str, str]) -> bool:
    """Avoid matching plain refractors to color refractors or vice versa."""
    requested = set(normalize_key(product.get("name", "")).split()) & VARIANT_COLOR_TOKENS
    candidate_colors = set(normalize_key(candidate.get("title", "")).split()) & VARIANT_COLOR_TOKENS
    if requested:
        return requested == candidate_colors
    # If the source only says "Refractors", a color refractor is a different
    # checklist line and should not win just because it shares the player name.
    return not candidate_colors


def enhanced_query_variants(product: dict[str, Any]) -> list[str]:
    """Add Beckett-specific fallback searches for checklist-prefix quirks."""
    title = lookup_title(normalize_spaces(product.get("name", "")))
    variants: list[str] = []

    def add(value: str) -> None:
        value = normalize_spaces(value)
        if value and value not in variants:
            variants.append(value)

    clean = re.sub(r"\([^)]*\)", " ", title)
    clean = re.sub(r"\s*/\s*\d{2,5}\b", " ", clean)
    clean = re.sub(r"\bAU\b", "Autograph", clean, flags=re.I)
    clean = re.sub(r"\bAutograph\b", "Autographs", clean, flags=re.I)
    clean = re.sub(r"\bProspect Autographs\b", "Prospect Autographs", clean, flags=re.I)
    clean = re.sub(r"\bRookie Autographs\b", "Rookie Autographs", clean, flags=re.I)
    add(clean)

    replacements = {
        "#BCP205": "#BCP205B",
        "#BPA-MS": "#MS",
        "#BPAPG": "#PG",
        "#BCA-KB": "#KB",
        "#AID-60": "#AID60",
        "#MM-MP": "#MP",
        "#TTRL-8": "#TTRL8",
    }
    for old, new in replacements.items():
        if old.lower() in clean.lower():
            add(re.sub(re.escape(old), new, clean, flags=re.I))
            add(re.sub(re.escape(old), new.replace("#", ""), clean, flags=re.I))

    if re.search(r"Draft Pick Autographs", clean, flags=re.I):
        add(re.sub(r"Draft Pick Autographs", "Draft Draft Pick Autographs", clean, flags=re.I))
        add(re.sub(r"#BCA-KB", "#KB", re.sub(r"Draft Pick Autographs", "Draft Draft Pick Autographs", clean, flags=re.I), flags=re.I))

    if re.search(r"#BCP205\b", clean, flags=re.I) and re.search(r"\bAutographs?\b|\bAU\b", title, flags=re.I):
        add(re.sub(r"#BCP205\b", "#BCP205B", clean, flags=re.I))

    number = parse_card_number(title)
    year = parse_year(title)
    player = player_from_title(title, year, number)
    set_guess = beckett.set_guess(product)
    if year and set_guess and number and player:
        add(f"{year} {set_guess} {number} {player}")
        for old, new in replacements.items():
            if number.upper() == old.replace("#", "").upper():
                add(f"{year} {set_guess} {new.replace('#', '')} {player}")
        if number.upper() == "BCP205":
            add(f"{year} {set_guess} BCP205B {player}")
    return variants


def candidate_year_compatible(product: dict[str, Any], candidate: dict[str, str]) -> bool:
    requested_year = str(product.get("year") or parse_year(product.get("name", "")))
    if not requested_year:
        return True
    candidate_years = beckett.YEAR_RE.findall(f"{candidate.get('title', '')} {candidate.get('url', '')}")
    if not candidate_years:
        return True
    return any(str(year).startswith(requested_year) for year in candidate_years)


def best_compatible_candidate(session: Any, product: dict[str, Any], cache: dict[str, Any]) -> beckett.SearchCandidate | None:
    """Search Beckett but reject mismatched card numbers before scoring."""
    override_url = MANUAL_BECKETT_OVERRIDES.get(normalize_key(product.get("name", "")))
    if override_url:
        page = beckett.fetch_card_page(session, override_url, cache, 0.03)
        return beckett.SearchCandidate(
            title=page.get("title") or normalize_spaces(product.get("name")),
            url=override_url,
            query="manual verified override",
            score=1.0,
        )

    requested_number = parse_card_number(product.get("name", ""))
    best: beckett.SearchCandidate | None = None
    queries: list[str] = []
    for query in [*beckett.query_variants(product), *enhanced_query_variants(product)]:
        if query not in queries:
            queries.append(query)
    for query in queries:
        try:
            candidates = beckett.search_beckett(session, query, cache, 0.03)
        except Exception as exc:
            cache.setdefault("search_errors", {})[query] = str(exc)
            continue
        for candidate in candidates:
            if not candidate_year_compatible(product, candidate):
                continue
            if not candidate_title_compatible(product, candidate):
                continue
            if not candidate_variant_compatible(product, candidate):
                continue
            if requested_number and not card_numbers_compatible(requested_number, candidate_card_number(candidate)):
                continue
            score = beckett.score_candidate(product, query, candidate)
            if best is None or score > best.score:
                best = beckett.SearchCandidate(
                    title=candidate["title"],
                    url=candidate["url"],
                    query=query,
                    score=score,
                )
        if best and best.score >= 0.86:
            break
    return best


def copy_asset(source: Path, subfolder: str) -> str:
    target_dir = ROOT / "assets" / "baseball-cards" / subfolder
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / safe_asset_name(source)
    if not target.exists() or target.stat().st_size != source.stat().st_size:
        shutil.copy2(source, target)
    return target.relative_to(ROOT).as_posix()


def write_thumbnail(asset_path: str) -> str | None:
    source = ROOT / asset_path
    if not source.exists() or source.suffix.lower() == ".svg":
        return None
    relative = Path(asset_path)
    if relative.parts and relative.parts[0].lower() == "assets":
        relative = Path(*relative.parts[1:])
    thumb = ROOT / "assets" / "thumbnails" / relative.with_suffix(".webp")
    thumb.parent.mkdir(parents=True, exist_ok=True)
    try:
        with Image.open(source) as image:
            image = image.convert("RGB")
            image.thumbnail((640, 640), Image.Resampling.LANCZOS)
            image.save(thumb, "WEBP", quality=78, method=6)
        return thumb.relative_to(ROOT).as_posix()
    except Exception:
        return None


def removal_base_title(value: str) -> tuple[str, int | None]:
    text = normalize_spaces(value)
    limit = 1 if re.search(r"remove\s+1|only\s+remove\s+1", text, flags=re.I) else None
    text = re.sub(r"\([^)]*remove[^)]*\)", "", text, flags=re.I)
    return normalize_spaces(text), limit


def removal_score(query: str, product: dict[str, Any]) -> float:
    title = str(product.get("name") or "")
    query_key = normalize_key(query)
    product_key = normalize_key(title)
    if not query_key or not product_key:
        return 0.0
    score = SequenceMatcher(None, query_key, product_key).ratio()
    query_tokens = set(query_key.split())
    product_tokens = set(product_key.split())
    score += len(query_tokens & product_tokens) / max(1, len(query_tokens))
    query_year = parse_year(query)
    query_number = parse_card_number(query)
    product_year = str(product.get("year") or parse_year(title))
    product_number = parse_card_number(title)
    if query_year:
        score += 0.25 if product_year == query_year else -0.25
    if query_number:
        score += 0.35 if product_number == query_number else -0.35
    query_grade, _ = parse_grade(query)
    product_grade, _ = parse_grade(product.get("condition") or title)
    if query_grade:
        score += 0.25 if query_grade == product_grade else -0.15
    return score


def remove_products(products: list[dict[str, Any]], remove_rows: list[dict[str, str]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    remaining = list(products)
    removed: list[dict[str, Any]] = []
    unresolved: list[dict[str, Any]] = []

    for row in remove_rows:
        query, limit = removal_base_title(row["title"])
        candidates = sorted(
            ((removal_score(query, product), product) for product in remaining),
            key=lambda item: item[0],
            reverse=True,
        )
        candidates = [item for item in candidates if item[0] >= 1.25]
        if not candidates:
            unresolved.append({**row, "normalizedTitle": query, "reason": "no safe matching product"})
            continue

        if limit == 1:
            generic_first = sorted(
                candidates,
                key=lambda item: (
                    0 if "guide range listed" in str(item[1].get("condition", "")).lower() else 1,
                    -item[0],
                ),
            )
            selected = [generic_first[0]]
        else:
            # Keep removals narrow: most workbook rows are one specific card.
            selected = [candidates[0]]

        selected_ids = {item[1]["id"] for item in selected}
        remaining = [product for product in remaining if product["id"] not in selected_ids]
        for score, product in selected:
            removed.append(
                {
                    "request": row["title"],
                    "productId": product["id"],
                    "name": product.get("name"),
                    "score": round(score, 4),
                    "sheet": row["sheet"],
                    "row": row["row"],
                }
            )
    return remaining, removed, unresolved


def player_from_title(title: str, year: str, card_number: str) -> str:
    override = PLAYER_OVERRIDES.get((year, card_number))
    if override:
        return override
    text = re.sub(r"\b(19\d{2}|20\d{2})(?:[-/]\d{2})?\b", " ", title)
    text = re.sub(r"#\s*[A-Za-z0-9]+", " ", text)
    text = re.sub(r"\b(Topps|Kelloggs?|Donruss|Baseball's|Best|Bronze|Rookies|UPD)\b", " ", text, flags=re.I)
    text = re.sub(r"\b(PSA|BGS|SGC|CGC|GEM|MINT|NM|EX|VG|MC|Rookie|Card|All-Star|AS)\b", " ", text, flags=re.I)
    text = re.sub(r"\([^)]*\)", " ", text)
    text = normalize_spaces(re.sub(r"[,/]+", " | ", text))
    text = re.sub(r"\s*\|\s*$", "", text)
    return text or ""


def source_page_for(year: str) -> str:
    if year.startswith("19"):
        decade = f"{year[:3]}0s"
    elif year.startswith("20"):
        decade = f"{year[:3]}0s"
    else:
        decade = "Legacy"
    return f"Baseball {decade}"


def condition_from_title(raw_title: str, price_basis: str) -> str:
    grade_label, _ = parse_grade(raw_title)
    if grade_label:
        detail = re.sub(r".*?\b" + re.escape(grade_label) + r"\b", "", raw_title, flags=re.I)
        detail = normalize_spaces(re.sub(r"[,()]+", " ", detail))
        useful = [token for token in detail.split() if token.upper() in {"GOOD", "VG", "EX", "NM", "MT", "MINT", "GEM", "PR"} or "-" in token]
        return normalize_spaces(f"{grade_label} {' '.join(useful)}")
    if AUTOGRAPH_RE.search(raw_title):
        return "Ungraded | On-card autograph"
    return "Ungraded | Guide range listed"


def details_from_beckett(raw_title: str, match: beckett.SearchCandidate, page: dict[str, Any], session: Any, cache: dict[str, Any]) -> dict[str, Any]:
    raw_prices = page.get("raw_prices") or {}
    graded_prices = page.get("graded_prices") or {}
    title = normalize_spaces(page.get("title") or match.title or raw_title)
    grade_label, grade_value = parse_grade(raw_title)
    condition = condition_from_title(raw_title, "")
    realtime = fetch_realtime_pricing(match.url, session, cache, 0.03)

    if grade_label:
        price_label, basis, note = beckett.matched_grade_price(condition, raw_prices, graded_prices)
        price_label = normalize_price_label(price_label)
        pricing_source = f"Beckett grade basis: {basis}" if basis else "Beckett graded pricing"
        if not price_label:
            market_range = (page.get("graded_market") or {}).get("range") or ""
            price_label = normalize_price_label(market_range)
            pricing_source = "Beckett graded market range fallback"
        if not price_label and normalize_price_label(realtime):
            price_label = normalize_price_label(realtime)
            pricing_source = "Beckett Real Time Pricing fallback; graded bucket unavailable"
        display_name = title if grade_label.lower() in title.lower() else f"{title} {grade_label}"
    else:
        price_label = normalize_price_label(realtime)
        pricing_source = "Beckett Real Time Pricing"
        note = ""
        if not price_label:
            market_range = (page.get("raw_market") or {}).get("range") or ""
            price_label = normalize_price_label(market_range or beckett.price_range(raw_prices))
            pricing_source = "Beckett raw fallback; real-time pricing unavailable"
        display_name = title

    return {
        "beckettTitle": title,
        "displayName": display_name,
        "priceLabel": price_label,
        "price": midpoint_from_label(price_label) or 0,
        "condition": condition,
        "gradeLabel": grade_label,
        "gradeValue": grade_value,
        "realtimePricing": normalize_price_label(realtime),
        "pricingSource": pricing_source,
        "pricingNote": note,
        "rawPriceRange": beckett.price_range(raw_prices),
        "gradedPriceRange": beckett.price_range(graded_prices),
    }


def add_products(products: list[dict[str, Any]], add_rows: list[dict[str, str]], session: Any, cache: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    files = image_files()
    next_id = max(product["id"] for product in products) + 1
    next_rank = max(int(product.get("sortRank") or 0) for product in products) + 1
    added: list[dict[str, Any]] = []
    unresolved: list[dict[str, Any]] = []
    existing_raw_titles = {
        normalize_key((product.get("metadata") or {}).get("cardsAddRawTitle"))
        for product in products
        if isinstance(product.get("metadata"), dict)
    }

    products_by_base: dict[str, dict[str, Any]] = {}
    for product in products:
        raw = ""
        metadata = product.get("metadata")
        if isinstance(metadata, dict):
            raw = normalize_spaces(metadata.get("cardsAddRawTitle"))
        key = normalize_key(base_copy_title(raw or product.get("name", "")))
        if key and key not in products_by_base:
            products_by_base[key] = product

    def add_cloned_copy(row: dict[str, str], reason: str) -> dict[str, Any] | None:
        nonlocal next_id, next_rank
        raw_title = row["title"]
        base_key = normalize_key(base_copy_title(raw_title))
        source = products_by_base.get(base_key)
        if not source:
            return None
        product = deepcopy(source)
        product["id"] = next_id
        product["sortRank"] = next_rank
        product["copyCount"] = int(product.get("copyCount") or 1)
        metadata = deepcopy(product.get("metadata") or {})
        metadata["cardsAddRawTitle"] = raw_title
        metadata["copiedFromProductId"] = source.get("id")
        metadata["copyReason"] = reason
        product["metadata"] = metadata
        image_source, image_score, image_reason = best_image_for_add(raw_title, files)
        if image_source:
            image_path = copy_asset(image_source, "legacy-additions")
            product["image"] = image_path
            product["imageGallery"] = [image_path]
            thumb = write_thumbnail(image_path)
            if thumb:
                product["metadata"]["thumbnailPath"] = thumb
            product["metadata"]["imageSource"] = str(image_source)
            product["metadata"]["imageMatchScore"] = round(image_score, 3)
            product["metadata"]["needsPhotoReview"] = False
        else:
            product["metadata"]["needsPhotoReview"] = True
            product["metadata"]["imageReason"] = image_reason
        products.append(product)
        products_by_base[base_key] = product
        added.append(
            {
                "request": raw_title,
                "productId": next_id,
                "name": product.get("name"),
                "priceLabel": product.get("priceLabel"),
                "beckettUrl": (product.get("metadata") or {}).get("beckettUrl") or product.get("htmlFullLink"),
                "beckettScore": (product.get("metadata") or {}).get("beckettMatchScore"),
                "image": product.get("image"),
                "imageSource": (product.get("metadata") or {}).get("imageSource", ""),
                "needsPhotoReview": bool((product.get("metadata") or {}).get("needsPhotoReview")),
                "imageReason": image_reason,
                "clonedFromProductId": source.get("id"),
            }
        )
        next_id += 1
        next_rank += 1
        return product

    for row in add_rows:
        raw_title = row["title"]
        search_title = lookup_title(raw_title)
        if normalize_key(raw_title) in existing_raw_titles:
            continue

        year = parse_year(search_title)
        card_number = parse_card_number(search_title)
        player = player_from_title(search_title, year, card_number)
        product_stub = {
            "name": search_title,
            "category": "Baseball",
            "sport": "Baseball",
            "playerAthlete": player,
            "year": int(year) if year else None,
        }
        match = best_compatible_candidate(session, product_stub, cache)
        if not match or match.score < 0.65:
            if add_cloned_copy(row, "No fresh Beckett match; duplicate inherited base card details."):
                continue
            unresolved.append({**row, "reason": "No confident Beckett match", "score": match.score if match else None})
            continue
        page = beckett.fetch_card_page(session, match.url, cache, 0.03)
        details = details_from_beckett(search_title, match, page, session, cache)
        if not details["priceLabel"]:
            if add_cloned_copy(row, "No fresh Beckett price; duplicate inherited base card details."):
                continue
            unresolved.append({**row, "reason": "No Beckett price found for matched card", "beckettUrl": match.url, "score": match.score})
            continue
        image_source, image_score, image_reason = best_image_for_add(search_title, files)
        image_path = copy_asset(image_source, "legacy-additions") if image_source else "assets/placeholder-baseball.svg"
        thumbnail_path = write_thumbnail(image_path)

        attributes: list[str] = []
        if AUTOGRAPH_RE.search(raw_title) or AUTOGRAPH_RE.search(details["displayName"]):
            attributes.append("Autograph")
        if re.search(r"\b(RC|Rookie)\b", raw_title, flags=re.I):
            attributes.append("Rookie")
        if re.search(r"\b(jersey|patch|relic|memorabilia|game[- ]used|game[- ]worn|materials?)\b", raw_title, flags=re.I):
            attributes.append("Memorabilia")
        if re.search(r"/\s*\d{2,5}\b", raw_title):
            attributes.append("Serial Numbered")

        team = TEAM_OVERRIDES.get((year, card_number), "")
        description = (
            f"Legacy baseball listing matched to Beckett as {details['beckettTitle']}. "
            "Please review the photos for the exact card you will receive."
        )
        if not image_source:
            description += " A matching product photo still needs to be added."

        product = {
            "id": next_id,
            "name": details["displayName"],
            "category": "Baseball",
            "team": team,
            "year": int(year) if year else None,
            "condition": details["condition"],
            "price": details["price"],
            "priceLabel": details["priceLabel"],
            "image": image_path,
            "imageGallery": [image_path],
            "description": description,
            "photoHostPageUrl": match.url,
            "legacyImageLabel": Path(image_path).stem if image_path else "",
            "sourcePage": source_page_for(year),
            "league": "MLB",
            "sport": "Baseball",
            "playerAthlete": player,
            "displayPrice": details["priceLabel"],
            "copyCount": 1,
            "itemPhotoUrl": "",
            "itemPhotoUrls": [],
            "htmlFullLink": match.url,
            "htmlImageUrls": [],
            "metadata": {
                "cardsAddRawTitle": raw_title,
                "cardsAddWorkbook": WORKBOOK_PATH.name,
                "beckettUrl": match.url,
                "beckettTitle": details["beckettTitle"],
                "beckettMatchScore": match.score,
                "beckettQuery": match.query,
                "beckettRealTimePricing": details["realtimePricing"],
                "beckettRawPriceRange": details["rawPriceRange"],
                "beckettGradedPriceRange": details["gradedPriceRange"],
                "pricingSource": details["pricingSource"],
                "pricingNote": details["pricingNote"],
                "imageSource": str(image_source) if image_source else "",
                "imageMatchScore": round(image_score, 3),
                "thumbnailPath": thumbnail_path or "",
                "needsPhotoReview": not bool(image_source),
            },
            "isFeatured": False,
            "isDeleted": False,
            "sortRank": next_rank,
        }
        if attributes:
            product["attributes"] = attributes

        products.append(product)
        products_by_base[normalize_key(base_copy_title(raw_title))] = product
        added.append(
            {
                "request": raw_title,
                "productId": next_id,
                "name": product["name"],
                "priceLabel": product["priceLabel"],
                "beckettUrl": match.url,
                "beckettScore": match.score,
                "image": image_path,
                "imageSource": str(image_source) if image_source else "",
                "needsPhotoReview": not bool(image_source),
                "imageReason": image_reason,
            }
        )
        next_id += 1
        next_rank += 1
    return products, added, unresolved


def replace_photos(products: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    replaced: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    files = [file for file in REPLACE_PHOTO_DIR.iterdir() if file.is_file() and file.suffix.lower() in IMAGE_EXTENSIONS]

    for file in files:
        year = parse_year(file.stem)
        number = parse_card_number(file.stem)
        file_grade, file_grade_value = parse_grade(file.stem)
        file_tokens = card_tokens(file.stem)
        candidates: list[tuple[float, dict[str, Any]]] = []
        for product in products:
            if product.get("category") != "Baseball":
                continue
            product_title = str(product.get("name") or "")
            product_year = str(product.get("year") or parse_year(product_title))
            product_number = parse_card_number(product_title)
            product_grade, product_grade_value = parse_grade(f"{product.get('condition')} {product_title}")
            if year and product_year != year:
                continue
            if number and product_number != number:
                continue
            if file_grade_value is not None and product_grade_value is not None and file_grade_value != product_grade_value:
                continue
            product_tokens = card_tokens(product_title)
            overlap = len(file_tokens & product_tokens)
            score = overlap + (3 if file_grade and file_grade == product_grade else 0)
            candidates.append((score, product))
        candidates.sort(key=lambda item: item[0], reverse=True)
        if not candidates or candidates[0][0] < 2:
            skipped.append({"file": str(file), "reason": "No exact year/card/grade product match"})
            continue
        product = candidates[0][1]
        old_image = product.get("image")
        new_path = copy_asset(file, "legacy-replacements")
        thumb = write_thumbnail(new_path)
        gallery = [new_path]
        for item in product.get("imageGallery") or []:
            if item and item not in {old_image, new_path}:
                gallery.append(item)
        product["image"] = new_path
        product["imageGallery"] = gallery
        metadata = deepcopy(product.get("metadata") or {})
        metadata["replacementPhotoSource"] = str(file)
        metadata["replacementPhotoAppliedAt"] = datetime.now().isoformat(timespec="seconds")
        if thumb:
            metadata["thumbnailPath"] = thumb
        product["metadata"] = metadata
        replaced.append({"productId": product["id"], "name": product.get("name"), "oldImage": old_image, "newImage": new_path})
    return replaced, skipped


def audit_products(products: list[dict[str, Any]], changed_ids: set[int]) -> dict[str, Any]:
    ids = [product["id"] for product in products]
    duplicate_ids = sorted({pid for pid in ids if ids.count(pid) > 1})
    missing_images = [
        {"id": product["id"], "name": product.get("name"), "image": product.get("image")}
        for product in products
        if product.get("image") and not (ROOT / product["image"]).exists() and not str(product["image"]).startswith("http")
    ]
    changed = [product for product in products if product["id"] in changed_ids]
    changed_missing_prices = [
        {"id": product["id"], "name": product.get("name")}
        for product in changed
        if not product.get("priceLabel")
    ]
    placeholder_changed = [
        {"id": product["id"], "name": product.get("name"), "image": product.get("image")}
        for product in changed
        if "placeholder" in str(product.get("image", "")).lower()
    ]
    return {
        "productCount": len(products),
        "duplicateIds": duplicate_ids,
        "missingImages": missing_images[:50],
        "missingImageCount": len(missing_images),
        "changedMissingPrices": changed_missing_prices,
        "changedPlaceholderImages": placeholder_changed,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="Write catalog changes.")
    parser.add_argument("--skip-beckett", action="store_true", help="Do not login to Beckett; useful for dry local checks.")
    args = parser.parse_args()

    adds, removes = workbook_rows()
    products = json.loads(PRODUCTS_PATH.read_text(encoding="utf-8"))
    original_count = len(products)

    products, removed, unresolved_removals = remove_products(products, removes)
    replaced, skipped_replacements = replace_photos(products)

    added: list[dict[str, Any]] = []
    unresolved_adds: list[dict[str, Any]] = []
    cache = beckett.load_cache(CACHE_PATH)
    session = None
    if not args.skip_beckett:
        email = os.environ.get("BECKETT_EMAIL")
        password = os.environ.get("BECKETT_PASSWORD")
        if not email or not password:
            raise SystemExit("Set BECKETT_EMAIL and BECKETT_PASSWORD to fetch Beckett details.")
        session = beckett.create_session(email, password)
        products, added, unresolved_adds = add_products(products, adds, session, cache)
        beckett.save_cache(CACHE_PATH, cache)

    changed_ids = {entry["productId"] for entry in removed + added + replaced if entry.get("productId") is not None}
    report = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "workbook": str(WORKBOOK_PATH),
        "applied": bool(args.apply),
        "originalProductCount": original_count,
        "finalProductCount": len(products),
        "requestedAdds": len(adds),
        "requestedRemovals": len(removes),
        "removed": removed,
        "unresolvedRemovals": unresolved_removals,
        "added": added,
        "unresolvedAdds": unresolved_adds,
        "photoReplacements": replaced,
        "skippedPhotoReplacements": skipped_replacements,
        "audit": audit_products(products, changed_ids),
    }

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    if args.apply:
        rebuild_product_files(products)
    print(json.dumps({k: report[k] for k in ["applied", "originalProductCount", "finalProductCount", "requestedAdds", "requestedRemovals"]}, indent=2))
    print(f"removed={len(removed)} added={len(added)} replaced={len(replaced)} unresolved_adds={len(unresolved_adds)} unresolved_removals={len(unresolved_removals)}")
    print(f"report={REPORT_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
