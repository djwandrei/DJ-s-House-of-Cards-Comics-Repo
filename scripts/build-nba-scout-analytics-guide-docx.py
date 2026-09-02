"""Build the NBA Scout Analytics Coaching Guide DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "NBA-Scout-Analytics-Coaching-Guide.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "59636F"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in CELL_MARGIN_DXA.items():
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, size=11, color="000000", bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    add_text(paragraph, text)
    return paragraph


def add_number(doc, number, text):
    # Use explicit numbering so each coaching workflow can restart at one;
    # Word's shared List Number style otherwise continues across sections.
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    add_text(paragraph, f"{number}.  {text}")
    return paragraph


def add_paragraph(doc, text, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    add_text(paragraph, text, italic=italic)
    return paragraph


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    add_text(paragraph, text, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
    return paragraph


def add_callout(doc, title, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    set_paragraph_shading(paragraph, LIGHT_GRAY)
    add_text(paragraph, f"{title}: ", size=10.5, color=NAVY, bold=True)
    add_text(paragraph, text, size=10.5, color=NAVY)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    header_cells = table.rows[0].cells
    for cell, text in zip(header_cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        add_text(paragraph, text, size=9.5, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, row)):
            if index == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_text(paragraph, text, size=9.25, color="202124", bold=index == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, color, before, after in [
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_page(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(header_p, "NBA SCOUT ANALYTICS", size=8.5, color=MUTED, bold=True)
    add_text(header_p, "  |  Coaching reference", size=8.5, color=MUTED)
    footer = section.footer
    footer_p = footer.paragraphs[0]
    # Keep the footer to a single page field. Word's PDF exporter can clip
    # adjacent text on some pages when a PAGE field shares the paragraph.
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer_p._p.append(field)


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(4)
    add_text(title, "NBA Scout Analytics", size=26, color=NAVY, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    add_text(subtitle, "A coaching reference for lineup decisions, scouting, player development, and future tools", size=13, color=MUTED)
    add_callout(doc, "Operating principle", "Use this package to make a better coaching question, not to make an unsupported conclusion. Every result should retain its sample, evidence level, and context.")

    add_heading(doc, "Purpose", 1)
    add_paragraph(doc, "This guide explains the local NBA Scout analytics package built from the licensed Sportradar play-by-play archive. It is designed for future Lineup Lab, scouting, roster-construction, game-prep, and player-development tools. The package distinguishes direct provider facts, reconstructed lineup facts, regularized estimates, and transparent proxies.")

    add_heading(doc, "Data scope and evidence levels", 1)
    add_paragraph(doc, "The current package covers the selected 2025-26 official NBA phases in the local archive: regular season, in-season tournament, play-in, and playoffs. It uses non-rescinded structured play-by-play statistics and a current-version reconstruction of lineups, stints, and possessions. Derivation runs offline with no network request or Supabase write.")
    add_table(doc, ["Evidence level", "Meaning", "Examples and correct use"], [
        ["Direct event fact", "Structured provider event reports it.", "Shot result, assist, rebound, steal, block, foul, player box score. Describe observed production and style."],
        ["Reconstructed lineup fact", "Verified five-player lineup is attached to a possession or stint.", "Exact five-man rating, player on/off exposure, rotation minutes. Evaluate combinations with sample context."],
        ["Regularized model", "Ridge regression estimates adjusted contribution.", "Net/offense/defense RAPM. Compare players after context adjustment; retain reliability."],
        ["Proxy or qualifier absence", "Clear rule approximates a concept not directly labeled in the archive.", "Garbage-time proxy, leverage proxy, no-provider-fastbreak context. Use to frame film questions, not as ground truth."],
    ], [1700, 2550, 5110])

    add_heading(doc, "The core coaching dashboard", 1)
    add_number(doc, 1, "Show the sample: games, possessions, minutes, reliability grade, and chosen phase/context.")
    add_number(doc, 2, "Show the unit: exact five-player lineup versus shared-floor pair, trio, or quartet.")
    add_number(doc, 3, "Show the output: ratings, four factors, shot profile, and possession outcomes.")
    add_number(doc, 4, "Show the next action: a film question, rotation option, matchup hypothesis, or development action.")

    add_heading(doc, "Team and lineup performance", 1)
    add_heading(doc, "Ratings and scoring margin", 2)
    for text in [
        "Offensive rating (ORtg): points scored per 100 offensive possessions; compares scoring efficiency across uneven minutes or games.",
        "Defensive rating (DRtg): points allowed per 100 defensive possessions; lower is better.",
        "Net rating: ORtg minus DRtg; a unit's possession-level scoring margin.",
        "Plus-minus per 100: score differential scaled by all offense and defense possessions; a descriptive companion to net rating.",
        "Reliability and 95% intervals: sample-stability aids, not proof that a player or lineup caused the result.",
    ]:
        add_bullet(doc, text)
    add_paragraph(doc, "Use ratings to compare a candidate unit's all-sample output with its clutch, venue, period, and recent-window output. Small samples should produce a test-on-film recommendation rather than a hard rotation rule.")

    add_heading(doc, "Exact lineups and co-presence groups", 2)
    for text in [
        "Five-player lineup: the exact five players verified at possession start. This is the strongest lineup record in the package.",
        "Two-, three-, and four-player combination: a shared-floor co-presence group. It answers whether those players played together; it is not a complete lineup.",
        "Exposure: games, team possessions, minutes, possessions per game, minutes per game, and pace per 48 minutes.",
        "Continuity: team possession share, team minute share, games used, and for exact five-player lineups only, starting and closing lineup counts/rates.",
    ]:
        add_bullet(doc, text)
    add_paragraph(doc, "Use five-man rows to identify units that deserve more testing, protection against a certain opponent style, or caution because the result is based on a thin sample. Use pair/trio rows to explore compatibility before choosing a full five-man unit.")

    add_heading(doc, "On/off and WOWY", 2)
    add_paragraph(doc, "Player on/off compares team performance in a player's on-court and same-game off-court possessions. This reduces schedule mix but does not remove role, lineup, opponent, or score-state selection effects. With-or-without-you (WOWY) partitions a pair into both-on, A-on/B-off, A-off/B-on, and both-off states.")
    add_callout(doc, "Example", "Use on/off to ask whether the team's turnover rate rises when a primary handler sits. Use WOWY to investigate whether two players are complementary, redundant, or simply deployed in different contexts. Pair both with minutes, lineup mix, and film.")

    add_heading(doc, "Adjusted player value", 1)
    add_heading(doc, "Regularized adjusted plus-minus", 2)
    for text in [
        "Net RAPM per 100: regularized adjusted net contribution.",
        "Offensive RAPM per 100: estimated offensive contribution relative to the fitted baseline.",
        "Defensive RAPM per 100: estimated defensive contribution; positive means better defense in this package.",
        "Combined O/D RAPM: the combined adjusted estimate.",
        "Venue control: O/D RAPM fits a signed home-court term separately and reports player-level home/away exposure balance. Strong imbalance remains a roster-venue confounding warning, not causal proof.",
        "Teammate and opponent context: exposure-weighted averages shown as context, not added a second time as an adjustment.",
        "Reliability: ridge/exposure proxies and sample tiers; these are not confidence intervals.",
    ]:
        add_bullet(doc, text)
    add_paragraph(doc, "RAPM is best for choosing players who merit more lineup experiments, estimating unseen five-man groups, and avoiding raw-plus-minus traps. It is not a final player ranking: ridge regularization shrinks noisy samples and lineup deployment is not random.")
    add_heading(doc, "Lineup projection", 2)
    add_paragraph(doc, "For an unseen five-player lineup, the neutral projection starts with the five net RAPM values. For an observed exact lineup, it adds a possession-shrunk residual synergy after opponent and home-court exposure adjustment. The home-court adjustment uses the net RAPM model's signed home-versus-away term, weighted by the lineup's actual home/away possession exposure, so it stays aligned with the fitted player values. Use it to prioritize controlled lineup trials, with observed net rating, possessions, synergy weight, and caveats visible beside it.")

    add_heading(doc, "Possession quality and style", 1)
    add_heading(doc, "Four factors and shooting profile", 2)
    add_table(doc, ["Metric", "Formula / definition", "Coaching use"], [
        ["eFG%", "(FGM + 0.5 x 3PM) / FGA", "Explain whether a rating gap came from shot value and conversion."],
        ["Turnover rate", "TOV / (FGA + 0.44 x FTA + TOV)", "Identify possession security and pressure response."],
        ["Offensive rebound %", "ORB / (ORB + opponent DRB)", "Evaluate second-shot access where rebound evidence is structured."],
        ["Free-throw attempt rate", "FTA / FGA", "Capture rim pressure and foul drawing at the unit level."],
        ["True shooting", "Points / (2 x (FGA + 0.44 x FTA))", "Compare scoring efficiency across shot and free-throw mixes."],
        ["Shot zones", "At rim 0-4 ft; short mid 5-14; long mid 15+", "Reveal location profile; unknown distances stay out of zone counts."],
    ], [1600, 3050, 4710])
    add_paragraph(doc, "Show the four-factor coverage status with every view. Missing event structure or outcome, nonparticipant attribution, and role-inconsistent team attribution make a sample partial rather than silently treating the play as zero.")
    add_heading(doc, "Playmaking, disruption, and possession extensions", 2)
    for text in [
        "Playmaking: assists, assists per 100, assisted-FG rate, assist-to-turnover ratio, fouls drawn, and fouls drawn per 100.",
        "Defensive disruption: steals, blocks, personal fouls, per-100 rates, block rate, and steal-forced-turnover rate.",
        "Second chance: only flags a possession when the structured event slice contains an offensive rebound. It reports possessions, points, rate, points per 100, and points per second-chance possession.",
        "Points off turnovers: only flags when the possession opening event contains a structured opponent turnover. It reports possession count, points, rate, points per 100, and points per possession after a turnover.",
        "Possession outcomes: counts and rates for 0, 1, 2, 3, and 4+ points.",
    ]:
        add_bullet(doc, text)
    add_paragraph(doc, "These measures support transition-defense preparation, offensive-glass priorities, pressure packages, end-of-quarter play calls, and player development. Absence of a structured event is never treated as a made-up tactical label.")

    add_heading(doc, "Game context and situation splits", 1)
    add_table(doc, ["Split", "Definition", "Practical question"], [
        ["Phase", "Regular, in-season tournament, play-in, or playoffs.", "Does this rotation travel to higher-stakes games?"],
        ["Venue", "Home or away team perspective.", "Is the performance portable?"],
        ["Period / half", "Q1-Q4, overtime; first half, second half, overtime.", "Which units start, close, or stabilize third-quarter runs?"],
        ["Rolling window", "Last 5, 10, and 20 eligible games at the archive snapshot.", "Is current form different from season-long output?"],
        ["Clutch", "Final five minutes of Q4/OT, margin five or fewer; missing inputs stay unclassified.", "What is the trusted late-game lineup and style?"],
        ["Transition", "Provider fast-break qualifier, no qualifier, or unclassified.", "How does the unit perform in confirmed fast break?"],
        ["Score / leverage", "Score-state bands plus transparent score-clock proxy buckets.", "Which units have been tested when possessions matter most?"],
    ], [1450, 3550, 4360])
    add_callout(doc, "Important caveat", "No-provider-fastbreak is an absence-of-qualifier proxy. It can narrow film review, but it must never be relabeled as verified half-court offense.")

    add_heading(doc, "Player production and rotation role cards", 1)
    add_paragraph(doc, "The direct player profile joins structured event totals to reconstructed on-court exposure for rostered players with a same-game on/off row. Coverage fields report unresolved shot outcomes, shot values, free-throw outcomes, and rebound types. Dependent percentages and rates remain unavailable instead of treating unknown shots as misses or unknown shot values as two-pointers.")
    for text in [
        "Box-score line: points; FGA/FGM; 2PA/2PM; 3PA/3PM; FTA/FTM; rebounds; assists; steals; blocks; turnovers; personal/drawn/technical/flagrant fouls; blocked attempts; ejections.",
        "Shooting line: FG%, 2P%, 3P%, FT%, eFG%, TS%, 3PA rate, blocked-attempt rate, observed distance, two-point zones, and provider shot-type/description profiles such as jumpers, layups, dunks, hooks, pull-ups, drives, step-backs, cuts, floaters, putbacks, and fadeaways.",
        "Rate line: per-36 and per-100 possession scoring, rebounding, playmaking, disruption, foul, and blocked-attempt indicators plus possession-ending involvement proxy.",
        "Role line: games appeared, reconstructed minutes, team possessions while on court, starter games/rate, and closer games/rate.",
    ]:
        add_bullet(doc, text)
    add_paragraph(doc, "Use player cards to build development plans. A low-rim-frequency, low-foul-drawn scorer may merit rim-pressure work; a high-assist but high-turnover player may merit decision-quality film; a strong defensive RAPM player with a modest box score may need matchup-focused film instead of a scoring-centric evaluation.")

    add_heading(doc, "Recommended future tool patterns", 1)
    for title_text, steps in [
        ("Rotation planner", [
            "Filter to an eligible phase and meaningful minimum possession threshold.",
            "Rank exact five-man groups by projected and observed net rating, with reliability visible.",
            "Compare continuity, starter/closer usage, availability, and role assumptions.",
            "Stress-test against clutch, venue, period, and opponent-context views.",
            "Produce a proposed 48-minute rotation as a hypothesis, then review film and coach constraints.",
        ]),
        ("Opponent scout", [
            "Start with team last-10 and last-20 views against season-wide output.",
            "Review eFG%, 3PA rate, free-throw rate, turnover rate, offensive rebound rate, and possession outcomes.",
            "Split by periods, clutch, score state, confirmed fast break, and venue.",
            "Identify the most-used starting and closing exact lineups.",
            "Use co-presence and player cards to frame matchup questions for film.",
        ]),
        ("Player development and recruitment board", [
            "Start with direct player production and rate statistics.",
            "Add on/off and RAPM only with reliability and exposure displayed.",
            "Compare role usage, lineup partners, and opponent context.",
            "Turn outliers into qualitative review prompts rather than automated evaluations.",
            "Keep a separate scout-note layer for video, medical, contract, and character data.",
        ]),
    ]:
        add_heading(doc, title_text, 2)
        for step_number, step in enumerate(steps, start=1):
            add_number(doc, step_number, step)

    add_heading(doc, "Future product architecture", 1)
    for text in [
        "Query one team shard at a time. Team JSON and gzip shards avoid loading the entire season into a browser.",
        "Keep raw archive and derived package separate. Raw play-by-play is provenance; compact derived output is the product input.",
        "Store source version, reconstruction version, metrics version, filters, and generated timestamp with every downstream snapshot.",
        "Use stable provider team/player IDs internally and resolve display names in the presentation layer.",
        "Preserve nulls and coverage fields. Never replace unavailable metrics with zero.",
        "Treat model estimates, proxies, and direct facts as different data types in UI copy and color treatment.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "What this package does not provide", 1)
    add_paragraph(doc, "The archive does not provide verified defender matchups, screen/action labels, player or ball tracking, contest distance, expected shot quality, video clips, injuries, medical data, contracts, or a true win-probability model. Those features require separately licensed tracking, video, medical, or roster sources and a documented join strategy.")

    add_heading(doc, "Release and interpretation checklist", 1)
    for text in [
        "Confirm season, phase, and rolling-window cutoff.",
        "Confirm whether the row is an exact lineup, co-presence group, direct event statistic, model estimate, or proxy.",
        "Show possessions, minutes, games, and reliability before ranking.",
        "Preserve source caveats in all dashboards and exports.",
        "Use player and lineup outputs to prioritize coaching questions, film review, and controlled experiments.",
        "Do not make medical, contractual, or character conclusions from basketball play-by-play analytics.",
    ]:
        add_bullet(doc, text)

    doc.core_properties.title = "NBA Scout Analytics Coaching Guide"
    doc.core_properties.subject = "Metric definitions, coaching uses, and future tool design"
    doc.core_properties.author = "DJ's House of Cards & Comics"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
