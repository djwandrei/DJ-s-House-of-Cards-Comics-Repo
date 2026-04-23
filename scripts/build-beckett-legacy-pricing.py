import json
import os
import re
import time
from collections import Counter
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import quote_plus, urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
PRODUCTS_PATH = ROOT / "products.json"
OUT_DIR = Path.home() / "Documents" / "eBay Docs" / "Listing Automation"
OUT_PATH = OUT_DIR / "Beckett Legacy Pricing Review.xlsx"
CACHE_PATH = OUT_DIR / "beckett_legacy_pricing_cache.json"

SPORT_IDS = {
    "Baseball": "185223",
    "Basketball": "185226",
    "Football": "185224",
}

SPORT_PATH = {
    "Baseball": "baseball",
    "Basketball": "basketball",
    "Football": "football",
}

COMMON_SETS = [
    "Fleer Ted Williams",
    "Topps Chrome",
    "Bowman Chrome",
    "Topps Heritage",
    "Topps Finest",
    "Upper Deck",
    "Fleer Tradition",
    "Press Pass",
    "Playoff",
    "Donruss Elite",
    "Donruss",
    "Panini",
    "Bowman",
    "Topps",
    "Fleer",
    "Score",
    "Leaf",
    "Hoops",
    "SP",
    "Prizm",
]

GRADE_WORD_TO_RAW = [
    ("NM-MT+", ["NM-MT+", "NMMT+", "NM/MT+"]),
    ("NM-MT", ["NM-MT", "NMMT", "NM/MT"]),
    ("EX-MT", ["EX-MT", "EXMT", "EX/MT"]),
    ("VG-EX", ["VG-EX", "VGEX", "VG/EX"]),
    ("MINT", ["MINT", "MT"]),
    ("NM", ["NM", "NEAR MINT", "NRMT"]),
    ("EX", ["EX", "EXCELLENT"]),
    ("VG", ["VG", "VERY GOOD"]),
    ("GOOD", ["GOOD"]),
    ("POOR", ["POOR"]),
]


def load_products():
    return json.loads(PRODUCTS_PATH.read_text(encoding="utf-8"))


def is_legacy(product):
    metadata = product.get("metadata") if isinstance(product.get("metadata"), dict) else {}
    return bool(
        product.get("sourcePage")
        or product.get("legacyImageLabel")
        or "legacy" in json.dumps(metadata, ensure_ascii=False).lower()
    )


def clean_query(value):
    text = str(value or "")
    text = re.sub(
        r"\b(PSA|BGS|SGC|HGA|CGC|BCCG|GAI)\s*[0-9](?:\.5)?(?:\s*[A-Za-z+\- ]+)?\b",
        " ",
        text,
        flags=re.I,
    )
    text = re.sub(r"\bRookie Card\b|\bGuide range listed\b", " ", text, flags=re.I)
    text = re.sub(r"\b(?:White|Gray|Grey) Back\b", " ", text, flags=re.I)
    text = text.replace("#", " ")
    text = re.sub(r"[/+|()]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def card_number(name):
    match = re.search(r"#\s*([A-Za-z0-9-]+)", str(name or ""))
    return match.group(1) if match else ""


def set_guess(name):
    name = str(name or "")
    for set_name in sorted(COMMON_SETS, key=len, reverse=True):
        if re.search(r"\b" + re.escape(set_name) + r"\b", name, flags=re.I):
            return set_name
    return ""


def first_player(product):
    player = product.get("playerAthlete") or ""
    if player:
        return re.split(r"\s*\|\s*|\s*/\s*", player)[0].strip()
    name = str(product.get("name") or "")
    return re.sub(r"\b\d{4}.*", "", name).strip()


def query_variants(product):
    year = str(product.get("year") or "")
    set_name = set_guess(product.get("name"))
    card = card_number(product.get("name"))
    player = first_player(product)
    name = clean_query(product.get("name"))
    variants = [
        f"{year} {set_name} {card} {player}",
        f"{player} {year} {set_name} {card}",
        f"{year} {set_name} {player}",
        name,
        f"{year} {set_name} {card}",
        f"{year} {name}",
    ]
    cleaned = []
    for variant in variants:
        variant = clean_query(variant)
        if variant and variant not in cleaned:
            cleaned.append(variant)
    return cleaned


def token_set(value):
    return set(re.findall(r"[a-z0-9]+", clean_query(value).lower()))


def candidate_score(product, text, href):
    name_tokens = token_set(product.get("name"))
    haystack = f"{text} {href}".lower()
    candidate_tokens = token_set(haystack)
    overlap = len(name_tokens & candidate_tokens) / max(1, len(name_tokens))
    ratio = SequenceMatcher(None, clean_query(product.get("name")).lower(), text.lower()).ratio() if text else 0
    year_bonus = 0.2 if product.get("year") and str(product["year"]) in haystack else 0
    player = first_player(product)
    player_bonus = 0.25 if player and token_set(player) <= candidate_tokens else 0
    number = card_number(product.get("name"))
    number_bonus = 0
    if number:
        if re.search(rf"(?:#|/){re.escape(number)}[a-z]?\b", haystack):
            number_bonus = 0.45
        elif number.lower() in candidate_tokens:
            number_bonus = 0.2
    set_name = set_guess(product.get("name"))
    set_bonus = 0.15 if set_name and token_set(set_name) <= candidate_tokens else 0
    return overlap + ratio + year_bonus + player_bonus + number_bonus + set_bonus


def login_session(email, password):
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36",
            "Accept-Language": "en-US,en;q=0.9",
        }
    )
    login = session.get("https://www.beckett.com/login", timeout=30)
    login.raise_for_status()
    soup = BeautifulSoup(login.text, "html.parser")
    form = soup.find("form", id="loginFrm")
    token = form.find("input", {"name": "login_token"}).get("value")
    response = session.post(
        "https://www.beckett.com/login/",
        data={
            "login_token": token,
            "email": email,
            "password": password,
            "redirect_url": "https://www.beckett.com/online-price-guide",
            "remember_me": "1",
        },
        timeout=30,
        allow_redirects=True,
    )
    response.raise_for_status()
    if "Logout" not in response.text and "is_user_logged_in = '1'" not in response.text:
        raise RuntimeError("Beckett login did not produce a logged-in session.")
    return session


def find_beckett_match(session, product):
    expected_path = SPORT_PATH.get(product.get("category"))
    if not expected_path:
        return {
            "status": "Not searched",
            "note": "Product is not a baseball, basketball, or football card row.",
        }

    candidates = []
    tried = []
    for query in query_variants(product):
        tried.append(query)
        url = f"https://www.beckett.com/search/?term={quote_plus(query)}"
        response = session.get(url, timeout=30)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        for anchor in soup.find_all("a", href=True):
            href = urljoin("https://www.beckett.com", anchor["href"]).split("?")[0]
            text = anchor.get_text(" ", strip=True)
            if f"beckett.com/{expected_path}/" not in href:
                continue
            if not re.search(r"-\d+$", href):
                continue
            candidates.append(
                {
                    "title": text,
                    "url": href,
                    "score": candidate_score(product, text, href),
                    "query": query,
                }
            )
        if candidates:
            break
        time.sleep(0.1)

    if not candidates:
        return {
            "status": "No match",
            "note": "No Beckett card result found from generated search queries.",
            "queries": tried,
        }

    best_by_url = {}
    for candidate in candidates:
        existing = best_by_url.get(candidate["url"])
        if not existing or candidate["score"] > existing["score"]:
            best_by_url[candidate["url"]] = candidate
    best = max(best_by_url.values(), key=lambda item: item["score"])
    best["status"] = "Matched" if best["score"] >= 1.05 else "Review match"
    best["queries"] = tried
    return best


def parse_money(value):
    match = re.search(r"\$[\d,]+(?:\.\d{2})?", str(value or ""))
    if not match:
        return None
    return float(match.group(0).replace("$", "").replace(",", ""))


def money_range(values):
    nums = [parse_money(value) for value in values if parse_money(value) is not None]
    if not nums:
        return ""
    low = min(nums)
    high = max(nums)
    return f"${low:,.2f}-${high:,.2f}" if low != high else f"${low:,.2f}"


def parse_price_tables(html):
    soup = BeautifulSoup(html, "html.parser")
    title = soup.find("h1").get_text(" ", strip=True) if soup.find("h1") else ""
    tables = []
    for table in soup.find_all("table"):
        rows = []
        for tr in table.find_all("tr"):
            cells = [cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])]
            if cells:
                rows.append(cells)
        if not rows:
            continue
        labels = []
        prices = []
        if len(rows) >= 2 and any("$" in cell for cell in rows[1]):
            labels = rows[0]
            prices = rows[1]
        else:
            flat = [cell for row in rows for cell in row]
            first_price = next((idx for idx, cell in enumerate(flat) if "$" in cell), None)
            if first_price:
                labels = flat[:first_price]
                prices = flat[first_price : first_price + len(labels)]
        if labels and prices and any("$" in price for price in prices):
            table_type = "graded" if any("(" in label and ")" in label for label in labels) else "raw"
            tables.append(
                {
                    "type": table_type,
                    "labels": labels,
                    "prices": prices,
                    "map": {label: price for label, price in zip(labels, prices)},
                    "range": money_range(prices),
                }
            )
    raw_table = next((table for table in tables if table["type"] == "raw"), None)
    graded_table = next((table for table in tables if table["type"] == "graded"), None)
    return title, raw_table, graded_table


def matching_raw_price(condition, raw_table):
    if not raw_table:
        return ""
    normalized = str(condition or "").upper().replace(" ", "")
    for raw_label, aliases in GRADE_WORD_TO_RAW:
        for alias in aliases:
            if alias.replace(" ", "") in normalized:
                for label, price in raw_table["map"].items():
                    if label.upper().replace(" ", "") == raw_label.replace(" ", ""):
                        return price
    return ""


def matching_graded_price(condition, graded_table):
    if not graded_table:
        return ""
    match = re.search(r"\b(?:PSA|BGS|SGC|HGA|CGC|BCCG|GAI)\s*([0-9](?:\.5)?|10)\b", str(condition or ""), flags=re.I)
    if not match:
        return ""
    grade = float(match.group(1))
    best = None
    for label, price in graded_table["map"].items():
        grade_match = re.search(r"\((\d+(?:\.\d+)?)\)", label)
        if not grade_match:
            continue
        label_grade = float(grade_match.group(1))
        distance = abs(label_grade - grade)
        if best is None or distance < best[0]:
            best = (distance, label, price)
    if not best:
        return ""
    suffix = "" if best[0] == 0 else f" (nearest Beckett listed grade: {best[1]})"
    return f"{best[2]}{suffix}"


def fetch_pricing(session, url):
    response = session.get(url, timeout=30)
    response.raise_for_status()
    title, raw_table, graded_table = parse_price_tables(response.text)
    return {
        "beckett_title": title,
        "raw_range": raw_table["range"] if raw_table else "",
        "graded_range": graded_table["range"] if graded_table else "",
        "raw_map": raw_table["map"] if raw_table else {},
        "graded_map": graded_table["map"] if graded_table else {},
    }


def format_original_price(product):
    if product.get("priceLabel"):
        return product["priceLabel"]
    if product.get("price") not in (None, ""):
        return f"${float(product['price']):,.2f}"
    return ""


def load_cache():
    if CACHE_PATH.exists():
        return json.loads(CACHE_PATH.read_text(encoding="utf-8"))
    return {"matches": {}, "prices": {}}


def save_cache(cache):
    CACHE_PATH.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


def build_workbook(rows, summary):
    wb = Workbook()
    ws = wb.active
    ws.title = "Legacy Beckett Pricing"
    headers = [
        "Product ID",
        "Category",
        "Source Page",
        "Title",
        "Original Price / Range",
        "Current Grade / Condition",
        "Beckett Match Status",
        "Beckett Matched Title",
        "Beckett Price Range",
        "Beckett Raw Range",
        "Beckett Graded Range",
        "Beckett Matching Raw Grade Price",
        "Beckett Matching Graded Price",
        "Beckett Price Page",
        "Match Confidence",
        "Notes",
    ]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="172033")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        ws.append([row.get(header, "") for header in headers])
        url = row.get("Beckett Price Page")
        if url:
            cell = ws.cell(ws.max_row, headers.index("Beckett Price Page") + 1)
            cell.hyperlink = url
            cell.style = "Hyperlink"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "A": 11,
        "B": 13,
        "C": 20,
        "D": 48,
        "E": 18,
        "F": 20,
        "G": 18,
        "H": 42,
        "I": 20,
        "J": 18,
        "K": 18,
        "L": 22,
        "M": 24,
        "N": 48,
        "O": 14,
        "P": 42,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    summary_ws = wb.create_sheet("Summary")
    summary_ws.append(["Metric", "Value"])
    for key, value in summary.items():
        summary_ws.append([key, value])
    for cell in summary_ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    summary_ws.column_dimensions["A"].width = 34
    summary_ws.column_dimensions["B"].width = 24
    summary_ws.freeze_panes = "A2"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)


def main():
    email = os.environ.get("BECKETT_EMAIL")
    password = os.environ.get("BECKETT_PASSWORD")
    if not email or not password:
        raise SystemExit("BECKETT_EMAIL and BECKETT_PASSWORD environment variables are required.")

    products = [product for product in load_products() if is_legacy(product)]
    cache = load_cache()
    session = login_session(email, password)
    rows = []
    counts = Counter()

    for index, product in enumerate(products, start=1):
        product_id = str(product.get("id"))
        match = cache["matches"].get(product_id)
        if not match:
            match = find_beckett_match(session, product)
            cache["matches"][product_id] = match
            save_cache(cache)
            time.sleep(0.15)

        pricing = {}
        url = match.get("url")
        if url:
            pricing = cache["prices"].get(url)
            if not pricing:
                try:
                    pricing = fetch_pricing(session, url)
                except Exception as error:
                    pricing = {"error": repr(error)}
                cache["prices"][url] = pricing
                save_cache(cache)
                time.sleep(0.15)

        raw_range = pricing.get("raw_range", "")
        graded_range = pricing.get("graded_range", "")
        price_range = graded_range or raw_range
        condition = product.get("condition") or ""
        raw_match = matching_raw_price(condition, {"map": pricing.get("raw_map", {})}) if pricing.get("raw_map") else ""
        graded_match = (
            matching_graded_price(condition, {"map": pricing.get("graded_map", {})})
            if pricing.get("graded_map")
            else ""
        )
        status = match.get("status", "")
        if pricing.get("error"):
            status = "Price page error"
        counts[status] += 1

        rows.append(
            {
                "Product ID": product.get("id"),
                "Category": product.get("category", ""),
                "Source Page": product.get("sourcePage", ""),
                "Title": product.get("name", ""),
                "Original Price / Range": format_original_price(product),
                "Current Grade / Condition": condition,
                "Beckett Match Status": status,
                "Beckett Matched Title": pricing.get("beckett_title") or match.get("title", ""),
                "Beckett Price Range": price_range,
                "Beckett Raw Range": raw_range,
                "Beckett Graded Range": graded_range,
                "Beckett Matching Raw Grade Price": raw_match,
                "Beckett Matching Graded Price": graded_match,
                "Beckett Price Page": url or "",
                "Match Confidence": round(float(match.get("score", 0)), 3) if match.get("score") else "",
                "Notes": pricing.get("error") or match.get("note", ""),
            }
        )

        if index % 25 == 0 or index == len(products):
            print(f"Processed {index}/{len(products)} legacy listings; matched={counts['Matched']}; review={counts['Review match']}; no_match={counts['No match']}")

    summary = {
        "Generated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Legacy rows reviewed": len(products),
        "Matched rows": counts["Matched"],
        "Review match rows": counts["Review match"],
        "No match rows": counts["No match"],
        "Not searched rows": counts["Not searched"],
        "Price page error rows": counts["Price page error"],
        "Cache file": str(CACHE_PATH),
    }
    build_workbook(rows, summary)
    print(f"Workbook saved: {OUT_PATH}")


if __name__ == "__main__":
    main()
