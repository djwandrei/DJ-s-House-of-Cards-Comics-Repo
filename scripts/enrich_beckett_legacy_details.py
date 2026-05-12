#!/usr/bin/env python3
"""Fill verified legacy product team/manufacturer details from Beckett pages.

The Beckett Legacy workbook decides which product IDs are confirmed matches.
This companion pass trusts those reviewed links and reads only item-detail
fields from the exact Beckett pages, then syncs the storefront metadata.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DOCS = Path.home() / "Documents" / "eBay Docs" / "Listing Automation"
WORKBOOK_PATH = DOCS / "Beckett Legacy Pricing Review.xlsx"
CACHE_PATH = DOCS / "beckett_legacy_pricing_cache.json"
REPORT_PATH = ROOT / "outputs" / "beckett-legacy-detail-enrichment-report.json"

sys.path.insert(0, str(ROOT / "scripts"))
import apply_beckett_legacy_updates as legacy  # noqa: E402
from ensure_beckett_legacy_site_sync import (  # noqa: E402
    card_features,
    manufacturer_from_title,
    normalize_spaces,
    rebuild_product_files,
)

SPORT_CATEGORIES = {"Baseball", "Basketball", "Football"}
ALLOWED_CARD_FEATURES = {"Autograph", "Memorabilia", "Serial Numbered"}

# Beckett intentionally omits Team/Brand on some oddball, college, and leaders
# pages. These reviewed fallbacks cover the small set left blank after direct
# Beckett detail extraction.
TEAM_FALLBACK_BY_PRODUCT_ID = {
    46: "National League",
    47: "National League",
    53: "National League",
    61: "National League",
    67: "National League",
    84: "National League",
    87: "National League",
    99: "National League",
    112: "Houston Astros | California Angels",
    287: "Arizona Diamondbacks",
    420: "New York Knicks",
    506: "Washington Wizards",
    507: "Washington Wizards",
    508: "Washington Wizards",
    552: "San Francisco 49ers",
    555: "Minnesota Vikings",
    557: "Denver Broncos",
    564: "Minnesota Vikings",
    571: "Dallas Cowboys",
    579: "Indianapolis Colts",
    580: "Indianapolis Colts",
    581: "Indianapolis Colts",
    582: "Indianapolis Colts",
    583: "Indianapolis Colts",
    584: "Indianapolis Colts",
    605: "Washington Commanders",
    608: "Washington Commanders",
    609: "Washington Commanders",
    610: "Washington Commanders",
    612: "Cleveland Browns",
    613: "Cleveland Browns",
    614: "Cleveland Browns",
    615: "Cleveland Browns",
    2881: "Chicago White Sox",
    2882: "Chicago White Sox",
    2889: "San Francisco Giants | Los Angeles Dodgers",
    2936: "Detroit Pistons | Milwaukee Bucks | Cleveland Cavaliers",
    2973: "Minnesota Timberwolves",
    2981: "Cincinnati Bengals",
}

MANUFACTURER_FALLBACK_BY_PRODUCT_ID = {
    118: "ASA",
    119: "ASA",
    132: "Perez-Steele",
    165: "Hillshire Farms",
    166: "Hillshire Farms",
    249: "BBM",
}


def load_cache() -> dict[str, Any]:
    if CACHE_PATH.exists():
        return json.loads(CACHE_PATH.read_text(encoding="utf-8"))
    return {"searches": {}, "cards": {}, "completed": {}}


def save_cache(cache: dict[str, Any]) -> None:
    CACHE_PATH.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


def clean_company_name(value: Any) -> str:
    text = normalize_spaces(value)
    text = re.sub(r"\s+Co\.?$", "", text, flags=re.I)
    text = re.sub(r"\s+Inc\.?$", "", text, flags=re.I)
    text = re.sub(r"\s+Company$", "", text, flags=re.I)
    text = re.sub(r"\s+America$", "", text, flags=re.I)
    text = re.sub(r"\s*\([^)]*\)\s*$", "", text)
    return text


def clean_team_name(value: Any) -> str:
    text = normalize_spaces(value)
    text = re.sub(r"\s+BB$", "", text, flags=re.I)
    return text


def parse_detail_items(html: str) -> dict[str, str]:
    """Extract Beckett detail-list values such as Team, Brand, and Manufacturer."""
    soup = BeautifulSoup(html, "html.parser")
    details: dict[str, str] = {}
    for item in soup.find_all("li"):
        strong = item.find("strong")
        if not strong:
            continue
        label = normalize_spaces(strong.get_text(" ", strip=True)).rstrip(":")
        if not label:
            continue
        full_text = normalize_spaces(item.get_text(" ", strip=True))
        value = full_text
        strong_text = normalize_spaces(strong.get_text(" ", strip=True))
        if value.startswith(strong_text):
            value = value[len(strong_text) :].strip(" :")
        if value:
            details[label] = value
    return details


def fetch_details(url: str, session: requests.Session, cache: dict[str, Any], delay: float) -> dict[str, str]:
    cards = cache.setdefault("cards", {})
    card_cache = cards.setdefault(url, {})
    cached_details = card_cache.get("details")
    if isinstance(cached_details, dict) and cached_details.get("detail_fetched_at"):
        return cached_details

    time.sleep(max(0.0, delay))
    response = session.get(url, timeout=45)
    response.raise_for_status()
    details = parse_detail_items(response.text)
    details["detail_fetched_at"] = datetime.now().isoformat(timespec="seconds")
    card_cache["details"] = details
    return details


def confirmed_rows_by_id() -> dict[int, dict[str, Any]]:
    wb = load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
    ws = wb["Legacy Beckett Pricing"]
    rows: dict[int, dict[str, Any]] = {}
    for row_index in range(2, ws.max_row + 1):
        row = legacy.row_to_dict(ws, row_index)
        pid = legacy.product_id_from_value(row.get("Product ID"))
        if pid is not None and legacy.has_confirmed_match(row):
            rows[pid] = row
    return rows


def normalized_brand(details: dict[str, str], fallback_title: str) -> str:
    manufacturer = clean_company_name(details.get("Manufacturer"))
    if manufacturer:
        return manufacturer
    brand = clean_company_name(details.get("Brand"))
    if brand:
        return brand
    return manufacturer_from_title(fallback_title)


def apply_detail_enrichment(product: dict[str, Any], row: dict[str, Any], details: dict[str, str]) -> dict[str, Any]:
    updated = deepcopy(product)
    product_id = int(updated["id"])
    category = normalize_spaces(updated.get("category"))
    title = normalize_spaces(updated.get("name") or row.get("Beckett Matched Title"))

    team = clean_team_name(details.get("Team")) or TEAM_FALLBACK_BY_PRODUCT_ID.get(product_id, "")
    if category in SPORT_CATEGORIES and team:
        updated["team"] = team

    metadata = deepcopy(updated.get("metadata") or {})
    brand = normalized_brand(details, title) or MANUFACTURER_FALLBACK_BY_PRODUCT_ID.get(product_id, "")
    manufacturer = clean_company_name(details.get("Manufacturer"))
    if brand:
        metadata["manufacturer"] = brand
    if manufacturer:
        metadata["manufacturerCompany"] = manufacturer
    for source_key, metadata_key in (
        ("Team", "beckettTeam"),
        ("Sport", "beckettSport"),
        ("Card Number", "beckettCardNumber"),
        ("Brand", "beckettBrand"),
        ("Manufacturer", "beckettManufacturer"),
    ):
        value = normalize_spaces(details.get(source_key))
        if value:
            metadata[metadata_key] = value
    updated["metadata"] = metadata

    features = [feature for feature in card_features(title, row.get("Title"), updated.get("condition")) if feature in ALLOWED_CARD_FEATURES]
    if features:
        updated["attributes"] = features
    else:
        updated.pop("attributes", None)
    return updated


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--delay", type=float, default=0.04)
    parser.add_argument("--limit", type=int, default=0, help="Optional maximum number of network fetches.")
    parser.add_argument(
        "--only-needs-enrichment",
        action="store_true",
        help="Fetch only matched rows with blank team/manufacturer/detail metadata.",
    )
    args = parser.parse_args()

    rows = confirmed_rows_by_id()
    products = json.loads((ROOT / "products.json").read_text(encoding="utf-8"))
    product_by_id = {int(product["id"]): product for product in products}
    cache = load_cache()
    session = requests.Session()
    session.headers.update({"User-Agent": "Mozilla/5.0 (compatible; DJ-House-of-Cards-Inventory-Sync/1.0)"})

    changed: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    fetched = 0
    used_cached = 0

    for pid in sorted(rows):
        product = product_by_id.get(pid)
        if not product:
            continue
        url = normalize_spaces(rows[pid].get("Beckett URL"))
        if not url:
            continue
        metadata = product.get("metadata") if isinstance(product.get("metadata"), dict) else {}
        if args.only_needs_enrichment:
            needs_team = product.get("category") in SPORT_CATEGORIES and not normalize_spaces(product.get("team"))
            needs_manufacturer = not normalize_spaces(metadata.get("manufacturer"))
            needs_beckett_details = not normalize_spaces(metadata.get("beckettTeam")) and not normalize_spaces(metadata.get("beckettBrand"))
            if not (needs_team or needs_manufacturer or needs_beckett_details):
                continue
        try:
            before_fetch_cached = bool(cache.get("cards", {}).get(url, {}).get("details", {}).get("detail_fetched_at"))
            if args.limit and not before_fetch_cached and fetched >= args.limit:
                continue
            details = fetch_details(url, session, cache, args.delay)
            if before_fetch_cached:
                used_cached += 1
            else:
                fetched += 1
        except Exception as error:  # noqa: BLE001 - report per-card failures.
            errors.append({"id": pid, "url": url, "error": str(error)})
            continue

        updated = apply_detail_enrichment(product, rows[pid], details)
        if updated != product:
            product_by_id[pid] = updated
            changed.append(
                {
                    "id": pid,
                    "name": updated.get("name"),
                    "team": updated.get("team"),
                    "manufacturer": (updated.get("metadata") or {}).get("manufacturer"),
                }
            )
        if fetched and fetched % 25 == 0:
            save_cache(cache)

    rebuild_product_files(list(product_by_id.values()))
    save_cache(cache)

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "matchedRows": len(rows),
        "fetchedPages": fetched,
        "usedCachedPages": used_cached,
        "changedProducts": len(changed),
        "changedSample": changed[:80],
        "errorCount": len(errors),
        "errors": errors[:50],
    }
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
