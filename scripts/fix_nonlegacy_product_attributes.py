#!/usr/bin/env python3
"""Correct non-legacy storefront attributes and keep the eBay workbook in sync.

The generated eBay descriptions previously treated every ``/123`` pattern as
"Serial Numbered". That is correct for most card titles, but it breaks on lot
titles such as ``+ /1990 Hoops Set`` where the slash value is actually a year.
This script applies a tighter serial-number detector to the storefront catalog,
rebuilds the static product bundles, and updates the matching eBay upload rows.
"""

from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = Path.home() / "Documents" / "eBay Docs" / "Listing Automation" / "Ebay Bulk Upload with HTML.xlsx"
BACKUP_ROOT = Path.home() / "Documents" / "eBay Docs" / "Listing Automation" / "backups"
REPORT_PATH = ROOT / "outputs" / "nonlegacy-product-attribute-fixes.json"

PRODUCT_FILES = {
    "products.json": lambda items: items,
    "products-baseball.json": lambda items: [item for item in items if item.get("category") == "Baseball"],
    "products-basketball.json": lambda items: [item for item in items if item.get("category") == "Basketball"],
    "products-football.json": lambda items: [item for item in items if item.get("category") == "Football"],
    "products-comics.json": lambda items: [item for item in items if item.get("category") == "Comics"],
    "products-collectibles.json": lambda items: [item for item in items if item.get("category") == "Collectibles"],
    "products-sports.json": lambda items: [
        item for item in items if item.get("category") in {"Baseball", "Basketball", "Football"}
    ],
    "products-featured.json": lambda items: [item for item in items if item.get("isFeatured")],
}

DATA_BUNDLE_FILES = {
    "products.json": "products-data-full.js",
    "products-baseball.json": "products-data-baseball.js",
    "products-basketball.json": "products-data-basketball.js",
    "products-football.json": "products-data-football.js",
    "products-comics.json": "products-data-comics.js",
    "products-collectibles.json": "products-data-collectibles.js",
    "products-sports.json": "products-data-sports.js",
    "products-featured.json": "products-data-featured.js",
}

FEATURE_ORDER = ["Rookie", "Short Print", "One of One", "Serial Numbered", "Memorabilia", "Parallel/Variety"]
FALSE_SERIAL_IDS = {854, 857, 1042}

SERIAL_SLASH_RE = re.compile(r"/\s*(\d{1,4})\b", re.I)
EXPLICIT_SERIAL_RE = re.compile(
    r"\bserial[- ](?:ly[- ])?numbered\b"
    r"|\bnumbered\s+(?:to|/)\s*\d+\b"
    r"|\blimited\s+to\s+\d+\b"
    r"|\bone\s+of\s+one\b"
    r"|\b1\s*of\s*1\b"
    r"|\b1/1\b",
    re.I,
)
YEAR_SERIAL_CONTEXT_RE = re.compile(
    r"\b(?:gold|silver|bronze|blue|red|green|purple|orange|pink|black|aqua|yellow|fuchsia|lime|cyan|white|"
    r"tie-dye|rainbow|platinum|foil|border|parallel|refractor|prizm|holo|shimmer|wave|lava|pulsar|"
    r"speckle|mojo|ice|glitter|chrome|optic|choice|cosmic|sapphire|x-?fractor|the finals|playoff ticket|"
    r"premium stock|masterpieces|limited|numbered|serial|short print|sp)\b",
    re.I,
)
CARD_COUNT_AFTER_RE = re.compile(r"^\s*(?:cards?|pcs?|boxes?|packs?)\b", re.I)
AUTO_RE = re.compile(r"\b(auto(?:s|graph(?:ed|s)?)?|autographed|signed|signature|signatures|on-card)\b", re.I)

JASON_MAXIELL_FIX = {
    "id": 874,
    "category": "Basketball",
    "sport": "Basketball",
    "league": "NBA",
    "team": "Detroit Pistons",
    "player": "Jason Maxiell",
}
PLAYER_CASE_FIXES = {
    "Zz Clark": "ZZ Clark",
}


def backup_file(path: Path) -> Path:
    BACKUP_ROOT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = BACKUP_ROOT / f"{path.stem}.backup-{stamp}{path.suffix}"
    shutil.copy2(path, backup)
    return backup


def split_features(value: Any) -> list[str]:
    parts = [part.strip() for part in str(value or "").split("|") if part and part.strip()]
    return list(dict.fromkeys(parts))


def join_features(features: list[str]) -> str:
    known = [feature for feature in FEATURE_ORDER if feature in features]
    extra = [feature for feature in features if feature not in FEATURE_ORDER]
    return "|".join(known + extra)


def is_serial_numbered_title(title: str) -> bool:
    """Return true only when the title has a real serial-number signal.

    Four-digit years after a plus sign are common in lot titles. We only accept
    19xx/20xx slash values when the surrounding wording says the card is a
    parallel/numbered variant, such as Gold Border /2023.
    """

    if EXPLICIT_SERIAL_RE.search(title):
        return True

    for segment in re.split(r"\s+\+\s+", title):
        for match in SERIAL_SLASH_RE.finditer(segment):
            denominator = int(match.group(1))
            after = segment[match.end() : match.end() + 30]
            if CARD_COUNT_AFTER_RE.match(after):
                continue

            before = segment[max(0, match.start() - 90) : match.start()]
            if 1900 <= denominator <= 2035 and not YEAR_SERIAL_CONTEXT_RE.search(before):
                continue

            return True

    return False


def desired_features(title: str, existing_value: Any) -> str:
    features = split_features(existing_value)
    has_serial = "Serial Numbered" in features
    should_have_serial = is_serial_numbered_title(title)

    if should_have_serial and not has_serial:
        features.append("Serial Numbered")
    elif has_serial and not should_have_serial:
        features = [feature for feature in features if feature != "Serial Numbered"]

    return join_features(features)


def sentence_join(parts: list[str]) -> str:
    if not parts:
        return ""
    if len(parts) == 1:
        return parts[0]
    if len(parts) == 2:
        return f"{parts[0]} and {parts[1]}"
    return f"{', '.join(parts[:-1])}, and {parts[-1]}"


def feature_appeals(features: list[str]) -> list[str]:
    labels = {
        "Rookie": "rookie-card appeal",
        "Short Print": "short-print appeal",
        "One of One": "one-of-one rarity",
        "Serial Numbered": "serial-numbered appeal",
        "Memorabilia": "memorabilia-card appeal",
        "Parallel/Variety": "parallel/variation appeal",
    }
    return [labels[feature] for feature in FEATURE_ORDER if feature in features]


def condition_summary(product: dict[str, Any]) -> str:
    condition = str(product.get("condition") or "").strip()
    if re.search(r"\b(PSA|BGS|SGC|CGC|HGA|BCCG|GAI|CSG)\b", condition, re.I):
        return condition
    return "Ungraded"


def build_description(product: dict[str, Any]) -> str:
    metadata = product.setdefault("metadata", {})
    excel_fields = metadata.setdefault("excelFields", {})
    features = split_features(excel_fields.get("C:Features"))
    set_name = excel_fields.get("C:Set") or product.get("name") or "This listing"
    sport = product.get("sport") or excel_fields.get("C:Sport") or product.get("category") or "cards"
    league = product.get("league") or metadata.get("league") or excel_fields.get("C:League") or ""
    player = product.get("playerAthlete") or metadata.get("playerAthlete") or excel_fields.get("C:Player/Athlete") or ""
    team = product.get("team") or excel_fields.get("C:Team") or ""

    subject = f" card featuring {player}" if player else " listing"
    team_text = f" with connections to {team}" if team else ""
    collector_text = f", {league} collectors" if league else " collectors"
    appeals = feature_appeals(features)
    appeal_text = f" thanks to its {sentence_join(appeals)}" if appeals else ""

    details = [
        ("Set", set_name),
        ("Year", product.get("year")),
        ("Player", player),
        ("Team", team),
        ("Sport", sport),
        ("League", league),
        ("Condition", condition_summary(product)),
        ("Features", " | ".join(features)),
    ]
    details_text = "; ".join(f"{label}: {value}" for label, value in details if value not in (None, ""))

    return (
        f"{set_name}{subject}{team_text}. "
        f"A strong addition for {str(sport).lower()} fans{collector_text}, and set builders{appeal_text}. "
        f"Details: {details_text}. "
        "Please review the photos for the exact card you will receive. "
        "Great for player, team, and vintage or modern card collections."
    )


def apply_jason_maxiell_fix(product: dict[str, Any]) -> bool:
    if product.get("id") != JASON_MAXIELL_FIX["id"]:
        return False

    changed = False
    metadata = product.setdefault("metadata", {})
    excel_fields = metadata.setdefault("excelFields", {})

    updates = {
        "category": JASON_MAXIELL_FIX["category"],
        "sport": JASON_MAXIELL_FIX["sport"],
        "league": JASON_MAXIELL_FIX["league"],
        "team": JASON_MAXIELL_FIX["team"],
        "playerAthlete": JASON_MAXIELL_FIX["player"],
    }
    for key, value in updates.items():
        if product.get(key) != value:
            product[key] = value
            changed = True

    metadata_updates = {
        "sport": JASON_MAXIELL_FIX["sport"],
        "league": JASON_MAXIELL_FIX["league"],
        "playerAthlete": JASON_MAXIELL_FIX["player"],
    }
    for key, value in metadata_updates.items():
        if metadata.get(key) != value:
            metadata[key] = value
            changed = True

    field_updates = {
        "C:Sport": JASON_MAXIELL_FIX["sport"],
        "C:League": JASON_MAXIELL_FIX["league"],
        "C:Team": JASON_MAXIELL_FIX["team"],
        "C:Player/Athlete": JASON_MAXIELL_FIX["player"],
    }
    for key, value in field_updates.items():
        if excel_fields.get(key) != value:
            excel_fields[key] = value
            changed = True

    return changed


def write_json(path: Path, data: list[dict[str, Any]]) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")


def write_data_bundle(source_name: str, bundle_path: Path, data: list[dict[str, Any]]) -> None:
    serialized = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    bundle = (
        f'window.DJ_PRELOADED_SOURCE = "{source_name}";\n'
        f"window.DJ_PRELOADED_PRODUCTS = {serialized}\n"
        ";\n"
    )
    bundle_path.write_text(bundle, encoding="utf-8")


def rebuild_product_files(products: list[dict[str, Any]]) -> None:
    for source_name, selector in PRODUCT_FILES.items():
        subset = selector(products)
        write_json(ROOT / source_name, subset)
        write_data_bundle(source_name, ROOT / DATA_BUNDLE_FILES[source_name], subset)


def update_products() -> tuple[list[dict[str, Any]], Counter, list[dict[str, Any]]]:
    products_path = ROOT / "products.json"
    products = json.loads(products_path.read_text(encoding="utf-8"))
    changes: Counter = Counter()
    touched: list[dict[str, Any]] = []

    for product in products:
        metadata = product.get("metadata") if isinstance(product.get("metadata"), dict) else {}
        excel_fields = metadata.get("excelFields") if isinstance(metadata.get("excelFields"), dict) else {}
        if not excel_fields:
            continue

        product_changed = apply_jason_maxiell_fix(product)
        for source, target in PLAYER_CASE_FIXES.items():
            for container, key in ((product, "playerAthlete"), (metadata, "playerAthlete"), (excel_fields, "C:Player/Athlete")):
                if container.get(key) == source:
                    container[key] = target
                    changes["product_player_case_updates"] += 1
                    product_changed = True

        title = str(product.get("name") or excel_fields.get("Title") or "")
        current_features = excel_fields.get("C:Features")
        next_features = desired_features(title, current_features)

        core_field_defaults = {
            "C:Sport": product.get("sport") or product.get("category") or "",
            "C:League": product.get("league") or metadata.get("league") or "",
            "C:Team": product.get("team") or "",
            "C:Player/Athlete": product.get("playerAthlete") or metadata.get("playerAthlete") or "",
            "C:Features": next_features,
        }
        for key, value in core_field_defaults.items():
            if key not in excel_fields and value not in (None, ""):
                excel_fields[key] = value
                changes["product_missing_core_fields_filled"] += 1
                product_changed = True

        if next_features != str(current_features or ""):
            excel_fields["C:Features"] = next_features
            changes["product_feature_rows_updated"] += 1
            product_changed = True

        next_auto = "Yes" if AUTO_RE.search(title) else "No"
        if excel_fields.get("C:Autographed") != next_auto:
            excel_fields["C:Autographed"] = next_auto
            changes["product_autograph_rows_updated"] += 1
            product_changed = True

        if product_changed:
            product["description"] = build_description(product)
            excel_fields["Description"] = product["description"].replace(". Details:", ".\nDetails:").replace(". Please review", ".\nPlease review")
            touched.append(
                {
                    "id": product.get("id"),
                    "title": product.get("name"),
                    "category": product.get("category"),
                    "features": excel_fields.get("C:Features"),
                }
            )

    rebuild_product_files(products)
    return products, changes, touched


def update_workbook(products_by_title: dict[str, dict[str, Any]]) -> tuple[Path, Counter]:
    backup = backup_file(WORKBOOK_PATH)
    wb = load_workbook(WORKBOOK_PATH)
    ws = wb["Sheet1"]
    headers = [ws.cell(row=1, column=column).value for column in range(1, ws.max_column + 1)]
    idx = {header: position + 1 for position, header in enumerate(headers)}
    changes: Counter = Counter()

    for row_num in range(2, ws.max_row + 1):
        title = str(ws.cell(row=row_num, column=idx["Title"]).value or "")
        product = products_by_title.get(title)
        if not product:
            continue

        metadata = product.get("metadata") if isinstance(product.get("metadata"), dict) else {}
        excel_fields = metadata.get("excelFields") if isinstance(metadata.get("excelFields"), dict) else {}
        if not excel_fields:
            continue

        for column_name in ["C:Features", "C:Autographed", "C:Sport", "C:League", "C:Team", "C:Player/Athlete", "Description"]:
            if column_name not in idx:
                continue
            desired = excel_fields.get(column_name)
            cell = ws.cell(row=row_num, column=idx[column_name])
            if str(cell.value or "") != str(desired or ""):
                cell.value = desired
                changes[f"workbook_{column_name}_updated"] += 1

    wb.save(WORKBOOK_PATH)
    return backup, changes


def main() -> int:
    products, product_changes, touched = update_products()
    workbook_backup, workbook_changes = update_workbook({str(product.get("name")): product for product in products})

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "workbook": str(WORKBOOK_PATH),
        "workbookBackup": str(workbook_backup),
        "productChanges": dict(product_changes),
        "workbookChanges": dict(workbook_changes),
        "touchedProducts": touched,
    }
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
