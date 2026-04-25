#!/usr/bin/env python3
"""Improve confirmed Beckett matches in the non-legacy pricing workbook.

The original Beckett workbook generator intentionally leaves many candidates in
review when it sees possible set, variant, or card-type mismatches. This pass is
more targeted: it retries unresolved/review rows with extra query variants and
only promotes a row to "Matched" when the best result has strong evidence and no
hard conflicts. The workbook is backed up before any edits are saved.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime
import json
import os
import re
import shutil
import sys
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from beckett_legacy_pricing import (  # noqa: E402
    BECKETT_BASE,
    SPORT_CATEGORIES,
    build_workbook,
    create_session,
    extract_card_number,
    extract_year,
    fetch_card_page,
    format_money,
    format_percent,
    is_graded,
    load_cache,
    matched_grade_price,
    midpoint_from_range,
    normalize_card_code,
    normalize_key,
    normalize_spaces,
    price_range,
    price_signal,
    query_variants,
    save_cache,
    score_candidate,
    search_beckett,
    selected_products,
)
from update_beckett_realtime_column import (  # noqa: E402
    normalize_realtime_label,
    parse_realtime_pricing,
)


DEFAULT_WORKBOOK = (
    Path.home()
    / "Documents"
    / "eBay Docs"
    / "Listing Automation"
    / "Beckett Non-Legacy Pricing Review.xlsx"
)
DEFAULT_CACHE = (
    Path.home()
    / "Documents"
    / "eBay Docs"
    / "Listing Automation"
    / "beckett_pricing_cache_all.json"
)
BACKUP_ROOT = Path.home() / "Documents" / "eBay Docs" / "Listing Automation" / "backups"
REPORT_PATH = ROOT / "outputs" / "beckett-nonlegacy-match-improvement-report.json"

MAIN_SHEET = "Non-Legacy Beckett Pricing"
REALTIME_HEADER = "Beckett Real Time Pricing"
MATCHED_STATUSES = {"Matched", "Not searched - non sports-card row"}

CARD_CODE_RE = re.compile(r"(?:#|no\.?\s*)([A-Za-z0-9]+(?:[-/][A-Za-z0-9]+)*)", re.I)
SERIAL_RE = re.compile(r"/\s*(\d{1,5})(?:\D|$)")
YEAR_RE = re.compile(r"\b((?:18|19|20)\d{2}(?:[-/](?:\d{2}|\d{4}))?)\b")

FEATURE_GROUPS = {
    "autograph": {
        "source": {"auto", "autos", "autograph", "autographs", "signed", "signature", "signatures"},
        "target": {
            "auto",
            "autos",
            "autograph",
            "autographs",
            "au",
            "ink",
            "signed",
            "signature",
            "signatures",
            "signings",
            "scripts",
            "lettermen",
            "playergraphs",
            "graphs",
            "chirography",
        },
    },
    "memorabilia": {
        "source": {
            "patch",
            "patches",
            "jersey",
            "jerseys",
            "jsy",
            "relic",
            "relics",
            "memorabilia",
            "swatch",
            "swatches",
            "material",
            "materials",
            "game",
            "worn",
            "used",
            "bat",
            "ball",
            "helmet",
            "glove",
        },
        "target": {
            "patch",
            "patches",
            "jersey",
            "jerseys",
            "jsy",
            "relic",
            "relics",
            "memorabilia",
            "swatch",
            "swatches",
            "material",
            "materials",
            "fabric",
            "fabrics",
            "threads",
            "authentics",
            "bat",
            "ball",
            "helmet",
            "glove",
            "caps",
            "manufactured",
        },
    },
}

VARIANT_WORDS = {
    "tiffany",
    "refractor",
    "refractors",
    "xfractor",
    "xfractors",
    "prizm",
    "mojo",
    "pulsar",
    "shimmer",
    "lava",
    "speckle",
    "cracked",
    "ice",
    "wave",
    "foil",
    "glitter",
    "black",
    "blue",
    "red",
    "green",
    "gold",
    "silver",
    "purple",
    "orange",
    "pink",
    "bronze",
    "platinum",
}

NOISE_QUERY_WORDS = re.compile(
    r"\b(SP|RC|rookie card|rookie|card|cards|set|lot|x\d+|serial numbered|near mint|mint|ungraded)\b",
    re.I,
)


def load_rows(workbook_path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    ws = wb[MAIN_SHEET]
    headers = [normalize_spaces(cell.value) for cell in ws[1]]
    rows: list[dict[str, Any]] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if raw[0] is None:
            continue
        rows.append({header: raw[idx] for idx, header in enumerate(headers)})
    return headers, rows


def backup_workbook(workbook_path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-nonlegacy-improve")
    backup_dir = BACKUP_ROOT / stamp
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / workbook_path.name
    shutil.copy2(workbook_path, backup_path)
    return backup_path


def price_label(product: dict[str, Any]) -> str:
    return str(product.get("priceLabel") or product.get("displayPrice") or product.get("price") or "")


def row_product(row: dict[str, Any], products_by_id: dict[int, dict[str, Any]]) -> dict[str, Any] | None:
    try:
        product_id = int(row.get("Product ID"))
    except (TypeError, ValueError):
        return None
    return products_by_id.get(product_id)


def tokens(value: Any) -> set[str]:
    return set(normalize_key(value).split())


def title_segments(title: str) -> list[str]:
    """Split lot titles into searchable single-card-like pieces."""

    cleaned = re.sub(r"\((?:x\s*)?\d+\)", " ", title, flags=re.I)
    cleaned = re.sub(r"\bset\s+of\s+\d+\b", " ", cleaned, flags=re.I)
    pieces = [cleaned]
    for separator in [" + ", " / ", " | ", ";"]:
        next_pieces: list[str] = []
        for piece in pieces:
            next_pieces.extend(piece.split(separator))
        pieces = next_pieces
    normalized: list[str] = []
    for piece in pieces:
        piece = normalize_spaces(piece)
        if len(piece) >= 8 and piece not in normalized:
            normalized.append(piece)
    return normalized[:5]


def normalized_card_codes(*values: Any) -> set[str]:
    codes: set[str] = set()
    for value in values:
        for match in CARD_CODE_RE.finditer(str(value or "")):
            code = normalize_card_code(match.group(1))
            if code:
                codes.add(code)
    return codes


def serial_denominators(value: Any) -> set[str]:
    return {match.group(1) for match in SERIAL_RE.finditer(str(value or ""))}


def enhanced_query_variants(product: dict[str, Any], row: dict[str, Any]) -> list[str]:
    title = normalize_spaces(row.get("Title") or product.get("name"))
    year = normalize_spaces(row.get("Site Year") or product.get("year") or extract_year(product))
    player = normalize_spaces(row.get("Site Player / Athlete") or product.get("playerAthlete"))
    card_codes = normalized_card_codes(title)
    values: list[str] = []

    def add(value: str) -> None:
        value = normalize_spaces(value)
        value = re.sub(r"\s+", " ", value)
        if value and value not in values and len(value) >= 4:
            values.append(value)

    for query in query_variants(product):
        add(query)

    for segment in title_segments(title):
        base = NOISE_QUERY_WORDS.sub(" ", segment)
        base = re.sub(r"\bCertified Autograph Issue\b", "Autographs", base, flags=re.I)
        base = re.sub(r"\bAuto\b", "Autographs", base, flags=re.I)
        base = normalize_spaces(base.replace("#", " "))
        add(segment)
        add(base)
        if year:
            add(f"{year} {base}")
        if player:
            add(f"{year} {player} {base}")
            add(f"{player} {year} {base}")
        for code in card_codes:
            dashed = " ".join(re.findall(r"[A-Za-z]+|\d+", code))
            add(f"{year} {player} {code}")
            add(f"{year} {player} {dashed}")
            add(f"{year} {base} {code}")
            add(f"{player} {code}")

    # Some Beckett card numbers collapse punctuation, e.g. #BA-JL -> #BAJL.
    for code in list(card_codes):
        if len(code) >= 3:
            add(f"{year} {player} {code}")
            add(f"{player} {code}")

    return values[:36]


def candidate_payload(title: str, url: str) -> dict[str, str]:
    return {"title": normalize_spaces(title), "url": normalize_spaces(url)}


def hard_conflicts(product: dict[str, Any], row: dict[str, Any], candidate: dict[str, str]) -> list[str]:
    source_title = normalize_spaces(row.get("Title") or product.get("name"))
    target_title = normalize_spaces(candidate.get("title"))
    target_url = normalize_spaces(candidate.get("url"))
    source_key = normalize_key(source_title)
    target_key = normalize_key(f"{target_title} {target_url}")
    source_tokens = tokens(source_title)
    target_tokens = tokens(f"{target_title} {target_url}")
    conflicts: list[str] = []

    category = normalize_spaces(row.get("Category") or product.get("category"))
    if category in SPORT_CATEGORIES and f"/{category.lower()}/" not in target_url.lower() and "/multisport/" not in target_url.lower():
        conflicts.append("sport path mismatch")

    expected_year = str(row.get("Site Year") or product.get("year") or "").split("-")[0].strip()
    if expected_year:
        years = YEAR_RE.findall(f"{target_title} {target_url}")
        if years and not any(str(year).startswith(expected_year) for year in years):
            conflicts.append("year mismatch")

    source_codes = normalized_card_codes(source_title)
    target_compact = normalize_card_code(f"{target_title} {target_url}")
    if source_codes and not any(code in target_compact for code in source_codes):
        # Numeric set/card numbers are often absent from lot titles; keep this as
        # a soft signal unless the source used an alpha code like BAJL or GMA-LH.
        if any(re.search(r"[A-Z]", code) and len(code) >= 3 for code in source_codes):
            conflicts.append("card number mismatch")

    source_serials = serial_denominators(source_title)
    target_serials = serial_denominators(f"{target_title} {target_url}")
    if source_serials and target_serials and not (source_serials & target_serials):
        conflicts.append("serial-number denominator mismatch")

    for group in FEATURE_GROUPS.values():
        if source_tokens & group["source"] and not (target_tokens & group["target"]):
            conflicts.append("feature/type mismatch")
            break

    # The broad memorabilia group above catches "this is some relic", but
    # Beckett often distinguishes jersey, patch, bat, ball, helmet, etc. Keep
    # those stricter so a jersey listing does not silently become a bat card.
    type_checks = [
        ({"bat", "bats"}, {"bat", "bats"}, "bat material mismatch"),
        ({"ball", "balls"}, {"ball", "balls"}, "ball material mismatch"),
        ({"helmet", "helmets"}, {"helmet", "helmets"}, "helmet material mismatch"),
        ({"glove", "gloves"}, {"glove", "gloves"}, "glove material mismatch"),
        (
            {"jersey", "jerseys", "jsy", "uniform", "uniforms"},
            {"jersey", "jerseys", "jsy", "uniform", "uniforms", "swatch", "swatches", "fabric", "fabrics"},
            "jersey material mismatch",
        ),
        (
            {"patch", "patches"},
            {"patch", "patches", "prime", "swatch", "swatches", "material", "materials"},
            "patch material mismatch",
        ),
    ]
    for source_terms, target_terms, label in type_checks:
        if source_tokens & source_terms and not (target_tokens & target_terms):
            conflicts.append(label)

    # Avoid accepting a named parallel/variant that the source never mentions.
    target_only_variants = (target_tokens & VARIANT_WORDS) - (source_tokens & VARIANT_WORDS)
    if target_only_variants:
        allowed = {"relic", "relics"} & source_tokens and {"swatch", "swatches", "material", "materials"} & target_tokens
        if not allowed:
            conflicts.append(f"variant mismatch: {', '.join(sorted(target_only_variants))}")

    # Reject obvious different-player results. This stays conservative: rows with
    # multi-player products are allowed if at least one named player matches.
    player_text = normalize_spaces(row.get("Site Player / Athlete") or product.get("playerAthlete"))
    if player_text:
        player_names = [normalize_spaces(name) for name in re.split(r"\s*\|\s*|\s*/\s*|;", player_text) if name.strip()]
        meaningful = [name for name in player_names if len(tokens(name)) >= 2]
        if meaningful and not any(tokens(name) <= target_tokens for name in meaningful):
            # The imported player field is sometimes polluted with set words, so
            # only block when the source title itself also names that player.
            if any(tokens(name) <= source_tokens for name in meaningful):
                conflicts.append("player mismatch")

    if "tiffany" in target_key and "tiffany" not in source_key:
        conflicts.append("tiffany variant mismatch")

    return conflicts


def deep_score(product: dict[str, Any], row: dict[str, Any], query: str, candidate: dict[str, str]) -> float:
    base = score_candidate(product, query, candidate)
    source_title = normalize_spaces(row.get("Title") or product.get("name"))
    target_blob = normalize_spaces(f"{candidate.get('title')} {candidate.get('url')}")
    source_tokens = tokens(source_title)
    target_tokens = tokens(target_blob)
    score = base

    source_codes = normalized_card_codes(source_title)
    target_compact = normalize_card_code(target_blob)
    if source_codes and any(code in target_compact for code in source_codes):
        score += 0.16

    player_text = normalize_spaces(row.get("Site Player / Athlete") or product.get("playerAthlete"))
    if player_text:
        player_names = [normalize_spaces(name) for name in re.split(r"\s*\|\s*|\s*/\s*|;", player_text) if name.strip()]
        if any(tokens(name) and tokens(name) <= target_tokens for name in player_names):
            score += 0.12

    expected_year = str(row.get("Site Year") or product.get("year") or "").split("-")[0].strip()
    if expected_year and expected_year in normalize_key(target_blob):
        score += 0.08

    for group in FEATURE_GROUPS.values():
        if source_tokens & group["source"] and target_tokens & group["target"]:
            score += 0.05

    if source_tokens & VARIANT_WORDS and (source_tokens & VARIANT_WORDS) <= (target_tokens | {"refractors"}):
        score += 0.04

    return round(max(0.0, min(1.25, score)), 4)


def best_enhanced_candidate(
    session: Any,
    product: dict[str, Any],
    row: dict[str, Any],
    cache: dict[str, Any],
    delay: float,
    existing_only: bool = False,
) -> tuple[dict[str, str] | None, str, float, list[str]]:
    candidates: list[tuple[float, str, dict[str, str], list[str]]] = []

    existing_title = normalize_spaces(row.get("Beckett Matched Title"))
    existing_url = normalize_spaces(row.get("Beckett URL"))
    if existing_title and existing_url:
        existing = candidate_payload(existing_title, existing_url)
        conflicts = hard_conflicts(product, row, existing)
        candidates.append((deep_score(product, row, normalize_spaces(row.get("Lookup Query")), existing), "existing workbook candidate", existing, conflicts))

    if existing_only:
        if not candidates:
            return None, "", 0.0, []
        score, query, candidate, conflicts = candidates[0]
        return candidate, query, score, conflicts

    for query in enhanced_query_variants(product, row):
        try:
            results = search_beckett(session, query, cache, delay)
        except Exception as exc:  # noqa: BLE001 - recorded in report by caller.
            cache.setdefault("search_errors", {})[query] = str(exc)
            continue
        for result in results[:20]:
            candidate = candidate_payload(result["title"], result["url"])
            conflicts = hard_conflicts(product, row, candidate)
            score = deep_score(product, row, query, candidate)
            candidates.append((score, query, candidate, conflicts))
        if candidates and max(item[0] for item in candidates) >= 1.02:
            break

    if not candidates:
        return None, "", 0.0, []

    candidates.sort(key=lambda item: (len(item[3]) == 0, item[0]), reverse=True)
    score, query, candidate, conflicts = candidates[0]
    return candidate, query, score, conflicts


def realtime_for_url(session: Any, url: str, cache: dict[str, Any], delay: float) -> str:
    cached = normalize_realtime_label(cache.get("cards", {}).get(url, {}).get("real_time_pricing"))
    if cached:
        return cached
    response = session.get(url, timeout=45)
    response.raise_for_status()
    label = parse_realtime_pricing(response.text)
    if label:
        cache.setdefault("cards", {}).setdefault(url, {})["real_time_pricing"] = label
        cache.setdefault("cards", {}).setdefault(url, {})["real_time_pricing_fetched_at"] = datetime.now().isoformat(timespec="seconds")
    if delay:
        time.sleep(delay)
    return label


def apply_card_data(
    row: dict[str, Any],
    product: dict[str, Any],
    candidate: dict[str, str],
    query: str,
    score: float,
    session: Any,
    cache: dict[str, Any],
    delay: float,
) -> None:
    card = fetch_card_page(session, candidate["url"], cache, delay)
    raw_prices = card.get("raw_prices", {})
    graded_prices = card.get("graded_prices", {})
    raw_market = card.get("raw_market", {})
    graded_market = card.get("graded_market", {})
    raw_range = price_range(raw_prices)
    graded_range = price_range(graded_prices)
    matched_price, matched_basis, grade_note = matched_grade_price(
        str(product.get("condition", "")),
        raw_prices,
        graded_prices,
    )

    row["Lookup Query"] = query
    row["Match Confidence"] = score
    row["Beckett Matched Title"] = normalize_spaces(card.get("title") or candidate["title"])
    row["Beckett URL"] = candidate["url"]
    row["Beckett Raw Range"] = raw_range
    row["Beckett Graded Range"] = graded_range
    row["Beckett Raw Market Range"] = raw_market.get("range", "")
    row["Beckett Graded Market Range"] = graded_market.get("range", "")
    row["Beckett Page Source"] = card.get("source_page_title", "")
    row["Beckett Matched Grade Price"] = matched_price
    row["Beckett Matched Grade Basis"] = matched_basis
    row["Beckett Price Range"] = graded_range if is_graded(str(product.get("condition", ""))) and graded_range else raw_range
    if raw_range or graded_range:
        row["Beckett Availability"] = "Guide pricing"
    elif row["Beckett Raw Market Range"] or row["Beckett Graded Market Range"]:
        row["Beckett Availability"] = "Market sales only"
    else:
        row["Beckett Availability"] = "No visible pricing"
    site_mid = midpoint_from_range(row.get("Original Price / Range") or price_label(product))
    beckett_mid = midpoint_from_range(row.get("Beckett Price Range"))
    delta_money, delta_pct, signal = price_signal(site_mid, beckett_mid)
    row["Site Price Midpoint"] = format_money(site_mid)
    row["Beckett Price Midpoint"] = format_money(beckett_mid)
    row["Midpoint Delta"] = delta_money
    row["Midpoint Delta %"] = format_percent(delta_pct) if isinstance(delta_pct, float) else delta_pct
    row["Price Signal"] = signal
    if REALTIME_HEADER in row:
        row[REALTIME_HEADER] = realtime_for_url(session, candidate["url"], cache, delay)
    notes = ["Enhanced non-legacy pass confirmed this Beckett match."]
    if grade_note:
        notes.append(grade_note)
    if row["Beckett Availability"] == "No visible pricing":
        notes.append("Beckett page matched, but no visible guide pricing table was parsed.")
    row["Notes"] = " ".join(notes)


def before_after_stats(rows: list[dict[str, Any]]) -> dict[str, Any]:
    statuses = Counter(normalize_spaces(row.get("Match Status")) for row in rows)
    active = len(rows)
    sports = sum(1 for row in rows if row.get("Category") in SPORT_CATEGORIES)
    matched = statuses["Matched"]
    return {
        "rows": active,
        "sportsRows": sports,
        "matched": matched,
        "matchRateAllRows": round((matched / active * 100) if active else 0, 2),
        "matchRateSportsRows": round((matched / sports * 100) if sports else 0, 2),
        "statuses": dict(statuses),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--products", type=Path, default=ROOT / "products.json")
    parser.add_argument("--cache", type=Path, default=DEFAULT_CACHE)
    parser.add_argument("--delay", type=float, default=0.12)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--save-every", type=int, default=25)
    parser.add_argument("--existing-only", action="store_true")
    parser.add_argument(
        "--statuses",
        nargs="*",
        default=[],
        help="Optional exact Match Status values to process.",
    )
    args = parser.parse_args()

    if not args.workbook.exists():
        raise SystemExit(f"Workbook does not exist: {args.workbook}")

    headers, rows = load_rows(args.workbook)
    if REALTIME_HEADER not in headers:
        headers.append(REALTIME_HEADER)
        for row in rows:
            row[REALTIME_HEADER] = ""

    products = selected_products(args.products, "nonlegacy")
    products_by_id = {int(product["id"]): product for product in products}
    before = before_after_stats(rows)
    cache = load_cache(args.cache)
    email = os.environ.get("BECKETT_EMAIL", "")
    password = os.environ.get("BECKETT_PASSWORD", "")
    if not email or not password:
        raise SystemExit("Set BECKETT_EMAIL and BECKETT_PASSWORD before running.")
    session = create_session(email, password)

    backup_path = None if args.dry_run else backup_workbook(args.workbook)
    changed: list[dict[str, Any]] = []
    searched = 0
    skipped = Counter()

    candidates_to_process = [
        row
        for row in rows
        if normalize_spaces(row.get("Match Status")) not in MATCHED_STATUSES
        and row.get("Category") in SPORT_CATEGORIES
        and (not args.statuses or normalize_spaces(row.get("Match Status")) in set(args.statuses))
    ]
    if args.limit:
        candidates_to_process = candidates_to_process[: args.limit]

    for index, row in enumerate(candidates_to_process, start=1):
        product = row_product(row, products_by_id)
        if not product:
            skipped["missing product"] += 1
            continue
        searched += 1
        original_status = normalize_spaces(row.get("Match Status"))
        original_title = normalize_spaces(row.get("Beckett Matched Title"))
        candidate, query, score, conflicts = best_enhanced_candidate(
            session,
            product,
            row,
            cache,
            args.delay,
            existing_only=args.existing_only,
        )
        if not candidate:
            skipped["no candidate"] += 1
            continue
        if conflicts:
            # Conflicted candidates are useful for research, but leaving them
            # out of the workbook prevents "review" rows from looking more
            # certain than they are.
            skipped["conflicts"] += 1
            continue

        if score < 0.86:
            skipped["below confidence threshold"] += 1
            continue

        try:
            apply_card_data(row, product, candidate, query, score, session, cache, args.delay)
        except Exception as exc:  # noqa: BLE001 - keep processing remaining rows.
            row["Notes"] = f"Enhanced candidate found, but pricing could not be refreshed: {exc}"
            skipped["pricing fetch failed"] += 1
            continue

        row["Match Status"] = "Matched"
        row["Review Bucket"] = "Enhanced confirmed match"
        changed.append(
            {
                "productId": row.get("Product ID"),
                "action": "promoted-to-matched",
                "fromStatus": original_status,
                "score": score,
                "oldBeckettTitle": original_title,
                "newBeckettTitle": row.get("Beckett Matched Title"),
            }
        )

        if index % args.save_every == 0:
            save_cache(args.cache, cache)
            print(f"Processed {index}/{len(candidates_to_process)} review/unresolved rows...", flush=True)

    after = before_after_stats(rows)
    report = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "workbook": str(args.workbook),
        "backup": str(backup_path) if backup_path else "",
        "dryRun": args.dry_run,
        "searchedRows": searched,
        "changedRows": len(changed),
        "promotedRows": len([item for item in changed if item["action"] == "promoted-to-matched"]),
        "updatedReviewCandidates": len([item for item in changed if item["action"] == "updated-review-candidate"]),
        "skipped": dict(skipped),
        "before": before,
        "after": after,
        "changesSample": changed[:100],
    }

    if not args.dry_run:
        # Rebuild derived workbook tabs from the canonical first sheet rows so
        # Review Needed / Action Queue no longer contain newly confirmed rows.
        build_workbook(rows, args.workbook, "nonlegacy")
        save_cache(args.cache, cache)

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
