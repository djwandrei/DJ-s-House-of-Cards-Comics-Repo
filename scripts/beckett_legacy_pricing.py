#!/usr/bin/env python3
"""Create a Beckett pricing workbook for legacy DJ's House of Cards listings.

The script intentionally reads Beckett credentials from environment variables so
passwords are never written into the repository or the generated workbook.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus, urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BECKETT_BASE = "https://www.beckett.com"
SPORT_CATEGORIES = {"Baseball", "Basketball", "Football"}
SPORT_PATHS = {
    "Baseball": "/baseball/",
    "Basketball": "/basketball/",
    "Football": "/football/",
}

MONEY_RE = re.compile(r"\$(?:\d[\d,]*(?:\.\d{1,2})?|\.\d{1,2})")
YEAR_RE = re.compile(r"\b((?:18|19|20)\d{2}(?:[-/](?:\d{2}|\d{4}))?)\b")
CARD_NO_RE = re.compile(r"(?:#|no\.?\s*)([A-Za-z0-9]+(?:[-/][A-Za-z0-9]+)*)", re.I)
GRADE_RE = re.compile(r"\b(PSA|BGS|SGC|CGC|HGA|BCCG|GAI)\s*([0-9](?:\.\d)?|10)\b", re.I)

NOISE_PATTERNS = [
    r"\bRookie Card\b",
    r"\bRC\b",
    r"\bAuto(?:graph)?\b",
    r"\bAutographed\b",
    r"\bSigned\b",
    r"\bPatch\b",
    r"\bRelic\b",
    r"\bJersey\b",
    r"\bRedemption\b",
    r"\bGame[- ]Used\b",
    r"\bGame[- ]Worn\b",
    r"\bAuthentic\b",
    r"\bBlue Sharpie\b",
    r"\bWith Sharpie\b",
    r"\bSerial Numbered\b",
    r"\bRefractor\b",
    r"\bPrizm\b",
    r"\bChrome\b",
    r"\bInsert\b",
    r"\bParallel\b",
    r"\bPSA\b\s*\d+(?:\.\d)?",
    r"\bBGS\b\s*\d+(?:\.\d)?",
    r"\bSGC\b\s*\d+(?:\.\d)?",
    r"\bCGC\b\s*\d+(?:\.\d)?",
    r"\bHGA\b\s*\d+(?:\.\d)?",
    r"\bVG\b",
    r"\bEX\b",
    r"\bGOOD\b",
    r"\bNM\b",
    r"\bNear Mint\b",
    r"\bMint\b",
    r"\bGraded\b",
    r"\bUngraded\b",
    r"\bGuide range listed\b",
    r"\bComplete Boxed Set\b",
]


@dataclass
class SearchCandidate:
    title: str
    url: str
    query: str
    score: float


def normalize_spaces(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize_key(value: Any) -> str:
    text = normalize_spaces(value).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return normalize_spaces(text)


def expanded_token_set(value: Any) -> set[str]:
    tokens = set(normalize_key(value).split())
    expanded = set(tokens)
    for token in list(tokens):
        if token.endswith("ies") and len(token) > 4:
            expanded.add(token[:-3] + "y")
        elif token.endswith("s") and len(token) > 3:
            expanded.add(token[:-1])
    return expanded


def remove_phrase_tokens(text: str, phrase: str) -> str:
    result = normalize_key(text)
    phrase_key = normalize_key(phrase)
    if not phrase_key:
        return result
    result = re.sub(rf"\b{re.escape(phrase_key)}\b", " ", result)
    for token in phrase_key.split():
        result = re.sub(rf"\b{re.escape(token)}\b", " ", result)
    return normalize_spaces(result)


def clean_listing_title(title: str) -> str:
    text = normalize_spaces(title)
    for pattern in NOISE_PATTERNS:
        text = re.sub(pattern, " ", text, flags=re.I)
    text = text.replace("#", " ")
    text = re.sub(r"\b/\d+\b", " ", text)
    text = re.sub(r"[^\w\s.-]", " ", text)
    return normalize_spaces(text)


def search_title_core(title: str) -> str:
    text = normalize_spaces(title)
    text = re.sub(r"\((?:variant|variation)[^)]+\)", " ", text, flags=re.I)
    text = re.sub(r"\(\s*\d+\s*/\s*\d+\s*\)", " ", text)
    text = re.sub(r"\b\d+\s*/\s*\d+\b", " ", text)
    text = re.sub(r"\bserial(?:ly)?\s+numbered\b", " ", text, flags=re.I)
    text = re.sub(r"\bblue sharpie\b|\bwith sharpie\b", " ", text, flags=re.I)
    text = re.sub(r"[()]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def descriptor_stripped_title(title: str) -> str:
    text = search_title_core(title)
    descriptor_patterns = [
        r"\bRookie Card\b",
        r"\bRookie\b",
        r"\bAutograph(?:s)?\b",
        r"\bSignature(?:s)?\b",
        r"\bSigned\b",
        r"\bGame[- ]Used\b",
        r"\bGame[- ]Worn\b",
        r"\bJersey\b",
        r"\bRelic(?:s)?\b",
        r"\bMemorabilia\b",
        r"\bFramed Mini\b",
        r"\bParallel\b",
        r"\bCard\b",
    ]
    for pattern in descriptor_patterns:
        text = re.sub(pattern, " ", text, flags=re.I)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def price_label(product: dict[str, Any]) -> str:
    if product.get("priceLabel"):
        return str(product["priceLabel"])
    price = product.get("price")
    if isinstance(price, (int, float)):
        return f"${price:,.2f}"
    return ""


def extract_year(product: dict[str, Any]) -> str:
    if product.get("year"):
        return str(product["year"])
    match = YEAR_RE.search(product.get("name", ""))
    return match.group(1) if match else ""


def extract_card_number(title: str) -> str:
    match = CARD_NO_RE.search(title or "")
    if match:
        number = match.group(1)
        return number.strip().upper()
    return ""


def player_guess(product: dict[str, Any]) -> str:
    player = normalize_spaces(product.get("playerAthlete"))
    if player:
        return normalize_spaces(player.split("|")[0].split(";")[0])
    return ""


def set_guess(product: dict[str, Any]) -> str:
    title = normalize_spaces(product.get("name"))
    year = extract_year(product)
    player = player_guess(product)
    card_number = extract_card_number(title)
    work = title
    if year:
        idx = work.lower().find(str(year).lower())
        if idx >= 0:
            work = work[idx + len(str(year)) :]
    work = re.sub(r"^\s*[-/]\d{2,4}\b", " ", work)
    end_positions = []
    if player:
        idx = work.lower().find(player.lower())
        if idx >= 0:
            end_positions.append(idx)
    if card_number:
        idx = work.find("#" + card_number)
        if idx >= 0:
            end_positions.append(idx)
    if end_positions:
        work = work[: min(end_positions)]
    work = clean_listing_title(work)
    work = re.sub(r"\b(Baseball|Basketball|Football|Card|Cards)\b", " ", work, flags=re.I)
    work = normalize_spaces(work)
    if len(work.split()) > 6:
        return ""
    return work


def query_variants(product: dict[str, Any]) -> list[str]:
    title = normalize_spaces(product.get("name"))
    year = extract_year(product)
    player = player_guess(product)
    set_name = set_guess(product)
    card_number = extract_card_number(title)
    cleaned = clean_listing_title(title)
    title_core = search_title_core(title)
    title_no_card = normalize_spaces(re.sub(r"\bcard\b", " ", title_core, flags=re.I))
    simplified_title = descriptor_stripped_title(title)
    values: list[str] = []

    def add(parts: list[str]) -> None:
        query = normalize_spaces(" ".join(part for part in parts if part))
        if query and query not in values:
            values.append(query)

    add([title_core])
    add([title_no_card])
    add([simplified_title])
    add([year, set_name, card_number, player])
    add([year, set_name, player, card_number])
    add([year, set_name, card_number])
    add([set_name, card_number, player])
    add([year, simplified_title])
    add([year, player, card_number])
    add([player, year, set_name, card_number])
    add([set_name, player, card_number])
    add([year, set_name, player])
    add([set_name, player])
    add([year, set_name])
    add([player, card_number])
    add([year, player])
    add([player])
    add([cleaned])
    add([cleaned])
    return [q for q in values if len(q) >= 4][:10]


def load_cache(cache_path: Path) -> dict[str, Any]:
    if cache_path.exists():
        try:
            return json.loads(cache_path.read_text(encoding="utf-8"))
        except Exception:
            backup = cache_path.with_suffix(f".corrupt-{datetime.now():%Y%m%d%H%M%S}.json")
            cache_path.replace(backup)
    return {"searches": {}, "cards": {}, "completed": {}}


def save_cache(cache_path: Path, cache: dict[str, Any]) -> None:
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = cache_path.with_suffix(".tmp")
    tmp.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(cache_path)


def create_session(email: str, password: str) -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/123.0 Safari/537.36",
            "Accept-Language": "en-US,en;q=0.9",
        }
    )
    login_page = session.get(f"{BECKETT_BASE}/login", timeout=30)
    login_page.raise_for_status()
    soup = BeautifulSoup(login_page.text, "html.parser")
    token_el = soup.find("input", {"name": "login_token"})
    token = token_el.get("value", "") if token_el else ""
    payload = {
        "login_token": token,
        "email": email,
        "password": password,
        "remember_me": "1",
        "redirect_url": "",
    }
    response = session.post(f"{BECKETT_BASE}/login/", data=payload, timeout=30, allow_redirects=True)
    response.raise_for_status()
    proof = response.text.lower()
    cookie_keys = list(session.cookies.get_dict().keys())
    has_auth_cookie = any(key.startswith("prod_") for key in cookie_keys)
    if not has_auth_cookie and "logout" not in proof and "my account" not in proof and "is_user_logged_in = '1'" not in proof:
        check = session.get(f"{BECKETT_BASE}/search/?term=Hank+Aaron+1954+Topps+128", timeout=30)
        check.raise_for_status()
        proof = check.text.lower()
        has_auth_cookie = any(key.startswith("prod_") for key in session.cookies.get_dict().keys())
    if not has_auth_cookie and "3609933" not in proof and "logout" not in proof and "my account" not in proof:
        raise RuntimeError("Beckett login did not appear to succeed.")
    return session


def search_beckett(session: requests.Session, query: str, cache: dict[str, Any], delay: float) -> list[dict[str, str]]:
    searches = cache.setdefault("searches", {})
    if query in searches:
        return searches[query]
    url = f"{BECKETT_BASE}/search/?term={quote_plus(query)}"
    response = session.get(url, timeout=45)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    candidates: list[dict[str, str]] = []
    seen: set[str] = set()
    for anchor in soup.find_all("a", href=True):
        href = urljoin(BECKETT_BASE, anchor.get("href", ""))
        if not re.search(r"beckett\.com/(baseball|basketball|football)/", href, flags=re.I):
            continue
        if not re.search(r"-\d+(?:\?.*)?$", href):
            continue
        title = normalize_spaces(anchor.get_text(" ", strip=True))
        if not title:
            title = normalize_spaces(href.rsplit("/", 1)[-1].replace("-", " "))
        if href in seen:
            continue
        seen.add(href)
        candidates.append({"title": title, "url": href})
    searches[query] = candidates
    if delay:
        time.sleep(delay)
    return candidates


def score_candidate(product: dict[str, Any], query: str, candidate: dict[str, str]) -> float:
    product_title = normalize_key(clean_listing_title(product.get("name", "")))
    candidate_text = normalize_key(candidate.get("title", "") + " " + candidate.get("url", ""))
    candidate_raw = normalize_spaces(candidate.get("title", "") + " " + candidate.get("url", ""))
    if not product_title or not candidate_text:
        return 0.0

    category = product.get("category")
    expected_path = SPORT_PATHS.get(category, "").lower()
    if expected_path and expected_path not in candidate.get("url", "").lower():
        return 0.0

    p_tokens = set(product_title.split())
    c_tokens = set(candidate_text.split())
    overlap = len(p_tokens & c_tokens) / max(1, len(p_tokens))
    ratio = SequenceMatcher(None, product_title, candidate_text).ratio()
    score = (overlap * 0.45) + (ratio * 0.25)

    year = extract_year(product)
    if year and year in candidate_text:
        score += 0.16
    elif year:
        score -= 0.12
        expected_year = year.split("-")[0].split("/")[0]
        candidate_years = YEAR_RE.findall(candidate_raw)
        if candidate_years and not any(str(candidate_year).startswith(expected_year) for candidate_year in candidate_years):
            score -= 0.18

    card_number = extract_card_number(product.get("name", ""))
    if card_number:
        number_key = re.escape(card_number.lower())
        number_matches = re.search(rf"(?:^|[^0-9a-z]){number_key}[a-z]?(?:[^0-9a-z]|$)", candidate_text)
        if number_matches:
            score += 0.22
        elif re.search(rf"/{number_key.lower()}[a-z]?-", candidate.get("url", "").lower()):
            score += 0.22
        else:
            score -= 0.05

    player = normalize_key(player_guess(product))
    if player:
        player_tokens = set(player.split())
        if player_tokens and player_tokens.issubset(c_tokens):
            score += 0.18
        elif player_tokens & c_tokens:
            score += 0.08
        else:
            score -= 0.08

    set_name = normalize_key(set_guess(product))
    if set_name:
        set_tokens = set(set_name.split())
        if set_tokens and set_tokens.issubset(c_tokens):
            score += 0.12
        elif set_tokens & c_tokens:
            score += 0.05

    if normalize_key(query) == normalize_key(candidate.get("title")):
        score += 0.05
    return round(max(0.0, min(1.0, score)), 4)


def best_candidate(session: requests.Session, product: dict[str, Any], cache: dict[str, Any], delay: float) -> SearchCandidate | None:
    best: SearchCandidate | None = None
    for query in query_variants(product):
        try:
            candidates = search_beckett(session, query, cache, delay)
        except Exception as exc:
            cache.setdefault("search_errors", {})[query] = str(exc)
            continue
        for candidate in candidates:
            score = score_candidate(product, query, candidate)
            if best is None or score > best.score:
                best = SearchCandidate(
                    title=candidate["title"],
                    url=candidate["url"],
                    query=query,
                    score=score,
                )
        if best and best.score >= 0.86:
            break
    return best


def parse_price_rows(rows: list[list[str]]) -> list[dict[str, str]]:
    maps: list[dict[str, str]] = []
    for idx, row in enumerate(rows):
        price_cells = [MONEY_RE.search(cell).group(0) for cell in row if MONEY_RE.search(cell)]
        if len(price_cells) < 2:
            continue
        labels: list[str] = []
        if idx > 0:
            labels = [cell for cell in rows[idx - 1] if not MONEY_RE.search(cell)]
        if len(labels) != len(price_cells):
            non_prices = [cell for cell in row if not MONEY_RE.search(cell)]
            if len(non_prices) == len(price_cells):
                labels = non_prices
        if len(labels) == len(price_cells):
            maps.append({normalize_spaces(label): price for label, price in zip(labels, price_cells)})

    if maps:
        return maps

    flat = [cell for row in rows for cell in row if cell]
    first_price = next((idx for idx, cell in enumerate(flat) if MONEY_RE.search(cell)), -1)
    if first_price > 0:
        labels = [normalize_spaces(cell) for cell in flat[:first_price]]
        prices = [MONEY_RE.search(cell).group(0) for cell in flat[first_price:] if MONEY_RE.search(cell)]
        if len(labels) >= len(prices):
            labels = labels[-len(prices) :]
        if len(labels) == len(prices) and len(prices) >= 2:
            maps.append(dict(zip(labels, prices)))
    return maps


def classify_price_maps(maps: list[dict[str, str]]) -> tuple[dict[str, str], dict[str, str]]:
    raw: dict[str, str] = {}
    graded: dict[str, str] = {}
    for price_map in maps:
        labels = " ".join(price_map.keys()).lower()
        if re.search(r"\(\s*(10|9|8|7|6|5|4|3|2|1)\s*\)", labels):
            graded.update(price_map)
        elif any(label.upper() in labels.upper() for label in ["NM-MT", "VG-EX", "MINT", "GOOD", "POOR"]):
            raw.update(price_map)
        elif not raw:
            raw.update(price_map)
    return raw, graded


def fetch_card_page(session: requests.Session, url: str, cache: dict[str, Any], delay: float) -> dict[str, Any]:
    cards = cache.setdefault("cards", {})
    if url in cards:
        return cards[url]
    response = session.get(url, timeout=45)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    title_el = soup.find("h1")
    title = normalize_spaces(title_el.get_text(" ", strip=True)) if title_el else ""
    source_page_title = ""
    for tr in soup.find_all("tr"):
        cells = [normalize_spaces(cell.get_text(" ", strip=True)) for cell in tr.find_all(["th", "td"])]
        if len(cells) >= 2 and cells[0].lower().startswith("source"):
            source_page_title = cells[1]
            break
    rows_by_table: list[list[list[str]]] = []
    for table in soup.find_all("table"):
        rows: list[list[str]] = []
        for tr in table.find_all("tr"):
            cells = [normalize_spaces(cell.get_text(" ", strip=True)) for cell in tr.find_all(["th", "td"])]
            cells = [cell for cell in cells if cell]
            if cells:
                rows.append(cells)
        if any(MONEY_RE.search(cell) for row in rows for cell in row):
            rows_by_table.append(rows)
    maps: list[dict[str, str]] = []
    for rows in rows_by_table:
        maps.extend(parse_price_rows(rows))
    raw, graded = classify_price_maps(maps)
    raw_market = {}
    graded_market = {}
    item_id_match = re.search(r"/(\d+)(?:\?.*)?$", url)
    if item_id_match:
        item_id = item_id_match.group(1)
        raw_market = fetch_market_range(session, item_id, "1", delay)
        graded_market = fetch_market_range(session, item_id, "2", delay)
    card_data = {
        "title": title,
        "source_page_title": source_page_title,
        "raw_prices": raw,
        "graded_prices": graded,
        "raw_market": raw_market,
        "graded_market": graded_market,
        "fetched_at": datetime.now().isoformat(timespec="seconds"),
    }
    cards[url] = card_data
    if delay:
        time.sleep(delay)
    return card_data


def fetch_market_range(session: requests.Session, item_id: str, ctype: str, delay: float) -> dict[str, Any]:
    endpoint = f"{BECKETT_BASE}/pgs_search/ajax_mkt_get_list/{item_id}/0/desc/grade/{ctype}/item"
    response = session.get(endpoint, timeout=45)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    prices = [money_to_float(td.get_text(" ", strip=True) or "") for td in soup.find_all("td")]
    prices = [price for price in prices if price is not None]
    result = {
        "count": len(prices),
        "range": f"${min(prices):,.2f}-${max(prices):,.2f}" if len(prices) >= 2 else (f"${prices[0]:,.2f}" if len(prices) == 1 else ""),
    }
    if delay:
        time.sleep(delay)
    return result


def money_to_float(value: str) -> float | None:
    match = MONEY_RE.search(str(value or ""))
    if not match:
        return None
    return float(match.group(0).replace("$", "").replace(",", ""))


def price_range(price_map: dict[str, str]) -> str:
    values = [money_to_float(value) for value in price_map.values()]
    values = [value for value in values if value is not None]
    if not values:
        return ""
    low, high = min(values), max(values)
    if low == high:
        return f"${low:,.2f}"
    return f"${low:,.2f}-${high:,.2f}"


def parse_money_values(text: str) -> list[float]:
    return [float(match.group(0).replace("$", "").replace(",", "")) for match in MONEY_RE.finditer(str(text or ""))]


def midpoint_from_range(text: str) -> float | None:
    values = parse_money_values(text)
    if not values:
        return None
    if len(values) == 1:
        return values[0]
    return (min(values) + max(values)) / 2


def format_money(value: float | None) -> str:
    return f"${value:,.2f}" if value is not None else ""


def format_percent(value: float | None) -> str:
    return f"{value:.1f}%" if value is not None else ""


def price_signal(site_mid: float | None, beckett_mid: float | None) -> tuple[str, str, str]:
    if site_mid is None and beckett_mid is None:
        return "", "", "No midpoint comparison"
    if site_mid is None:
        return "", "", "Site price not numeric"
    if beckett_mid is None:
        return "", "", "Beckett price unavailable"
    delta = beckett_mid - site_mid
    pct = (delta / site_mid * 100) if site_mid else None
    if pct is None:
        signal = "No midpoint comparison"
    elif pct >= 50:
        signal = "Beckett much higher"
    elif pct >= 15:
        signal = "Beckett higher"
    elif pct <= -50:
        signal = "Beckett much lower"
    elif pct <= -15:
        signal = "Beckett lower"
    else:
        signal = "Close"
    return format_money(delta), format_percent(pct), signal


def is_graded(condition: str) -> bool:
    return bool(GRADE_RE.search(condition or ""))


def raw_grade_basis(condition: str) -> str:
    text = (condition or "").upper()
    ordered = [
        ("NM-MT+", ["NM-MT+", "NEAR MINT-MINT+"]),
        ("NM-MT", ["NM-MT", "NEAR MINT-MINT"]),
        ("MINT", ["MINT", "MT"]),
        ("NM", ["NEAR MINT", "NRMT", "NM"]),
        ("EX-MT", ["EX-MT", "EXMT"]),
        ("EX", ["EXCELLENT", " EX ", "EX"]),
        ("VG-EX", ["VG-EX", "VGEX"]),
        ("VG", ["VERY GOOD", " VG ", "VG"]),
        ("GOOD", ["GOOD"]),
        ("POOR", ["POOR"]),
    ]
    padded = f" {text} "
    for label, needles in ordered:
        if any(needle in padded or needle in text for needle in needles):
            return label
    return ""


def matched_grade_price(condition: str, raw_prices: dict[str, str], graded_prices: dict[str, str]) -> tuple[str, str, str]:
    condition = condition or ""
    graded_match = GRADE_RE.search(condition)
    if graded_match and graded_prices:
        grade_value = float(graded_match.group(2))
        exact_labels = [label for label in graded_prices if re.search(rf"\(\s*{int(grade_value)}\s*\)", label)]
        if grade_value.is_integer() and exact_labels:
            label = exact_labels[0]
            return graded_prices[label], label, ""
        lower_labels: list[tuple[float, str]] = []
        for label in graded_prices:
            match = re.search(r"\(\s*(\d+)\s*\)", label)
            if match:
                lower_labels.append((float(match.group(1)), label))
        if lower_labels:
            lower_labels.sort(reverse=True)
            nearest = next((item for item in lower_labels if item[0] <= grade_value), lower_labels[-1])
            return graded_prices[nearest[1]], nearest[1], f"Exact {grade_value:g} grade not listed; used nearest Beckett grade {nearest[1]}."

    raw_basis = raw_grade_basis(condition)
    if raw_basis and raw_prices:
        for label, price in raw_prices.items():
            if normalize_key(label) == normalize_key(raw_basis):
                return price, label, ""
        for label, price in raw_prices.items():
            if raw_basis.replace("-", " ") in normalize_key(label):
                return price, label, ""
    return "", "", ""


def is_legacy_product(product: dict[str, Any]) -> bool:
    return bool(
        product.get("sourcePage")
        or product.get("legacyImageLabel")
        or "legacy" in str(product.get("metadata", "")).lower()
    )


def selected_products(products_path: Path, scope: str) -> list[dict[str, Any]]:
    data = json.loads(products_path.read_text(encoding="utf-8"))
    if scope == "legacy":
        return [product for product in data if is_legacy_product(product)]
    if scope == "nonlegacy":
        return [product for product in data if not is_legacy_product(product)]
    return list(data)


def year_mismatch(product: dict[str, Any], matched_title: str, matched_url: str) -> bool:
    year = extract_year(product)
    if not year:
        return False
    expected_year = year.split("-")[0].split("/")[0]
    candidate_years = YEAR_RE.findall(f"{matched_title} {matched_url}")
    if not candidate_years:
        return False
    return not any(str(candidate_year).startswith(expected_year) for candidate_year in candidate_years)


def set_mismatch_note(product: dict[str, Any], matched_title: str, matched_url: str) -> str:
    guessed = normalize_key(set_guess(product))
    if not guessed:
        return ""
    source_tokens = {
        token
        for token in expanded_token_set(guessed)
        if token
        and token
        not in {
            "topps",
            "bowman",
            "upper",
            "deck",
            "panini",
            "donruss",
            "fleer",
            "leaf",
            "press",
            "pass",
            "sage",
            "score",
            "hoops",
            "select",
            "optic",
            "prizm",
            "chrome",
            "cards",
            "card",
            "mini",
            "framed",
            "series",
            "edition",
            "collection",
            "select",
            "choice",
            "mark",
            "marks",
            "signature",
            "signatures",
            "autograph",
            "autographs",
            "auto",
            "autos",
            "relic",
            "relics",
            "memorabilia",
            "jersey",
            "jerseys",
            "swatch",
            "swatches",
            "patch",
            "patches",
            "thread",
            "threads",
            "authentic",
            "authentics",
            "future",
            "phenom",
            "phenoms",
            "rookie",
            "rookies",
            "prospect",
            "prospects",
            "parallel",
            "refactor",
            "refractor",
            "refractors",
            "gold",
            "silver",
            "blue",
            "red",
            "green",
            "black",
            "purple",
            "orange",
            "pink",
            "sepia",
            "holo",
            "holofoil",
            "bronze",
            "platinum",
            "prime",
            "booklet",
            "modern",
            "paint",
            "fresh",
            "signing",
            "signings",
            "inscription",
            "inscriptions",
            "letter",
            "letters",
            "campus",
            "id",
            "time",
            "shine",
            "baseball",
            "basketball",
            "football",
        }
    }
    if len(source_tokens) < 2:
        return ""
    target_tokens = expanded_token_set(f"{matched_title} {matched_url}")
    overlap = len(source_tokens & target_tokens) / len(source_tokens)
    if overlap < 0.34:
        return "Beckett result appears to come from a different set than the site title."
    return ""


def type_mismatch_notes(product: dict[str, Any], matched_title: str, matched_url: str) -> list[str]:
    source = product.get("name", "")
    target = f"{matched_title} {matched_url}"
    notes: list[str] = []
    target_key = normalize_key(target)
    checks = [
        ({"bat", "bats"}, {"bat", "bats"}, "Site title says bat, but the Beckett result does not clearly indicate a bat card."),
        (
            {"patch", "patches"},
            {"patch", "patches", "prime", "swatch", "swatches", "letter", "letters", "jsy", "mem"},
            "Site title says patch, but the Beckett result does not clearly indicate a patch card.",
        ),
        (
            {"jersey", "jerseys"},
            {
                "jersey",
                "jerseys",
                "jsy",
                "uniform",
                "uniforms",
                "uni",
                "swatch",
                "swatches",
                "thread",
                "threads",
                "fabric",
                "fabrics",
                "material",
                "materials",
                "mem",
                "relic",
                "relics",
                "memorabilia",
                "letter",
                "letters",
                "caps",
                "authentics",
                "helmet",
                "helmets",
                "glove",
                "gloves",
                "ball",
                "balls",
                "manufactured",
            },
            "Site title says jersey, but the Beckett result does not clearly indicate a jersey/uniform card.",
        ),
        (
            {"autograph", "autographs", "signature", "signatures", "signed", "auto", "autos"},
            {
                "autograph",
                "autographs",
                "auto",
                "autos",
                "au",
                "signature",
                "signatures",
                "signed",
                "sigs",
                "sweet",
                "ink",
                "playergraphs",
                "graphs",
                "chirography",
                "inscription",
                "inscriptions",
                "signing",
                "signings",
                "lettermen",
            },
            "Site title says autograph/signature, but the Beckett result does not clearly indicate an autograph card.",
        ),
        (
            {"checklist", "checklists"},
            {"checklist", "checklists"},
            "Site title says checklist, but the Beckett result does not clearly indicate a checklist card.",
        ),
        (
            {"relic", "relics", "memorabilia"},
            {
                "relic",
                "relics",
                "memorabilia",
                "bat",
                "bats",
                "jersey",
                "jerseys",
                "uniform",
                "uniforms",
                "uni",
                "patch",
                "patches",
                "swatch",
                "swatches",
                "thread",
                "threads",
                "fabric",
                "fabrics",
                "material",
                "materials",
                "mem",
                "letter",
                "letters",
                "jsy",
                "caps",
                "authentics",
                "helmet",
                "helmets",
                "glove",
                "gloves",
                "ball",
                "balls",
                "manufactured",
            },
            "Site title says relic/memorabilia, but the Beckett result does not clearly indicate memorabilia.",
        ),
    ]
    source_tokens = expanded_token_set(source)
    target_tokens = expanded_token_set(target)
    phrase_overrides = [
        (
            {"autograph", "autographs", "signature", "signatures", "signed", "auto", "autos"},
            [
                "autograph",
                "autographs",
                "by the letter",
                "signatures",
                "signings",
                "inscriptions",
                "modern marks",
                "marks of brilliance",
                "fresh paint",
                "scripts",
                "rookie signatures",
                "lettermen autographs",
                "choice au",
                " pc au",
                "time to shine",
            ],
        ),
        (
            {"jersey", "jerseys", "relic", "relics", "memorabilia"},
            [
                "jsy",
                "jersey",
                "relic",
                "relics",
                "swatch",
                "swatches",
                "patch",
                "patches",
                "memorabilia",
                "caps",
                "authentics",
                "helmet",
                "glove",
                "ball",
                "material",
                "materials",
                "manufactured patch",
            ],
        ),
    ]
    for expected, accepted, note in checks:
        if not (source_tokens & expected):
            continue
        if target_tokens & accepted:
            continue
        if any(source_tokens & override_expected and any(phrase in target_key for phrase in phrases) for override_expected, phrases in phrase_overrides):
            continue
        if source_tokens & expected and not (target_tokens & accepted):
            notes.append(note)
    return notes


def variant_mismatch_notes(product: dict[str, Any], matched_title: str, matched_url: str) -> list[str]:
    player = normalize_spaces(product.get("playerAthlete", ""))
    source = remove_phrase_tokens(product.get("name", ""), player)
    target = remove_phrase_tokens(f"{matched_title} {matched_url}", player)
    source_key = normalize_key(source)
    target_key = normalize_key(target)
    source_tokens = expanded_token_set(source)
    target_tokens = expanded_token_set(target)

    def has_variant(text_key: str, token_set: set[str], terms: set[str]) -> bool:
        for term in terms:
            cleaned = normalize_key(term)
            if " " in cleaned:
                if cleaned and cleaned in text_key:
                    return True
            elif cleaned in token_set:
                return True
        return False

    variant_checks = [
        ({"tiffany"}, {"tiffany"}, "Beckett result includes the Tiffany variant, which is not named in the site title."),
        ({"class 1"}, {"class 1", "c1"}, "Beckett result includes Class 1 wording not shown in the site title."),
        ({"class 2"}, {"class 2", "c2"}, "Beckett result includes Class 2 wording not shown in the site title."),
        ({"class 3"}, {"class 3", "c3"}, "Beckett result includes Class 3 wording not shown in the site title."),
        (
            {"refractor", "refractors", "xfractor", "xfractors"},
            {"refractor", "refractors", "xfractor", "xfractors"},
            "Beckett result includes refractor wording not shown in the site title.",
        ),
        ({"prizm", "prizms"}, {"prizm", "prizms"}, "Beckett result includes Prizm wording not shown in the site title."),
        ({"mojo"}, {"mojo"}, "Beckett result includes Mojo wording not shown in the site title."),
        ({"pulsar"}, {"pulsar"}, "Beckett result includes Pulsar wording not shown in the site title."),
        ({"shimmer"}, {"shimmer"}, "Beckett result includes Shimmer wording not shown in the site title."),
        ({"lava"}, {"lava"}, "Beckett result includes Lava wording not shown in the site title."),
        ({"speckle"}, {"speckle"}, "Beckett result includes Speckle wording not shown in the site title."),
        ({"cracked ice"}, {"cracked ice"}, "Beckett result includes Cracked Ice wording not shown in the site title."),
        ({"wave"}, {"wave"}, "Beckett result includes Wave wording not shown in the site title."),
        ({"foil"}, {"foil"}, "Beckett result includes Foil wording not shown in the site title."),
        ({"glitter"}, {"glitter"}, "Beckett result includes Glitter wording not shown in the site title."),
        ({"black"}, {"black"}, "Beckett result includes the Black variant, which is not named in the site title."),
        ({"blue"}, {"blue"}, "Beckett result includes the Blue variant, which is not named in the site title."),
        ({"red"}, {"red"}, "Beckett result includes the Red variant, which is not named in the site title."),
        ({"green"}, {"green"}, "Beckett result includes the Green variant, which is not named in the site title."),
        ({"gold"}, {"gold"}, "Beckett result includes the Gold variant, which is not named in the site title."),
        ({"silver"}, {"silver"}, "Beckett result includes the Silver variant, which is not named in the site title."),
        ({"purple"}, {"purple"}, "Beckett result includes the Purple variant, which is not named in the site title."),
        ({"orange"}, {"orange"}, "Beckett result includes the Orange variant, which is not named in the site title."),
        ({"pink"}, {"pink"}, "Beckett result includes the Pink variant, which is not named in the site title."),
    ]
    notes: list[str] = []
    for target_terms, source_terms, note in variant_checks:
        if has_variant(target_key, target_tokens, target_terms) and not has_variant(source_key, source_tokens, source_terms):
            notes.append(note)
    return notes


def lookup_product(
    session: requests.Session | None,
    product: dict[str, Any],
    cache: dict[str, Any],
    delay: float,
    refresh_nonmatched: bool = False,
) -> dict[str, Any]:
    completed = cache.setdefault("completed", {})
    product_key = str(product.get("id"))
    if product_key in completed:
        existing = completed[product_key]
        if not refresh_nonmatched or existing.get("Match Status") in {
            "Matched",
            "Not searched - non sports-card legacy row",
        }:
            return existing
        completed.pop(product_key, None)

    row = {
        "Product ID": product.get("id"),
        "Category": product.get("category", ""),
        "Source Page": product.get("sourcePage", ""),
        "Title": product.get("name", ""),
        "Original Price / Range": price_label(product),
        "Current Grade / Condition": product.get("condition", ""),
        "Site Team / Publisher": product.get("team", ""),
        "Site Year": product.get("year", ""),
        "Site Player / Athlete": product.get("playerAthlete", ""),
        "Lookup Query": "",
        "Match Status": "",
        "Review Bucket": "",
        "Match Confidence": "",
        "Beckett Matched Title": "",
        "Beckett URL": "",
        "Beckett Price Range": "",
        "Beckett Raw Range": "",
        "Beckett Graded Range": "",
        "Beckett Raw Market Range": "",
        "Beckett Graded Market Range": "",
        "Beckett Availability": "",
        "Site Price Midpoint": "",
        "Beckett Price Midpoint": "",
        "Midpoint Delta": "",
        "Midpoint Delta %": "",
        "Price Signal": "No midpoint comparison",
        "Beckett Page Source": "",
        "Beckett Matched Grade Price": "",
        "Beckett Matched Grade Basis": "",
        "Notes": "",
    }

    if product.get("category") not in SPORT_CATEGORIES:
        row["Match Status"] = "Not searched - non sports-card row"
        row["Review Bucket"] = "Non-sports"
        row["Beckett Availability"] = "Not searched"
        row["Notes"] = "Non-sports item included for completeness; Beckett sports-card OPG lookup was not applied."
        completed[product_key] = row
        return row

    if session is None:
        row["Match Status"] = "Skipped - no Beckett session"
        row["Review Bucket"] = "Session issue"
        row["Beckett Availability"] = "Session issue"
        completed[product_key] = row
        return row

    best = best_candidate(session, product, cache, delay)
    if not best:
        row["Match Status"] = "No Beckett result found"
        row["Review Bucket"] = "No Beckett result"
        row["Beckett Availability"] = "No match assigned"
        row["Lookup Query"] = " | ".join(query_variants(product))
        completed[product_key] = row
        return row

    if best.score < 0.55:
        row["Lookup Query"] = best.query
        row["Match Confidence"] = best.score
        row["Match Status"] = "No confident Beckett match"
        row["Review Bucket"] = "Weak search result"
        row["Beckett Availability"] = "No match assigned"
        row["Notes"] = "Search returned only weak Beckett candidates, so no pricing link was assigned automatically."
        completed[product_key] = row
        return row

    row["Lookup Query"] = best.query
    row["Match Confidence"] = best.score
    row["Beckett URL"] = best.url
    row["Beckett Matched Title"] = best.title
    if best.score >= 0.82:
        row["Match Status"] = "Matched"
    elif best.score >= 0.58:
        row["Match Status"] = "Low confidence match - review"
    else:
        row["Match Status"] = "Uncertain match - review"
    row["Review Bucket"] = "Exact/near match" if row["Match Status"] == "Matched" else "Needs review"

    try:
        card = fetch_card_page(session, best.url, cache, delay)
    except Exception as exc:
        row["Notes"] = f"Matched search result, but card page pricing could not be parsed: {exc}"
        completed[product_key] = row
        return row

    if card.get("title"):
        row["Beckett Matched Title"] = card["title"]
    raw_prices = card.get("raw_prices", {})
    graded_prices = card.get("graded_prices", {})
    raw_market = card.get("raw_market", {})
    graded_market = card.get("graded_market", {})
    row["Beckett Page Source"] = card.get("source_page_title", "")
    raw_range = price_range(raw_prices)
    graded_range = price_range(graded_prices)
    row["Beckett Raw Market Range"] = raw_market.get("range", "")
    row["Beckett Graded Market Range"] = graded_market.get("range", "")
    matched_price, matched_basis, grade_note = matched_grade_price(
        str(product.get("condition", "")),
        raw_prices,
        graded_prices,
    )
    row["Beckett Raw Range"] = raw_range
    row["Beckett Graded Range"] = graded_range
    row["Beckett Price Range"] = graded_range if is_graded(str(product.get("condition", ""))) and graded_range else raw_range
    row["Beckett Matched Grade Price"] = matched_price
    row["Beckett Matched Grade Basis"] = matched_basis
    if raw_range or graded_range:
        row["Beckett Availability"] = "Guide pricing"
    elif row["Beckett Raw Market Range"] or row["Beckett Graded Market Range"]:
        row["Beckett Availability"] = "Market sales only"
    else:
        row["Beckett Availability"] = "No visible pricing"
    site_mid = midpoint_from_range(row["Original Price / Range"])
    beckett_mid = midpoint_from_range(row["Beckett Price Range"])
    delta_money, delta_pct, signal = price_signal(site_mid, beckett_mid)
    row["Site Price Midpoint"] = format_money(site_mid)
    row["Beckett Price Midpoint"] = format_money(beckett_mid)
    row["Midpoint Delta"] = delta_money
    row["Midpoint Delta %"] = delta_pct
    row["Price Signal"] = signal
    notes = []
    if year_mismatch(product, row["Beckett Matched Title"], row["Beckett URL"]):
        if row["Match Status"] == "Matched":
            row["Match Status"] = "Low confidence match - review"
            row["Review Bucket"] = "Year mismatch"
        notes.append("Matched Beckett result does not share the same year shown in the site listing; review the card carefully.")
    set_note = set_mismatch_note(product, row["Beckett Matched Title"], row["Beckett URL"])
    if set_note:
        if row["Match Status"] == "Matched":
            row["Match Status"] = "Low confidence match - review"
            row["Review Bucket"] = "Set mismatch"
        notes.append(set_note)
    mismatch_notes = type_mismatch_notes(product, row["Beckett Matched Title"], row["Beckett URL"])
    if mismatch_notes and row["Match Status"] == "Matched":
        row["Match Status"] = "Low confidence match - review"
        row["Review Bucket"] = "Type mismatch"
    notes.extend(mismatch_notes)
    variant_notes = variant_mismatch_notes(product, row["Beckett Matched Title"], row["Beckett URL"])
    if variant_notes and row["Match Status"] == "Matched":
        row["Match Status"] = "Low confidence match - review"
        row["Review Bucket"] = "Variant mismatch"
    notes.extend(variant_notes)
    if grade_note:
        notes.append(grade_note)
    if row["Beckett Page Source"] and normalize_key(row["Beckett Page Source"]) != normalize_key(row["Beckett Matched Title"]):
        notes.append("Beckett page source label differs from the matched title shown on the page.")
    if not raw_range and not graded_range:
        if row["Beckett Raw Market Range"] or row["Beckett Graded Market Range"]:
            notes.append("Guide pricing was not visible, but Beckett market-sales data was captured.")
        else:
            notes.append("Beckett page matched, but no visible pricing table was parsed.")
    if row["Price Signal"] == "Site price not numeric" and row["Original Price / Range"]:
        notes.append("Site price text was not numeric, so midpoint comparison was skipped.")
    elif row["Price Signal"] == "Beckett price unavailable":
        notes.append("No Beckett guide price range was available for midpoint comparison.")
    if row["Match Status"] != "Matched":
        notes.append("Review match before relying on the price.")
        if row["Review Bucket"] == "Needs review":
            row["Review Bucket"] = "General review"
    row["Notes"] = " ".join(notes)
    completed[product_key] = row
    return row


def build_workbook(rows: list[dict[str, Any]], output_path: Path, scope: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    headers = list(rows[0].keys()) if rows else []
    header_fill = PatternFill("solid", fgColor="14213D")
    header_font = Font(color="FFFFFF", bold=True)
    widths = {
        "A": 12,
        "B": 14,
        "C": 18,
        "D": 48,
        "E": 20,
        "F": 22,
        "G": 22,
        "H": 12,
        "I": 24,
        "J": 35,
        "K": 20,
        "L": 14,
        "M": 42,
        "N": 55,
        "O": 22,
        "P": 22,
        "Q": 22,
        "R": 22,
        "S": 22,
        "T": 20,
        "U": 18,
        "V": 18,
        "W": 18,
        "X": 16,
        "Y": 20,
        "Z": 32,
        "AA": 24,
        "AB": 52,
    }

    def safe_float(value: Any, default: float = 0.0) -> float:
        """Keep workbook rebuilds resilient when legacy review cells contain text like SOLD."""
        try:
            return float(str(value).replace("%", "").strip() or default)
        except (TypeError, ValueError):
            return default

    def populate_sheet(ws, data_rows: list[dict[str, Any]]) -> None:
        ws.append(headers)
        for row in data_rows:
            ws.append([row.get(header, "") for header in headers])
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        url_col = headers.index("Beckett URL") + 1 if "Beckett URL" in headers else None
        status_col = headers.index("Match Status") + 1 if "Match Status" in headers else None
        bucket_col = headers.index("Review Bucket") + 1 if "Review Bucket" in headers else None
        price_signal_col = headers.index("Price Signal") + 1 if "Price Signal" in headers else None
        for row_idx in range(2, ws.max_row + 1):
            if url_col:
                cell = ws.cell(row=row_idx, column=url_col)
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
            if status_col:
                status = str(ws.cell(row=row_idx, column=status_col).value or "")
                fill = None
                if status == "Matched":
                    fill = PatternFill("solid", fgColor="D9EAD3")
                elif "Low confidence" in status or "Uncertain" in status:
                    fill = PatternFill("solid", fgColor="FFF2CC")
                elif "No Beckett" in status or "No confident" in status or "could not" in status:
                    fill = PatternFill("solid", fgColor="FCE5CD")
                elif "Not searched" in status:
                    fill = PatternFill("solid", fgColor="D9EAF7")
                if fill:
                    ws.cell(row=row_idx, column=status_col).fill = fill
            if bucket_col:
                bucket = str(ws.cell(row=row_idx, column=bucket_col).value or "")
                fill = None
                if bucket in {"Variant mismatch", "Set mismatch", "Year mismatch", "Type mismatch"}:
                    fill = PatternFill("solid", fgColor="FCE5CD")
                elif bucket in {"Weak search result", "No Beckett result", "General review"}:
                    fill = PatternFill("solid", fgColor="FFF2CC")
                elif bucket in {"Exact/near match"}:
                    fill = PatternFill("solid", fgColor="D9EAD3")
                elif bucket in {"Non-sports"}:
                    fill = PatternFill("solid", fgColor="D9EAF7")
                if fill:
                    ws.cell(row=row_idx, column=bucket_col).fill = fill
            if price_signal_col:
                signal = str(ws.cell(row=row_idx, column=price_signal_col).value or "")
                fill = None
                if signal == "Beckett much higher":
                    fill = PatternFill("solid", fgColor="F4CCCC")
                elif signal == "Beckett higher":
                    fill = PatternFill("solid", fgColor="FCE5CD")
                elif signal == "Beckett lower":
                    fill = PatternFill("solid", fgColor="D9EAD3")
                elif signal == "Beckett much lower":
                    fill = PatternFill("solid", fgColor="B6D7A8")
                elif signal == "Close":
                    fill = PatternFill("solid", fgColor="D9EAF7")
                elif signal in {"No midpoint comparison", "Site price not numeric", "Beckett price unavailable"}:
                    fill = PatternFill("solid", fgColor="EEEEEE")
                if fill:
                    ws.cell(row=row_idx, column=price_signal_col).fill = fill
            for cell in ws[row_idx]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

    ws = wb.active
    scope_title = {
        "legacy": "Legacy Beckett Pricing",
        "nonlegacy": "Non-Legacy Beckett Pricing",
        "all": "Full Beckett Pricing",
    }.get(scope, "Beckett Pricing")
    ws.title = scope_title
    populate_sheet(ws, rows)

    review_rows = [row for row in rows if "review" in str(row.get("Match Status", "")).lower()]
    unresolved_rows = [
        row
        for row in rows
        if row.get("Match Status") in {"No Beckett result found", "No confident Beckett match"}
    ]
    review_rows = sorted(
        review_rows,
        key=lambda row: (
            str(row.get("Review Bucket", "")),
            str(row.get("Category", "")),
            safe_float(row.get("Match Confidence")),
            str(row.get("Title", "")),
        ),
    )
    unresolved_rows = sorted(
        unresolved_rows,
        key=lambda row: (str(row.get("Category", "")), str(row.get("Title", ""))),
    )
    if review_rows:
        populate_sheet(wb.create_sheet("Review Needed"), review_rows)
    if unresolved_rows:
        populate_sheet(wb.create_sheet("No Beckett Match"), unresolved_rows)
    action_rows = [
        row
        for row in rows
        if row.get("Match Status") != "Matched" and row.get("Match Status") != "Not searched - non sports-card row"
    ]
    action_rows = sorted(
        action_rows,
        key=lambda row: (
            str(row.get("Category", "")),
            str(row.get("Review Bucket", "")),
            99 if row.get("Match Confidence") is None else 0,
            safe_float(row.get("Match Confidence")),
            str(row.get("Title", "")),
        ),
    )
    if action_rows:
        populate_sheet(wb.create_sheet("Action Queue"), action_rows)
    price_review_rows = [
        row
        for row in rows
        if row.get("Price Signal") in {"Beckett much higher", "Beckett higher", "Beckett lower", "Beckett much lower"}
    ]
    price_review_rows = sorted(
        price_review_rows,
        key=lambda row: abs(safe_float(row.get("Midpoint Delta %", "0"))),
        reverse=True,
    )
    if price_review_rows:
        populate_sheet(wb.create_sheet("Pricing Review"), price_review_rows)

    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary_counts = Counter(row.get("Match Status", "") for row in rows)
    sport_rows = sum(1 for row in rows if row.get("Category") in SPORT_CATEGORIES)
    total_label = {
        "legacy": "Total legacy rows",
        "nonlegacy": "Total non-legacy rows",
        "all": "Total catalog rows",
    }.get(scope, "Total rows")
    summary_rows = [
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        (total_label, len(rows)),
        ("Sports-card rows searched", sport_rows),
        ("Non-sports rows included", len(rows) - sport_rows),
    ]
    summary_rows.extend(sorted(summary_counts.items(), key=lambda item: str(item[0] or "")))
    for item in summary_rows:
        summary.append(list(item))
    summary.append([])
    summary.append(["Review bucket", "Count"])
    from collections import Counter as _Counter
    for bucket, count in sorted(_Counter(row.get("Review Bucket", "") for row in rows).items(), key=lambda item: str(item[0] or "")):
        summary.append([bucket, count])
    summary.append([])
    summary.append(["Availability", "Count"])
    for availability, count in sorted(
        _Counter(row.get("Beckett Availability", "") for row in rows).items(),
        key=lambda item: str(item[0] or ""),
    ):
        summary.append([availability, count])
    summary.append([])
    summary.append(["Price signal", "Count"])
    for signal, count in sorted(_Counter(row.get("Price Signal", "") for row in rows).items(), key=lambda item: str(item[0] or "")):
        summary.append([signal, count])
    summary.append([])
    summary.append(["Category", "Matched", "Review", "Unresolved", "Non-sports"])
    category_counts = _Counter()
    for row in rows:
        category = row.get("Category", "")
        status = str(row.get("Match Status", ""))
        if status == "Matched":
            category_counts[(category, "Matched")] += 1
        elif "review" in status.lower():
            category_counts[(category, "Review")] += 1
        elif status in {"No Beckett result found", "No confident Beckett match"}:
            category_counts[(category, "Unresolved")] += 1
        elif "non sports-card" in status.lower():
            category_counts[(category, "Non-sports")] += 1
    for category in sorted({row.get("Category", "") for row in rows}, key=lambda value: str(value or "")):
        summary.append(
            [
                category,
                category_counts[(category, "Matched")],
                category_counts[(category, "Review")],
                category_counts[(category, "Unresolved")],
                category_counts[(category, "Non-sports")],
            ]
        )
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row_num in range(1, summary.max_row + 1):
        if summary.cell(row=row_num, column=1).value in {"Review bucket", "Availability", "Price signal"}:
            for col_num in (1, 2):
                summary.cell(row=row_num, column=col_num).fill = header_fill
                summary.cell(row=row_num, column=col_num).font = header_font
        if summary.cell(row=row_num, column=1).value == "Category":
            for col_num in range(1, 6):
                summary.cell(row=row_num, column=col_num).fill = header_fill
                summary.cell(row=row_num, column=col_num).font = header_font
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 24
    summary.column_dimensions["C"].width = 14
    summary.column_dimensions["D"].width = 14
    summary.column_dimensions["E"].width = 14

    wb.save(output_path)


def verify_workbook(output_path: Path, expected_rows: int, scope: str) -> None:
    wb = load_workbook(output_path, read_only=False, data_only=False)
    scope_title = {
        "legacy": "Legacy Beckett Pricing",
        "nonlegacy": "Non-Legacy Beckett Pricing",
        "all": "Full Beckett Pricing",
    }.get(scope, "Beckett Pricing")
    ws = wb[scope_title]
    actual_rows = ws.max_row - 1
    if actual_rows != expected_rows:
        raise RuntimeError(f"Workbook row count mismatch: expected {expected_rows}, found {actual_rows}")
    if not wb["Summary"].max_row > 1:
        raise RuntimeError("Summary sheet did not populate.")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--products", default="products.json")
    parser.add_argument("--output", required=True)
    parser.add_argument("--cache", required=True)
    parser.add_argument("--scope", choices=["legacy", "nonlegacy", "all"], default="legacy")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--delay", type=float, default=0.2)
    parser.add_argument("--save-every", type=int, default=10)
    parser.add_argument("--refresh-nonmatched", action="store_true")
    parser.add_argument("--refresh-all", action="store_true")
    args = parser.parse_args()

    products_path = Path(args.products)
    output_path = Path(args.output)
    cache_path = Path(args.cache)
    email = os.environ.get("BECKETT_EMAIL", "")
    password = os.environ.get("BECKETT_PASSWORD", "")

    products = selected_products(products_path, args.scope)
    if args.limit:
        products = products[: args.limit]
    cache = load_cache(cache_path)
    session = None
    if any(product.get("category") in SPORT_CATEGORIES for product in products):
        if not email or not password:
            raise RuntimeError("Set BECKETT_EMAIL and BECKETT_PASSWORD environment variables before running.")
        session = create_session(email, password)

    rows: list[dict[str, Any]] = []
    for idx, product in enumerate(products, start=1):
        if args.refresh_all:
            cache.setdefault("completed", {}).pop(str(product.get("id")), None)
        row = lookup_product(
            session,
            product,
            cache,
            args.delay,
            refresh_nonmatched=args.refresh_nonmatched or args.refresh_all,
        )
        rows.append(row)
        if idx % args.save_every == 0:
            save_cache(cache_path, cache)
            print(f"Processed {idx}/{len(products)} legacy rows", flush=True)
    save_cache(cache_path, cache)
    build_workbook(rows, output_path, args.scope)
    verify_workbook(output_path, len(rows), args.scope)
    counts = Counter(row.get("Match Status", "") for row in rows)
    print(json.dumps({"output": str(output_path), "rows": len(rows), "status_counts": counts}, default=dict, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("Interrupted.", file=sys.stderr)
        raise SystemExit(130)
