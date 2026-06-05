#!/usr/bin/env python3
"""Ensure confirmed Beckett Legacy rows exist in the storefront catalog.

The Beckett Legacy workbook is the reviewed source of truth for legacy card
matches. This reconciliation pass fills a gap that the normal updater does not:
if a confirmed workbook row is missing from products.json, it restores the
record from an older catalog snapshot when possible, or builds a new storefront
record from the workbook row plus the matching legacy photo.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = Path.home() / "Documents" / "eBay Docs" / "Listing Automation" / "Beckett Legacy Pricing Review.xlsx"
OLD_CATALOG_ROOT = Path("C:/Users/djwan/Downloads/djshouseofcards-next-fixes-applied")
PHOTO_ROOT = Path("H:/My Drive/Unused Assets/unused-legacy-photos")
REPORT_PATH = ROOT / "outputs" / "beckett-legacy-site-sync-report.json"

sys.path.insert(0, str(ROOT / "scripts"))
import apply_beckett_legacy_updates as legacy  # noqa: E402

PRODUCT_FILES = legacy.PRODUCT_FILES
DATA_BUNDLE_FILES = legacy.DATA_BUNDLE_FILES

SPORT_CONFIG = {
    "Baseball": {"league": "MLB", "asset_dir": "baseball-cards"},
    "Basketball": {"league": "NBA", "asset_dir": "basketball-cards"},
    "Football": {"league": "NFL", "asset_dir": "football-cards"},
}

NEW_PHOTO_DIRS = {
    "Baseball": [
        PHOTO_ROOT / "baseball_cardx" / "New Site Listings",
        PHOTO_ROOT / "baseball_cardx" / "Replace Photos",
        PHOTO_ROOT / "baseball_cardx",
    ],
    "Basketball": [PHOTO_ROOT / "basketball_cards"],
    "Football": [PHOTO_ROOT / "football_cards"],
}

# Explicit team corrections are deliberately narrow. These rows were newly
# matched in Beckett but not present in the storefront, so the workbook does not
# yet carry team values for all of them.
TEAM_BY_PRODUCT_ID = {
    2983: "Arizona Diamondbacks",
    2984: "Arizona Diamondbacks",
    2985: "Chicago White Sox",
    2986: "Chicago White Sox",
    2987: "Chicago White Sox",
    2988: "Chicago White Sox",
    2989: "Chicago White Sox",
    2990: "Texas Rangers",
    2991: "Minnesota Twins",
    2992: "New York Yankees",
    2993: "Milwaukee Braves",
    2994: "Arizona Diamondbacks",
    2995: "New York Yankees",
    2996: "Washington Nationals",
    2997: "Washington Nationals",
    2998: "Washington Nationals",
    2999: "Minnesota Twins",
    3000: "Minnesota Twins",
    3001: "Minnesota Twins",
    3002: "Minnesota Twins",
    3003: "Minnesota Twins",
    3004: "Minnesota Twins",
    3005: "Minnesota Twins",
    3006: "Minnesota Twins",
    3007: "Minnesota Twins",
    3008: "Minnesota Twins",
    3009: "Minnesota Twins",
    3010: "Minnesota Twins",
    3011: "Minnesota Twins",
    3012: "Minnesota Twins",
    3013: "Minnesota Twins",
    3014: "Minnesota Twins",
    3015: "Minnesota Twins",
    3016: "Minnesota Twins",
    3017: "Minnesota Twins",
    3018: "Minnesota Twins",
    3019: "Minnesota Twins",
    3020: "Los Angeles Angels",
    3021: "Minnesota Twins",
    3022: "Minnesota Twins",
    3023: "Washington Nationals",
    3024: "Minnesota Twins",
    3025: "Arizona Diamondbacks",
    3026: "Los Angeles Dodgers",
    3027: "Seattle Mariners",
    3028: "Philadelphia Phillies",
    3029: "Seattle Mariners",
    3030: "St. Louis Cardinals",
    3031: "Pittsburgh Pirates",
    3032: "Arizona Diamondbacks",
}

MANUFACTURER_PATTERNS = [
    ("Bowman", r"\bBowman\b"),
    ("Topps", r"\bTopps\b"),
    ("Donruss", r"\bDonruss\b"),
    ("Panini", r"\bPanini\b"),
    ("Leaf", r"\bLeaf\b"),
    ("Fleer", r"\bFleer\b"),
    ("Upper Deck", r"\bUpper\s+Deck\b|\bUD\b"),
    ("Score", r"\bScore\b"),
    ("SAGE", r"\bSAGE\b"),
    ("Press Pass", r"\bPress\s+Pass\b"),
    ("O-Pee-Chee", r"\bO-?Pee-?Chee\b"),
]

INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1F]')
IMAGE_EXTENSIONS = {".avif", ".gif", ".jpeg", ".jpg", ".png", ".webp"}
GRADE_RE = re.compile(r"\b(PSA|BGS|SGC|CGC|HGA|BCCG|GAI)\s*(10|[1-9](?:\.\d)?)\b", re.I)


def normalize_spaces(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def safe_filename_stem(value: str, max_length: int = 180) -> str:
    cleaned = INVALID_FILENAME_CHARS.sub(" - ", value)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .-_")
    return (cleaned or "legacy-product-photo")[:max_length].rstrip(" .-_")


def product_id(row: dict[str, Any]) -> int | None:
    return legacy.product_id_from_value(row.get("Product ID"))


def parse_year(*values: Any) -> int | None:
    for value in values:
        match = re.search(r"\b(19\d{2}|20\d{2})\b", str(value or ""))
        if match:
            return int(match.group(1))
    return None


def parse_grade_label(*values: Any) -> str:
    for value in values:
        match = GRADE_RE.search(str(value or ""))
        if match:
            grade = float(match.group(2))
            grade_label = str(int(grade)) if grade.is_integer() else str(grade).rstrip("0").rstrip(".")
            return f"{match.group(1).upper()} {grade_label}"
    return ""


def append_grade_to_name(name: str, grade_label: str) -> str:
    if not grade_label or grade_label.lower() in name.lower():
        return name
    return f"{name} {grade_label}"


def manufacturer_from_title(title: str) -> str:
    for manufacturer, pattern in MANUFACTURER_PATTERNS:
        if re.search(pattern, title, flags=re.I):
            return manufacturer
    return ""


def card_features(*values: Any) -> list[str]:
    text = " ".join(str(value or "") for value in values).lower()
    features: list[str] = []
    if re.search(r"\b(auto|autograph|autographs|signatures?|signed)\b", text):
        features.append("Autograph")
    if re.search(r"\b(relic|relics|jersey|jsy|patch|bat|swatch|materials?|memorabilia)\b", text):
        features.append("Memorabilia")
    if re.search(r"(?:/|\bnumbered\s+to\s+)(?:10|25|50|75|99|100|150|199|249|299|399|499|600|999)\b", text):
        features.append("Serial Numbered")
    return features


def player_from_title(title: str) -> str:
    # Beckett titles generally end with the player/subject after the checklist
    # number. Keep multi-player slashes because the UI can show that cleanly.
    cleaned = re.sub(r"\b(19\d{2}|20\d{2})(?:-\d{2})?\b", " ", title)
    cleaned = re.sub(r"#\s*[A-Za-z0-9.-]+[A-Za-z]?", " ", cleaned)
    cleaned = re.sub(
        r"\b(Topps|Bowman|Chrome|Heritage|Donruss|Elite|Leaf|Valiant|Draft|Panini|Prizm|Finest|Gypsy|Queen|Allen|Ginter|Stadium|Club|Museum|Collection|Tribute|Autographs?|Relics?|Refractors?|Rookie|Prospect|Framed|Mini|Blue|Green|Orange|Purple|White|Ink|Sepia|Gold|Black|Update|Now|Diamond|Icons|Inception|Supreme|Tier|One|Archives|Snapshots|Five|Star|Transcendent|Definitive|Buybacks|Clubhouse)\b",
        " ",
        cleaned,
        flags=re.I,
    )
    cleaned = re.sub(r"\b(SP|RC|BW|EXCH|USA|A|B|C)\b", " ", cleaned)
    cleaned = re.sub(r"/\d+\b", " ", cleaned)
    return normalize_spaces(cleaned.replace("&", "/"))


def price_for_row(row: dict[str, Any], cache: dict[str, Any]) -> tuple[str, float | None]:
    label, basis = legacy.best_price_for_row(row, cache)
    if basis == "real_time_na":
        return "N/A", None
    return label, legacy.midpoint_from_label(label)


def row_product_name(row: dict[str, Any]) -> str:
    name = normalize_spaces(row.get("Beckett Matched Title") or row.get("Title"))
    grade_label = parse_grade_label(row.get("Current Grade / Condition"), row.get("Title"))
    return append_grade_to_name(name, grade_label)


def desired_condition(row: dict[str, Any], name: str) -> str:
    condition = normalize_spaces(row.get("Current Grade / Condition"))
    if condition:
        return condition
    grade_label = parse_grade_label(name, row.get("Title"))
    if grade_label:
        return grade_label
    return "Ungraded | Near Mint or Better"


def matching_photo_files(category: str) -> list[Path]:
    files: list[Path] = []
    for folder in NEW_PHOTO_DIRS.get(category, []):
        if folder.exists():
            files.extend(path for path in folder.iterdir() if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS)
    return files


def best_photo_for_row(row: dict[str, Any], files: list[Path], claimed: set[Path]) -> tuple[Path | None, str]:
    raw_title = normalize_spaces(row.get("Title"))
    beckett_title = normalize_spaces(row.get("Beckett Matched Title"))
    raw_key = normalize_key(raw_title)
    beckett_key = normalize_key(beckett_title)

    exact = [path for path in files if normalize_key(path.stem) == raw_key and path not in claimed]
    if exact:
        return exact[0], "exact raw title"

    exact = [path for path in files if normalize_key(path.stem) == beckett_key and path not in claimed]
    if exact:
        return exact[0], "exact Beckett title"

    wanted_tokens = set((raw_key + " " + beckett_key).split())
    best: tuple[float, Path | None] = (0.0, None)
    for path in files:
        if path in claimed:
            continue
        candidate_tokens = set(normalize_key(path.stem).split())
        if not candidate_tokens:
            continue
        score = len(wanted_tokens & candidate_tokens) / len(wanted_tokens | candidate_tokens)
        if score > best[0]:
            best = (score, path)
    if best[1] and best[0] >= 0.55:
        return best[1], f"token score {best[0]:.2f}"
    return None, "no confident photo match"


def copy_asset(source: Path, relative_target: str) -> None:
    target = ROOT / relative_target
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.exists():
        shutil.copy2(source, target)


def target_path_for_source_photo(product: dict[str, Any], source: Path, claimed_targets: dict[str, str]) -> str:
    """Map a matched source photo to a stable, title-based storefront path.

    Multiple physical cards can legitimately share one Beckett title. When that
    happens, keep the first photo at the exact title filename and suffix the
    later copies so one card cannot overwrite another card's image.
    """
    category = normalize_spaces(product.get("category")) or "Baseball"
    config = SPORT_CONFIG.get(category, SPORT_CONFIG["Baseball"])
    stem = safe_filename_stem(normalize_spaces(product.get("name")) or f"Product {product.get('id')}")
    suffix = source.suffix.lower()
    source_key = str(source).lower()
    raw_title = normalize_spaces((product.get("metadata") or {}).get("cardsAddRawTitle"))
    copy_marker = re.search(r"\(#?(\d+)\)", raw_title)

    candidate_stems = [stem]
    if copy_marker:
        candidate_stems.append(f"{stem} - {copy_marker.group(1)}")
    candidate_stems.append(f"{stem} - Product {product.get('id')}")

    for candidate_stem in candidate_stems:
        candidate = str(PurePosixPath("assets", config["asset_dir"], "legacy-additions", f"{candidate_stem}{suffix}"))
        claimed_source = claimed_targets.get(candidate)
        if claimed_source is None or claimed_source == source_key:
            claimed_targets[candidate] = source_key
            return candidate

    # The loop above should always return, but keep a deterministic fallback.
    fallback = str(PurePosixPath("assets", config["asset_dir"], "legacy-additions", f"{stem} - Product {product.get('id')}{suffix}"))
    claimed_targets[fallback] = source_key
    return fallback


def retarget_source_photos(products_by_id: dict[int, dict[str, Any]]) -> list[dict[str, Any]]:
    retargeted: list[dict[str, Any]] = []
    claimed_targets: dict[str, str] = {}
    for pid in sorted(products_by_id):
        product = products_by_id[pid]
        metadata = product.get("metadata") if isinstance(product.get("metadata"), dict) else {}
        source_value = normalize_spaces(metadata.get("sourcePhoto") if metadata else "")
        if not source_value:
            continue
        source = Path(source_value)
        if not source.exists():
            continue
        target = target_path_for_source_photo(product, source, claimed_targets)
        copy_asset(source, target)
        old_image = product.get("image")
        product["image"] = target
        gallery = product.get("imageGallery") if isinstance(product.get("imageGallery"), list) else []
        if gallery:
            product["imageGallery"] = [target if item == old_image or index == 0 else item for index, item in enumerate(gallery)]
        else:
            product["imageGallery"] = [target]
        if old_image != target:
            retargeted.append({"id": pid, "oldImage": old_image, "newImage": target})
    return retargeted


def copy_old_product_assets(product: dict[str, Any], old_root: Path) -> list[str]:
    copied: list[str] = []
    paths = [product.get("image"), *(product.get("imageGallery") or [])]
    for path in {normalize_spaces(item) for item in paths if item}:
        if not path.startswith("assets/"):
            continue
        target = ROOT / path
        source = old_root / path
        if not target.exists() and source.exists():
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)
            copied.append(path)
    return copied


def build_missing_product(row: dict[str, Any], source_photo: Path | None, source_reason: str) -> dict[str, Any]:
    pid = product_id(row)
    if pid is None:
        raise ValueError("Cannot build product without Product ID")

    category = normalize_spaces(row.get("Category")) or "Baseball"
    config = SPORT_CONFIG.get(category, SPORT_CONFIG["Baseball"])
    name = row_product_name(row)
    year = legacy.product_id_from_value(row.get("Site Year")) or parse_year(name, row.get("Title"))
    player = normalize_spaces(row.get("Site Player / Athlete")) or player_from_title(name)
    team = normalize_spaces(row.get("Site Team / Publisher")) or TEAM_BY_PRODUCT_ID.get(pid, "")
    condition = desired_condition(row, name)
    manufacturer = manufacturer_from_title(name)
    cache = legacy.load_cache()
    price_label, price_value = price_for_row(row, cache)

    if source_photo:
        target_name = f"{safe_filename_stem(name)}{source_photo.suffix.lower()}"
        image_path = str(PurePosixPath("assets", config["asset_dir"], "legacy-additions", target_name)).replace("\\", "/")
        copy_asset(source_photo, image_path)
        gallery = [image_path]
    else:
        placeholder_name = f"placeholder-{category.lower()}.svg"
        image_path = str(PurePosixPath("assets", placeholder_name))
        gallery = [image_path]

    features = card_features(name, row.get("Title"), condition)
    description = (
        f"Legacy {category.lower()} listing matched to Beckett as {normalize_spaces(row.get('Beckett Matched Title'))}. "
        "Please review the photos for the exact card you will receive."
    )

    return {
        "id": pid,
        "name": name,
        "category": category,
        "team": team,
        "year": year,
        "condition": condition,
        "price": price_value,
        "priceLabel": price_label,
        "image": image_path,
        "imageGallery": gallery,
        "description": description,
        "photoHostPageUrl": "",
        "legacyImageLabel": safe_filename_stem(normalize_spaces(row.get("Title"))),
        "sourcePage": normalize_spaces(row.get("Source Page")) or f"{category} Cards",
        "league": config["league"],
        "sport": category,
        "playerAthlete": player,
        "displayPrice": price_label,
        "copyCount": 1,
        "itemPhotoUrl": "",
        "itemPhotoUrls": [],
        "htmlFullLink": normalize_spaces(row.get("Beckett URL")),
        "htmlImageUrls": [],
        "metadata": {
            "cardsAddRawTitle": normalize_spaces(row.get("Title")),
            "beckettUrl": normalize_spaces(row.get("Beckett URL")),
            "beckettTitle": normalize_spaces(row.get("Beckett Matched Title")),
            "manufacturer": manufacturer,
            "sourcePhoto": str(source_photo) if source_photo else "",
            "sourcePhotoReason": source_reason,
        },
        "isFeatured": False,
        "isDeleted": False,
        "sortRank": pid,
        **({"attributes": features} if features else {}),
    }


def enrich_product(product: dict[str, Any], row: dict[str, Any], cache: dict[str, Any]) -> dict[str, Any]:
    updated = legacy.update_product(product, row, cache)
    pid = int(updated["id"])
    name = row_product_name(row)
    if name:
        updated["name"] = name
    updated["condition"] = desired_condition(row, updated.get("name", ""))

    year = legacy.product_id_from_value(row.get("Site Year")) or parse_year(updated.get("name"), row.get("Title"))
    if year:
        updated["year"] = year

    player = normalize_spaces(row.get("Site Player / Athlete")) or updated.get("playerAthlete") or player_from_title(updated.get("name", ""))
    if player:
        updated["playerAthlete"] = player

    team = normalize_spaces(row.get("Site Team / Publisher")) or TEAM_BY_PRODUCT_ID.get(pid) or updated.get("team", "")
    updated["team"] = normalize_spaces(team)
    category = normalize_spaces(updated.get("category")) or normalize_spaces(row.get("Category")) or "Baseball"
    config = SPORT_CONFIG.get(category, SPORT_CONFIG["Baseball"])
    updated["league"] = updated.get("league") or config["league"]
    updated["sport"] = updated.get("sport") or category
    updated["displayPrice"] = updated.get("priceLabel", "")

    metadata = deepcopy(updated.get("metadata") or {})
    metadata["beckettUrl"] = normalize_spaces(row.get("Beckett URL"))
    metadata["beckettTitle"] = normalize_spaces(row.get("Beckett Matched Title"))
    manufacturer = manufacturer_from_title(updated.get("name", ""))
    if manufacturer:
        metadata["manufacturer"] = manufacturer
    updated["metadata"] = metadata

    features = card_features(updated.get("name"), row.get("Title"), updated.get("condition"))
    if features:
        updated["attributes"] = features
    elif "attributes" in updated:
        updated.pop("attributes", None)
    return updated


def rebuild_product_files(products: list[dict[str, Any]]) -> None:
    products = sorted(products, key=lambda item: int(item["id"]))
    for source_name, selector in PRODUCT_FILES.items():
        subset = selector(products)
        (ROOT / source_name).write_text(json.dumps(subset, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
        serialized = json.dumps(subset, ensure_ascii=False, separators=(",", ":"))
        (ROOT / DATA_BUNDLE_FILES[source_name]).write_text(
            f'window.DJ_PRELOADED_SOURCE = "{source_name}";\n'
            f"window.DJ_PRELOADED_PRODUCTS = {serialized}\n"
            ";\n",
            encoding="utf-8",
        )


def main() -> int:
    cache = legacy.load_cache()
    products = json.loads((ROOT / "products.json").read_text(encoding="utf-8"))
    product_by_id = {int(product["id"]): product for product in products}
    old_products: dict[int, dict[str, Any]] = {}
    if (OLD_CATALOG_ROOT / "products.json").exists():
        old_products = {
            int(product["id"]): product
            for product in json.loads((OLD_CATALOG_ROOT / "products.json").read_text(encoding="utf-8"))
        }

    wb = load_workbook(WORKBOOK_PATH, data_only=True)
    ws = wb["Legacy Beckett Pricing"]
    rows_by_id: dict[int, dict[str, Any]] = {}
    for row_index in range(2, ws.max_row + 1):
        row = legacy.row_to_dict(ws, row_index)
        pid = product_id(row)
        if pid is not None and legacy.has_confirmed_match(row):
            rows_by_id[pid] = row

    photo_files = {sport: matching_photo_files(sport) for sport in SPORT_CONFIG}
    claimed_photos: set[Path] = set()
    restored: list[dict[str, Any]] = []
    built: list[dict[str, Any]] = []
    changed: list[dict[str, Any]] = []
    copied_assets: list[str] = []
    unresolved_photos: list[dict[str, Any]] = []

    for pid in sorted(rows_by_id):
        row = rows_by_id[pid]
        before = deepcopy(product_by_id.get(pid))
        if before is None and pid in old_products:
            restored_product = deepcopy(old_products[pid])
            copied_assets.extend(copy_old_product_assets(restored_product, OLD_CATALOG_ROOT))
            product_by_id[pid] = restored_product
            restored.append({"id": pid, "name": restored_product.get("name")})
        elif before is None:
            category = normalize_spaces(row.get("Category")) or "Baseball"
            source_photo, reason = best_photo_for_row(row, photo_files.get(category, []), claimed_photos)
            if source_photo:
                claimed_photos.add(source_photo)
            else:
                unresolved_photos.append({"id": pid, "title": row.get("Title"), "reason": reason})
            product_by_id[pid] = build_missing_product(row, source_photo, reason)
            built.append({"id": pid, "name": product_by_id[pid].get("name"), "photoReason": reason})

        updated = enrich_product(product_by_id[pid], row, cache)
        if updated != product_by_id[pid] or before is None:
            product_by_id[pid] = updated
            changed.append(
                {
                    "id": pid,
                    "oldName": (before or {}).get("name"),
                    "newName": updated.get("name"),
                    "team": updated.get("team"),
                    "year": updated.get("year"),
                    "priceLabel": updated.get("priceLabel"),
                }
            )

    retargeted_photos = retarget_source_photos(product_by_id)
    rebuild_product_files(list(product_by_id.values()))

    report = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "workbook": str(WORKBOOK_PATH),
        "matchedWorkbookRows": len(rows_by_id),
        "restoredFromOldCatalog": len(restored),
        "builtFromWorkbook": len(built),
        "changedProducts": len(changed),
        "copiedOldAssets": copied_assets,
        "retargetedPhotoCount": len(retargeted_photos),
        "retargetedPhotos": retargeted_photos,
        "unresolvedPhotoCount": len(unresolved_photos),
        "unresolvedPhotos": unresolved_photos,
        "restored": restored,
        "built": built,
        "changed": changed,
        "finalProductCount": len(product_by_id),
    }
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not unresolved_photos else 1


if __name__ == "__main__":
    raise SystemExit(main())
