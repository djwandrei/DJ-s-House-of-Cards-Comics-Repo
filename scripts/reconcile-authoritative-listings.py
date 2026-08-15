#!/usr/bin/env python3
"""Reconcile all non-legacy products to the authoritative Listings worksheet."""

from __future__ import annotations

import argparse
import copy
import hashlib
import html
import json
import re
import shutil
import subprocess
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit

from openpyxl import load_workbook


DEFAULT_WORKBOOK = Path(
    r"C:\Users\djwan\Downloads\Ebay Bulk Upload (Final) - Photo Links Updated 8-15-25 2.xlsx"
)
SOURCE_WORKBOOK = DEFAULT_WORKBOOK.name
SOURCE_PAGE = "Non-Legacy Listings"
SOURCE_SHEET = "Listings"
EXPECTED_HEADER_COUNT = 45
REQUIRED_HEADERS = {
    "Category ID",
    "Category Name",
    "Title",
    "Start price",
    "Quantity",
    "Item photo URL",
    "HTML Full Link",
    "Condition ID",
    "Description",
    "C:Manufacturer",
    "C:Set",
    "C:Year Manufactured",
    "C:Player/Athlete",
    "C:Sport",
    "C:League",
    "C:Team",
    "C:Autographed",
    "C:Type",
}
URL_RE = re.compile(r"https?://[^\s\"'<>|,]+", re.IGNORECASE)
HREF_RE = re.compile(r"""href\s*=\s*["']([^"']+)["']""", re.IGNORECASE)
YEAR_RE = re.compile(r"\b(18\d{2}|19\d{2}|20\d{2})\b")
CONDITION_SUFFIX_RE = re.compile(r"\s*-\s*\(ID:\s*[^)]+\)\s*$", re.IGNORECASE)
NON_WORD_RE = re.compile(r"[^a-z0-9]+")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workbook",
        default=str(DEFAULT_WORKBOOK),
    )
    parser.add_argument("--products", default="products.json")
    parser.add_argument("--output-dir", default="outputs")
    parser.add_argument("--apply", action="store_true")
    return parser.parse_args()


def json_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalized(value: Any) -> str:
    return NON_WORD_RE.sub(" ", text(value).lower()).strip()


def normalize_url(url: Any) -> str:
    value = html.unescape(text(url)).rstrip(".,;)")
    if not value:
        return ""
    try:
        parts = urlsplit(value)
        return urlunsplit(
            (parts.scheme.lower(), parts.netloc.lower(), parts.path, parts.query, "")
        )
    except ValueError:
        return value


def split_urls(value: Any) -> list[str]:
    if isinstance(value, list):
        candidates = value
    elif isinstance(value, dict):
        candidates = list(value.values())
    else:
        candidates = URL_RE.findall(html.unescape(text(value)))
    output: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        url = normalize_url(candidate)
        if url and url not in seen:
            seen.add(url)
            output.append(url)
    return output


def first_href(value: Any) -> str:
    match = HREF_RE.search(text(value))
    return normalize_url(match.group(1)) if match else ""


def read_listings(workbook_path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise RuntimeError(f"{workbook_path} has no {SOURCE_SHEET!r} sheet")
    sheet = workbook[SOURCE_SHEET]
    rows = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    if len(headers) != EXPECTED_HEADER_COUNT:
        raise RuntimeError(
            f"Expected {EXPECTED_HEADER_COUNT} {SOURCE_SHEET!r} headers; found {len(headers)}"
        )
    missing_headers = sorted(REQUIRED_HEADERS - set(headers))
    if missing_headers:
        raise RuntimeError(
            f"{SOURCE_SHEET!r} is missing required headers: {', '.join(missing_headers)}"
        )
    listings: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows, start=2):
        fields = {
            header: json_value(row[index] if index < len(row) else None)
            for index, header in enumerate(headers)
        }
        if not text(fields.get("Title")):
            continue
        if not text(fields.get("Category ID")):
            raise RuntimeError(f"{SOURCE_SHEET}!{row_number} has no Category ID")
        if not text(fields.get("Condition ID")):
            raise RuntimeError(f"{SOURCE_SHEET}!{row_number} has no Condition ID")
        listings.append(
            {
                "row": row_number,
                "fields": fields,
                "title": text(fields.get("Title")),
                "title_norm": normalized(fields.get("Title")),
                "photos": split_urls(fields.get("Item photo URL")),
            }
        )
    workbook.close()
    return headers, listings


def is_nonlegacy(product: dict[str, Any]) -> bool:
    metadata = product.get("metadata")
    return isinstance(metadata, dict) and isinstance(metadata.get("excelFields"), dict)


def product_photos(product: dict[str, Any]) -> list[str]:
    candidates: list[Any] = []
    for key in ("itemPhotoUrls", "imageGallery", "htmlImageUrls"):
        candidates.extend(split_urls(product.get(key)))
    candidates.extend(split_urls(product.get("itemPhotoUrl")))
    candidates.extend(split_urls(product.get("image")))
    output: list[str] = []
    seen: set[str] = set()
    for url in candidates:
        if url.startswith("http") and url not in seen:
            seen.add(url)
            output.append(url)
    return output


def existing_local_media(product: dict[str, Any] | None) -> list[str]:
    """Preserve verified local display media while workbook URLs stay authoritative metadata."""
    if not product:
        return []
    gallery = product.get("imageGallery")
    values = gallery if isinstance(gallery, list) else []
    output: list[str] = []
    seen: set[str] = set()
    for value in [product.get("image"), *values]:
        reference = text(value).replace("\\", "/")
        if (
            not reference.lower().startswith("assets/")
            or "placeholder" in reference.lower()
            or reference in seen
            or not Path(reference).is_file()
        ):
            continue
        seen.add(reference)
        output.append(reference)
    return output


def product_identity(product: dict[str, Any]) -> dict[str, str]:
    fields = (product.get("metadata") or {}).get("excelFields") or {}
    return {
        "title": normalized(product.get("name")),
        "sport": normalized(product.get("sport") or fields.get("C:Sport")),
        "player": normalized(product.get("playerAthlete") or fields.get("C:Player/Athlete")),
        "team": normalized(product.get("team") or fields.get("C:Team")),
        "league": normalized(product.get("league") or fields.get("C:League")),
        "year": normalized(product.get("year") or fields.get("C:Year Manufactured")),
        "price": normalized(product.get("price") or fields.get("Start price")),
        "cert": normalized(fields.get("CDA:Certification Number - (ID: 27503)")),
        "set": normalized(fields.get("C:Set")),
    }


def listing_identity(listing: dict[str, Any]) -> dict[str, str]:
    fields = listing["fields"]
    return {
        "title": listing["title_norm"],
        "sport": normalized(fields.get("C:Sport")),
        "player": normalized(fields.get("C:Player/Athlete")),
        "team": normalized(fields.get("C:Team")),
        "league": normalized(fields.get("C:League")),
        "year": normalized(fields.get("C:Year Manufactured")),
        "price": normalized(fields.get("Start price")),
        "cert": normalized(fields.get("CDA:Certification Number - (ID: 27503)")),
        "set": normalized(fields.get("C:Set")),
    }


def pair_score(listing: dict[str, Any], product: dict[str, Any]) -> tuple[float, dict[str, Any]]:
    left = listing_identity(listing)
    right = product_identity(product)
    title_ratio = SequenceMatcher(None, left["title"], right["title"]).ratio()
    shared = [
        key
        for key in ("sport", "player", "team", "league", "year", "price", "cert", "set")
        if left[key] and left[key] == right[key]
    ]
    differing = [
        key
        for key in ("sport", "player", "team", "league", "year", "cert", "set")
        if left[key] and right[key] and left[key] != right[key]
    ]
    score = title_ratio * 100 + len(shared) * 12 - len(differing) * 8
    if left["cert"] and left["cert"] == right["cert"]:
        score += 80
    return score, {
        "titleRatio": round(title_ratio, 4),
        "sharedFields": shared,
        "differingFields": differing,
    }


def match_rows(
    listings: list[dict[str, Any]], products: list[dict[str, Any]]
) -> tuple[dict[int, tuple[int, str, dict[str, Any]]], list[int], list[int]]:
    matches: dict[int, tuple[int, str, dict[str, Any]]] = {}
    unmatched_rows = {listing["row"] for listing in listings}
    unmatched_ids = {int(product["id"]) for product in products}
    listing_by_row = {listing["row"]: listing for listing in listings}
    product_by_id = {int(product["id"]): product for product in products}
    photos_by_row = {listing["row"]: set(listing["photos"]) for listing in listings}
    photos_by_id = {int(product["id"]): set(product_photos(product)) for product in products}

    def accept(row: int, product_id: int, method: str, evidence: dict[str, Any]) -> None:
        if row not in unmatched_rows or product_id not in unmatched_ids:
            return
        matches[row] = (product_id, method, evidence)
        unmatched_rows.remove(row)
        unmatched_ids.remove(product_id)

    tuple_rows: dict[tuple[str, ...], list[int]] = defaultdict(list)
    tuple_ids: dict[tuple[str, ...], list[int]] = defaultdict(list)
    for row, photos in photos_by_row.items():
        if photos:
            tuple_rows[tuple(sorted(photos))].append(row)
    for product_id, photos in photos_by_id.items():
        if photos:
            tuple_ids[tuple(sorted(photos))].append(product_id)
    for photo_tuple, rows in tuple_rows.items():
        ids = tuple_ids.get(photo_tuple, [])
        if len(rows) == 1 and len(ids) == 1:
            accept(rows[0], ids[0], "exact-photo-tuple", {"photoCount": len(photo_tuple)})

    overlap_pairs: list[tuple[int, float, int, int, dict[str, Any]]] = []
    url_to_ids: dict[str, set[int]] = defaultdict(set)
    for product_id in unmatched_ids:
        for url in photos_by_id[product_id]:
            url_to_ids[url].add(product_id)
    for row in unmatched_rows:
        candidate_ids: set[int] = set()
        for url in photos_by_row[row]:
            candidate_ids.update(url_to_ids.get(url, set()))
        for product_id in candidate_ids:
            overlap = len(photos_by_row[row] & photos_by_id[product_id])
            score, evidence = pair_score(listing_by_row[row], product_by_id[product_id])
            evidence["photoOverlap"] = overlap
            overlap_pairs.append((-overlap, -score, row, product_id, evidence))
    for _, _, row, product_id, evidence in sorted(overlap_pairs):
        accept(row, product_id, "photo-overlap", evidence)

    title_rows: dict[str, list[int]] = defaultdict(list)
    title_ids: dict[str, list[int]] = defaultdict(list)
    for row in unmatched_rows:
        title_rows[listing_by_row[row]["title_norm"]].append(row)
    for product_id in unmatched_ids:
        title_ids[normalized(product_by_id[product_id].get("name"))].append(product_id)
    for title_key in sorted(set(title_rows) & set(title_ids)):
        pairs: list[tuple[float, int, int, dict[str, Any]]] = []
        for row in title_rows[title_key]:
            for product_id in title_ids[title_key]:
                score, evidence = pair_score(listing_by_row[row], product_by_id[product_id])
                pairs.append((-score, row, product_id, evidence))
        for _, row, product_id, evidence in sorted(pairs):
            accept(row, product_id, "exact-title", evidence)

    fuzzy_pairs: list[tuple[float, int, int, dict[str, Any]]] = []
    for row in unmatched_rows:
        listing = listing_by_row[row]
        for product_id in unmatched_ids:
            score, evidence = pair_score(listing, product_by_id[product_id])
            strong_identity = (
                evidence["titleRatio"] >= 0.82
                or (
                    evidence["titleRatio"] >= 0.62
                    and len(evidence["sharedFields"]) >= 3
                    and len(evidence["differingFields"]) <= 1
                )
                or (
                    "cert" in evidence["sharedFields"]
                    and evidence["titleRatio"] >= 0.45
                )
            )
            if strong_identity:
                fuzzy_pairs.append((-score, row, product_id, evidence))
    for _, row, product_id, evidence in sorted(fuzzy_pairs):
        accept(row, product_id, "identity-score", evidence)

    return matches, sorted(unmatched_rows), sorted(unmatched_ids)


def parse_year(fields: dict[str, Any]) -> int | None:
    candidate = fields.get("C:Year Manufactured")
    try:
        year = int(candidate)
        return year if 1800 <= year <= 2200 else None
    except (TypeError, ValueError):
        match = YEAR_RE.search(text(fields.get("Title")))
        return int(match.group(1)) if match else None


def clean_condition(fields: dict[str, Any]) -> str:
    grader = text(fields.get("CD:Professional Grader - (ID: 27501)"))
    grade = text(fields.get("CD:Grade - (ID: 27502)"))
    if grader and grade:
        return f"{grader} {grade}".strip()
    raw = text(fields.get("CD:Card Condition - (ID: 40001)"))
    return CONDITION_SUFFIX_RE.sub("", raw).strip() or "Ungraded"


def category_for(fields: dict[str, Any]) -> str:
    sport = text(fields.get("C:Sport")).lower()
    if sport == "baseball":
        return "Baseball"
    if sport == "basketball":
        return "Basketball"
    if sport == "football":
        return "Football"
    return "Collectibles"


def feature_attributes(fields: dict[str, Any]) -> list[str]:
    raw = text(fields.get("C:Features"))
    values = re.split(r"[,;|]+", raw) if raw else []
    if text(fields.get("C:Autographed")).lower() == "yes":
        values.append("Autograph")
    output: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = text(value)
        key = normalized(cleaned)
        if cleaned and key and key not in seen:
            seen.add(key)
            output.append(cleaned)
    return output


def price_value(value: Any) -> float:
    try:
        price = round(float(value), 2)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(f"Invalid non-legacy Start price: {value!r}") from exc
    if price <= 0:
        raise RuntimeError(f"Non-legacy Start price must be positive: {value!r}")
    return price


def quantity_value(value: Any) -> int:
    try:
        number = float(value)
        quantity = int(number)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(f"Invalid non-legacy Quantity: {value!r}") from exc
    if quantity <= 0 or number != quantity:
        raise RuntimeError(f"Non-legacy Quantity must be a positive integer: {value!r}")
    return quantity


def build_product(
    listing: dict[str, Any],
    product_id: int,
    old_product: dict[str, Any] | None,
    match_method: str,
    match_evidence: dict[str, Any],
) -> dict[str, Any]:
    fields = copy.deepcopy(listing["fields"])
    photos = listing["photos"]
    local_media = existing_local_media(old_product)
    display_media = local_media or photos
    price = price_value(fields.get("Start price"))
    quantity = quantity_value(fields.get("Quantity"))
    price_label = f"${price:,.2f}"
    product = {
        "id": product_id,
        "name": text(fields.get("Title")),
        "category": category_for(fields),
        "team": text(fields.get("C:Team")),
        "year": parse_year(fields),
        "condition": clean_condition(fields),
        "price": price,
        "description": text(fields.get("Description")),
        "image": display_media[0] if display_media else "",
        "imageGallery": display_media,
        "sport": text(fields.get("C:Sport")),
        "league": text(fields.get("C:League")),
        "playerAthlete": text(fields.get("C:Player/Athlete")),
        "attributes": feature_attributes(fields),
        "htmlFullLink": text(fields.get("HTML Full Link")),
        "itemPhotoUrls": photos,
        "itemPhotoUrl": photos[0] if photos else "",
        "htmlImageUrls": photos,
        "photoHostPageUrl": first_href(fields.get("HTML Full Link")),
        "sourcePage": text((old_product or {}).get("sourcePage")) or SOURCE_PAGE,
        "metadata": {
            "sport": text(fields.get("C:Sport")),
            "league": text(fields.get("C:League")),
            "playerAthlete": text(fields.get("C:Player/Athlete")),
            "htmlFullLink": text(fields.get("HTML Full Link")),
            "photoHostPageUrl": first_href(fields.get("HTML Full Link")),
            "itemPhotoUrls": photos,
            "htmlImageUrls": photos,
            "sourceWorkbook": SOURCE_WORKBOOK,
            "sourceSheet": SOURCE_SHEET,
            "excelRowNumber": listing["row"],
            "metadataMatchMethod": match_method,
            "matchEvidence": match_evidence,
            "excelFields": fields,
        },
        "isFeatured": bool((old_product or {}).get("isFeatured", False)),
        "isDeleted": False,
        "sortRank": int((old_product or {}).get("sortRank") or 0),
        "priceLabel": price_label,
        "displayPrice": price_label,
    }
    if quantity > 1:
        product["copyCount"] = quantity
    for field in ("featured", "featuredRank", "featuredSortRank"):
        if old_product and field in old_product:
            product[field] = copy.deepcopy(old_product[field])
    if product["year"] is None:
        product.pop("year")
    return product


def remote_product(product: dict[str, Any]) -> dict[str, Any]:
    mapping = {
        "id": "id",
        "name": "name",
        "category": "category",
        "team": "team",
        "year": "year",
        "condition": "condition",
        "price": "price",
        "priceLabel": "price_label",
        "displayPrice": "display_price",
        "image": "image",
        "imageGallery": "image_gallery",
        "description": "description",
        "photoHostPageUrl": "photo_host_page_url",
        "legacyImageLabel": "legacy_image_label",
        "sourcePage": "source_page",
        "league": "league",
        "sport": "sport",
        "playerAthlete": "player_athlete",
        "copyCount": "copy_count",
        "itemPhotoUrl": "item_photo_url",
        "itemPhotoUrls": "item_photo_urls",
        "htmlFullLink": "html_full_link",
        "htmlImageUrls": "html_image_urls",
        "metadata": "metadata",
        "isFeatured": "is_featured",
        "isDeleted": "is_deleted",
        "sortRank": "sort_rank",
    }
    return {remote: product.get(local) for local, remote in mapping.items()}


def write_json(path: Path, value: Any, compact: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(
            value,
            handle,
            ensure_ascii=True,
            indent=None if compact else 2,
            separators=(",", ":") if compact else None,
        )
        handle.write("\n")


def main() -> int:
    global SOURCE_WORKBOOK
    args = parse_args()
    workbook_path = Path(args.workbook).resolve()
    SOURCE_WORKBOOK = workbook_path.name
    products_path = Path(args.products).resolve()
    output_dir = Path(args.output_dir).resolve()
    headers, listings = read_listings(workbook_path)
    products = json.loads(products_path.read_text(encoding="utf-8"))
    active_nonlegacy = [
        product for product in products if is_nonlegacy(product) and product.get("isDeleted") is not True
    ]
    all_nonlegacy_ids = {int(product["id"]) for product in products if is_nonlegacy(product)}
    legacy_products = [product for product in products if not is_nonlegacy(product)]
    matches, unmatched_rows, delete_active_ids = match_rows(listings, active_nonlegacy)
    old_by_id = {int(product["id"]): product for product in active_nonlegacy}
    next_id = max(int(product["id"]) for product in products) + 1
    new_ids: list[int] = []
    authoritative: list[dict[str, Any]] = []
    match_method_counts: Counter[str] = Counter()
    match_records: list[dict[str, Any]] = []
    for listing in listings:
        match = matches.get(listing["row"])
        if match:
            product_id, method, evidence = match
            old_product = old_by_id[product_id]
        else:
            product_id = next_id
            next_id += 1
            new_ids.append(product_id)
            method = "new-authoritative-row"
            evidence = {}
            old_product = None
        match_method_counts[method] += 1
        authoritative.append(
            build_product(listing, product_id, old_product, method, evidence)
        )
        match_records.append(
            {
                "excelRowNumber": listing["row"],
                "productId": product_id,
                "matchMethod": method,
                "oldTitle": text((old_product or {}).get("name")),
                "authoritativeTitle": listing["title"],
                "evidence": evidence,
            }
        )

    delete_ids = sorted(all_nonlegacy_ids - {int(product["id"]) for product in authoritative})
    final_products = sorted(legacy_products + authoritative, key=lambda product: int(product["id"]))
    authoritative_rows = {
        int(product["metadata"]["excelRowNumber"]) for product in authoritative
    }
    validation = {
        "authoritativeListingCount": len(listings),
        "finalNonlegacyCount": len(authoritative),
        "uniqueAuthoritativeRows": len(authoritative_rows),
        "uniqueFinalProductIds": len({int(product["id"]) for product in final_products}),
        "finalProductCount": len(final_products),
        "finalActiveProductCount": sum(product.get("isDeleted") is not True for product in final_products),
        "missingPhotoRows": sum(not listing["photos"] for listing in listings),
        "deletedActiveNonlegacyCount": len(delete_active_ids),
        "deletedAllNonlegacyCount": len(delete_ids),
    }
    if validation["finalNonlegacyCount"] != len(listings):
        raise RuntimeError("Final non-legacy count does not match Listings count")
    if validation["uniqueAuthoritativeRows"] != len(listings):
        raise RuntimeError("Authoritative row numbers are not unique")
    if validation["uniqueFinalProductIds"] != len(final_products):
        raise RuntimeError("Final product IDs are not unique")
    if validation["missingPhotoRows"]:
        raise RuntimeError("One or more authoritative rows has no photo URL")

    report = {
        "generatedAt": datetime.now().astimezone().isoformat(),
        "workbook": str(workbook_path),
        "workbookSha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "sheet": SOURCE_SHEET,
        "headerCount": len(headers),
        "before": {
            "productCount": len(products),
            "activeProductCount": sum(product.get("isDeleted") is not True for product in products),
            "activeNonlegacyCount": len(active_nonlegacy),
            "allNonlegacyCount": len(all_nonlegacy_ids),
        },
        "matching": {
            "methodCounts": dict(sorted(match_method_counts.items())),
            "newProductIds": new_ids,
            "unmatchedAuthoritativeRowsBeforeNewIds": unmatched_rows,
            "activeNonlegacyIdsNotMatched": delete_active_ids,
            "allNonlegacyIdsToDelete": delete_ids,
        },
        "validation": validation,
        "matches": match_records,
    }
    write_json(output_dir / "authoritative-listings-reconciliation.json", report)
    write_json(output_dir / "authoritative-listings-delete-ids-current.json", delete_ids)
    if args.apply:
        write_json(output_dir / "authoritative-listings-delete-ids.json", delete_ids)
    write_json(
        output_dir / "authoritative-listings-supabase-upsert.json",
        [remote_product(product) for product in authoritative],
        compact=True,
    )

    if args.apply:
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup = output_dir / f"products-before-authoritative-sync-{timestamp}.json"
        shutil.copy2(products_path, backup)
        write_json(products_path, final_products)
        subprocess.run(
            [
                "node",
                str(products_path.parent / "scripts" / "build-public-catalog.mjs"),
                "--optimize-segments",
            ],
            cwd=products_path.parent,
            check=True,
        )
        report["applied"] = True
        report["backup"] = str(backup)
        write_json(output_dir / "authoritative-listings-reconciliation.json", report)

    print(json.dumps({"applied": args.apply, **report["before"], **validation, **report["matching"]["methodCounts"]}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
